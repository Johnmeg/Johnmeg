# -*- coding: utf-8 -*-
"""Genera la PLANTILLA / ESQUEMA de una Especificación Funcional (Documento Funcional)
para la solución de planeación financiera en SAP Analytics Cloud (Fanalca, Ciudad Limpia,
Transprensa). Cada sección trae una guía «Qué debe contener» y tablas esqueleto."""
import sys, os
HERE=os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'generador_diseno'))
from docx_helpers import Builder, NAVY, NAVY2, BLUE_D, RED, GREEN, GREY
from docx.shared import Pt, RGBColor

TPL=os.path.join(HERE,'template_funcional.docx')
LOGO_BUILD=os.path.join(HERE,'logos','build.png')
LOGO_FAN=os.path.join(HERE,'logos','fanalca.png')

b=Builder(TPL)
b.setup_headers(LOGO_FAN)
# ajusta el texto del encabezado al tipo de documento
hp=b.doc.sections[0].header.paragraphs[0]
for r in hp.runs:
    if 'Documento de Diseño' in r.text:
        r.text='Especificación Funcional (Documento Funcional) · SAP Analytics Cloud'
    elif 'Programa SAP BUILD' in r.text:
        r.text='\tPlantilla y guía de contenido — Grupo Fanalca'

# ---------- helpers ----------
def fm(text):
    p=b.doc.add_paragraph()
    p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5)
    r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(13); r.font.bold=True
    r.font.color.rgb=RGBColor.from_string(BLUE_D); b._bottom_border(p); return p

def guia(lines):
    b.callout("Qué debe contener", lines, accent=NAVY2, fill="EAF1FB", icon="✔")

def ejemplo(text):
    b.para("Ejemplo — "+text, italic=True, size=8.8, color=GREY, space_after=8)

def EMPTY(ncol, k=2):
    return [[""]*ncol for _ in range(k)]

# =====================================================================
# PORTADA
# =====================================================================
b.spacer(30)
b.cover_logo(LOGO_BUILD, width=3.0)
b.spacer(16)
b.cover_title([
    ("ESPECIFICACIÓN FUNCIONAL", 22, BLUE_D, True, 4),
    ("Documento Funcional — Plantilla y guía de contenido", 14, NAVY, False, 3),
    ("Modelos de Planeación Financiera en SAP Analytics Cloud", 12, NAVY, False, 2),
    ("Programa SAP BUILD · Grupo Fanalca", 11, GREY, False, 20),
])
b.meta_table([
    ("Documento", "Especificación Funcional (plantilla para diligenciar)"),
    ("Sociedades", "Fanalca S.A.  ·  Ciudad Limpia  ·  Transprensa S.A.S."),
    ("Proceso", "Planeación financiera: Ingresos · Costos y Gastos · Estados Financieros"),
    ("Plataforma", "SAP Analytics Cloud (SAC) · SAP DataSphere · SAP S/4HANA"),
    ("Versión", "0.1 — Plantilla"),
    ("Fecha", "(por diligenciar)"),
    ("Estado", "Plantilla para diligenciar"),
    ("Clasificación", "Confidencial"),
])
b.spacer(28)
b.cover_logo(LOGO_FAN, width=2.9)
b.page_break()

# =====================================================================
# FRONT MATTER
# =====================================================================
b.callout("Cómo usar esta plantilla", [
    "Este documento define el ESQUEMA de una Especificación Funcional para la solución de planeación en SAC. Cada sección incluye un recuadro «Qué debe contener» con la guía de lo que se espera; reemplácelo por el contenido real del proyecto.",
    "Las filas y textos marcados como «Ejemplo» ilustran el dato esperado; reemplácelos o elimínelos al diligenciar.",
    "Complete una ficha por cada objeto funcional (modelo, Data Action, plantilla, reporte, interfaz) en la sección 5, y enlace todo en la matriz de trazabilidad (sección 8).",
], accent=NAVY2)

fm("Control de versiones")
b.table(["Versión","Fecha","Autor","Descripción del cambio","Estado"],
        [["0.1","(fecha)","(autor)","Creación de la plantilla","Plantilla"]]+EMPTY(5,3),
        widths=[0.8,1.0,1.4,2.6,0.8])

fm("Aprobaciones")
b.table(["Rol","Nombre","Responsabilidad","Fecha","Firma"],
        EMPTY(5,4), widths=[1.5,1.6,1.9,0.8,0.8])

fm("Distribución")
b.table(["Nombre / Área","Rol","Medio"], EMPTY(3,3), widths=[2.6,2.4,1.6])

fm("Documentos relacionados y referencias")
b.table(["ID","Documento","Versión","Ubicación / enlace"],
        [["DOC-01","Documento de Diseño de Solución — SAC (Fanalca, Ciudad Limpia, Transprensa)","2.0","(repositorio)"],
         ["DOC-02","Plan de seguimiento de la configuración de SAC","1.0","(repositorio)"]]+EMPTY(4,2),
        widths=[0.7,2.9,0.8,2.2])

fm("Contenido")
b.toc()
b.page_break()

# =====================================================================
# 1. INTRODUCCIÓN
# =====================================================================
b.h1("Introducción")
b.h2("Propósito del documento")
guia(["Objetivo del documento y qué solución funcional describe.",
      "Relación con el documento de diseño y con el proceso de planeación (presupuesto y forecast).",
      "A qué preguntas responde y para qué fase del proyecto sirve."])
ejemplo("Especifica, a nivel funcional, los modelos de planeación (Ingresos, Costos/Gastos y EEFF) a construir en SAC y las reglas que los desarrolladores/configuradores deben implementar.")

b.h2("Alcance funcional")
guia(["Procesos, sociedades y modelos incluidos.",
      "Lo que queda explícitamente EXCLUIDO (para evitar supuestos).",
      "Fases o entregables cubiertos por este documento."])
b.table(["Incluye","Excluye"],
        [["Modelos Ingresos, Costos/Gastos y EEFF de las 3 sociedades; carga del plan, cálculos, versiones, reportes y seguridad.",
          "Cierre contable y ejecución transaccional (permanecen en el ERP); Talento Humano y Flujo de Caja diario (en diseño)."]]+EMPTY(2,2),
        widths=[3.3,3.3])

b.h2("Audiencia y uso")
guia(["Perfiles a los que va dirigido (negocio, funcional, técnico, QA).",
      "Cómo lo usa cada perfil (validar, construir, probar)."])

b.h2("Definiciones, acrónimos y glosario")
guia(["Términos de SAP/SAC y del negocio necesarios para leer el documento.",
      "Remita al glosario extendido del anexo para la lista completa."])
b.table(["Término / Acrónimo","Definición"],
        [["SAC","SAP Analytics Cloud: plataforma de planificación, análisis y reporte."],
         ["Data Action","Proceso de cálculo en SAC (Advanced Formulas / Cross-Model Copy)."],
         ["P×Q","Precio × Cantidad: metodología de cálculo de ingresos (tarifa × kilos/toneladas)."],
         ["DAC","Data Access Control: control de acceso a datos por dimensión."]]+EMPTY(2,2),
        widths=[1.7,4.9])

b.h2("Supuestos y restricciones")
guia(["Supuestos sobre datos, sistemas, disponibilidad de personas o decisiones tomadas.",
      "Restricciones técnicas, de licenciamiento o de negocio.",
      "Impacto si un supuesto no se cumple."])
b.table(["ID","Supuesto / Restricción","Impacto si no se cumple"],
        [["SUP-01","Los maestros provienen de S/4HANA vía DataSphere y están depurados.","Retrabajo en la carga y cuadres."]]+EMPTY(3,2),
        widths=[0.7,3.7,2.2])

b.h2("Dependencias")
guia(["Dependencias internas (entre modelos/objetos) y externas (otros equipos/sistemas).",
      "Responsable y fecha comprometida de cada dependencia."])
b.table(["ID","Dependencia","Tipo (interna/externa)","Responsable"],
        EMPTY(4,3), widths=[0.7,3.3,1.4,1.2])

# =====================================================================
# 2. CONTEXTO Y OBJETIVOS DEL NEGOCIO
# =====================================================================
b.h1("Contexto y objetivos del negocio")
b.h2("Descripción del proceso actual (AS-IS)")
guia(["El proceso tal como opera hoy: sistemas, roles y pasos.",
      "Puntos de dolor que la solución debe resolver.",
      "Volúmenes, periodicidad y responsables.",
      "Detalle por sociedad en el documento de diseño (referencia DOC-01)."])

b.h2("Objetivos y beneficios esperados")
guia(["Objetivos empresariales (OE) medibles que persigue la solución.",
      "Beneficio esperado y métrica para verificarlo."])
b.table(["ID (OE)","Objetivo","Beneficio esperado","Métrica"],
        [["OE-01","Unificar el presupuesto en una plataforma corporativa.","Menos dispersión y reprocesos en Excel.","N.º de archivos / horas de consolidación"]]+EMPTY(4,2),
        widths=[0.8,2.4,2.2,1.2])

b.h2("Indicadores de éxito (KPIs)")
guia(["KPIs que medirán el éxito del proyecto y de la operación.",
      "Definición, meta y fuente del dato de cada KPI."])
b.table(["KPI","Definición","Meta","Fuente"], EMPTY(4,3), widths=[1.7,2.6,0.9,1.4])

b.h2("Actores, roles y responsabilidades (RACI)")
guia(["Quién es Responsable (R), quién Aprueba (A), quién es Consultado (C) e Informado (I).",
      "Una fila por actividad o entregable relevante del proceso."])
b.table(["Actividad / Entregable","R","A","C","I"],
        [["Carga del presupuesto","Planeación","Gerencia","Comercial","Dirección"]]+EMPTY(5,3),
        widths=[3.0,0.9,0.9,0.9,0.9])

# =====================================================================
# 3. REQUERIMIENTOS
# =====================================================================
b.h1("Requerimientos")
b.h2("Requerimientos funcionales")
guia(["Un requerimiento por fila, con identificador único (RF-###).",
      "Prioridad con criterio MoSCoW (Must / Should / Could / Won't).",
      "Origen (acta, sesión, reglamentación) y criterio de aceptación verificable."])
b.table(["ID","Requerimiento","Modelo / Proceso","Prioridad","Origen","Criterio de aceptación"],
        [["RF-001","Calcular ingresos por P×Q por regional y servicio.","Ingresos (Transprensa)","Must","Sesión AS-IS","El cálculo coincide con el histórico N-1 ±0%"]]+EMPTY(6,2),
        widths=[0.7,1.9,1.1,0.8,0.8,1.3])

b.h2("Reglas de negocio")
guia(["Reglas y fórmulas que gobiernan los cálculos y validaciones.",
      "Identificador (RN-###), descripción/fórmula y a qué modelo/objeto aplica."])
b.table(["ID","Regla","Descripción / fórmula","Aplica a"],
        [["RN-001","Crecimiento de tarifa","Tarifa N = Tarifa N-1 × (1 + IPC + ICTC)","Ingresos Transprensa"]]+EMPTY(4,2),
        widths=[0.7,1.6,3.1,1.2])

b.h2("Requerimientos no funcionales")
guia(["Requisitos de rendimiento, seguridad, disponibilidad, usabilidad, auditoría y escalabilidad.",
      "Cada uno con una métrica o meta objetiva."])
b.table(["ID","Categoría","Requerimiento","Métrica / meta"],
        [["RNF-01","Rendimiento","Tiempo de ejecución de los cálculos","< 60 s por Data Action"],
         ["RNF-02","Seguridad","Acceso por sociedad y rol (DAC)","100% de modelos con DAC"]]+EMPTY(4,2),
        widths=[0.7,1.3,2.8,1.8])

# =====================================================================
# 4. DISEÑO FUNCIONAL DE LA SOLUCIÓN (TO-BE)
# =====================================================================
b.h1("Diseño funcional de la solución (TO-BE)")
b.h2("Descripción de la solución")
guia(["Cómo funcionará el proceso una vez implementada la solución (TO-BE).",
      "Componentes SAC involucrados y cómo se relacionan.",
      "Diferencias clave frente al AS-IS."])

b.h2("Diagrama de proceso / flujo funcional")
guia(["Diagrama del flujo funcional (fuentes → carga → cálculo → consolidación → reporte).",
      "Incluir puntos de decisión, actores y sistemas."])
b.para("[ Insertar aquí el diagrama de flujo funcional ]", italic=True, color=GREY, align='center', size=9)

b.h2("Modelo de datos (dimensiones, medidas y jerarquías)")
guia(["Dimensiones de cada modelo, su tipo, jerarquías y origen del maestro.",
      "Medidas/cuentas y su significado.",
      "Dimensiones compartidas entre modelos."])
b.table(["Dimensión","Tipo","Descripción","Jerarquía","Origen del maestro"],
        [["Cuenta","Cuenta","Plan de cuentas / rubros","Sí","S/4HANA (SKA1/SKB1)"],
         ["CEBE","Genérica","Centro de beneficio / componente","Sí","S/4HANA (CEPC)"]]+EMPTY(5,2),
        widths=[1.1,0.9,2.2,0.8,1.6])

b.h2("Lógica de cálculo (Data Actions)")
guia(["Cada cálculo con identificador (DA-###), descripción/fórmula, entradas y salida.",
      "Orden de ejecución y versión destino.",
      "Referencie las reglas de negocio (RN) que implementa."])
b.table(["ID","Cálculo","Descripción / fórmula","Entradas","Salida (medida / versión)"],
        [["DA-001","P×Q de ingresos","Importe = Kilos × Tarifa (pesos/kg)","Kilos, Tarifa","Importe / Presupuesto"]]+EMPTY(5,2),
        widths=[0.7,1.5,2.2,1.0,1.2])

b.h2("Versiones y escenarios")
guia(["Versiones por categoría (Actual, Presupuesto, Forecast) y su uso.",
      "Escenarios/simulaciones y reglas de copia y bloqueo (data locking)."])
b.table(["Versión / Categoría","Público/Privado","Uso","Reglas (copia / bloqueo)"],
        [["Presupuesto","Público","Cifra oficial aprobada","Bloqueo al aprobar"]]+EMPTY(4,2),
        widths=[1.6,1.2,2.0,1.8])

b.h2("Integración e interfaces")
guia(["Interfaces de entrada/salida: fuente, objeto, destino, frecuencia y método.",
      "El real se integra desde S/4HANA vía DataSphere; el plan se carga por plantillas.",
      "Detalle del mapeo de campos en la sección 5.2."])
b.table(["ID","Fuente","Objeto / vista","Destino (modelo)","Frecuencia","Método"],
        [["INT-01","S/4HANA","CDS de saldos (real)","EEFF / Costos","Diaria","DataSphere"]]+EMPTY(6,2),
        widths=[0.7,1.1,1.5,1.4,0.9,1.0])

b.h2("Plantillas de entrada (Input Forms)")
guia(["Cada plantilla de captura: modelo, dimensiones en filas/columnas, medidas y opción de carga.",
      "Opciones de carga: Reemplazar / Acumular / Borrar y recargar."])
b.table(["ID","Plantilla","Modelo","Filas / Columnas","Medidas","Opción de carga"],
        EMPTY(6,3), widths=[0.7,1.6,1.2,1.7,0.8,0.9])

b.h2("Reportes y dashboards")
guia(["Cada salida (RPT-###): tipo (Story/Analytic App), audiencia, filtros y KPIs.",
      "Formato de exportación requerido (PDF/PPT/Excel)."])
b.table(["ID","Reporte / Dashboard","Tipo","Audiencia","Dimensiones / Filtros","Medidas / KPIs"],
        [["RPT-01","Presupuesto vs. Real","Story","Gerencia","Sociedad, CEBE, Periodo","Importe, Variación %"]]+EMPTY(6,2),
        widths=[0.7,1.6,0.8,1.1,1.5,0.9])

b.h2("Seguridad y autorizaciones")
guia(["Roles, modelos a los que acceden, nivel (Lectura/Escritura), DAC y tipo de licencia.",
      "Regla de segregación de funciones si aplica."])
b.table(["Rol","Descripción","Modelos","Acceso","DAC (dimensión: valores)","Licencia"],
        [["Planeación","Construye el plan","Los 3","Escritura","Sociedad: la propia","Planificación"],
         ["Gerencia","Consulta","Los 3","Lectura","Sociedad: la propia","BI / lectura"]]+EMPTY(6,2),
        widths=[1.0,1.4,0.8,0.9,1.6,1.0])

b.h2("Monedas y conversión")
guia(["Moneda de cada modelo, necesidad de conversión y tabla de tasas (TRM).",
      "Reglas de redondeo y moneda de reporte."])

b.h2("Validaciones, mensajes y manejo de errores")
guia(["Validaciones de datos y reglas de consistencia.",
      "Mensajes al usuario y comportamiento ante error."])
b.table(["ID","Validación / mensaje","Condición","Acción / mensaje al usuario"],
        EMPTY(4,2), widths=[0.7,2.1,1.9,1.9])

# =====================================================================
# 5. ESPECIFICACIÓN DETALLADA POR OBJETO
# =====================================================================
b.h1("Especificación detallada por objeto")
b.h2("Ficha de objeto (repetir por cada objeto funcional)")
guia(["Diligencie una ficha por cada modelo, Data Action, plantilla, reporte o interfaz.",
      "Debe permitir a un configurador construir el objeto sin ambigüedad."])
b.table(["Campo","Contenido (a diligenciar)"],
        [["ID del objeto",""],["Nombre",""],
         ["Tipo","Modelo / Data Action / Plantilla / Reporte / Interfaz"],
         ["Sociedad / Modelo",""],["Descripción funcional",""],
         ["Entradas (datos/dimensiones)",""],["Lógica / proceso (paso a paso)",""],
         ["Salidas (medidas/versión)",""],["Dependencias",""],
         ["Reglas de negocio asociadas (RN)",""],["Criterio de aceptación",""],
         ["Estado",""]],
        widths=[2.2,4.4])

b.h2("Mapeo de campos (origen → destino)")
guia(["Correspondencia campo a campo entre la fuente y el modelo SAC.",
      "Incluya transformaciones, tipo de dato y obligatoriedad."])
b.table(["Campo origen","Tipo","Transformación / regla","Campo destino","Obligatorio"],
        [["WERKS","Texto","Mapear a regional","CEBE","Sí"]]+EMPTY(5,2),
        widths=[1.3,0.9,2.1,1.5,0.8])

# =====================================================================
# 6. CASOS DE USO Y ESCENARIOS
# =====================================================================
b.h1("Casos de uso y escenarios")
b.h2("Casos de uso funcionales (repetir por caso)")
guia(["Un caso de uso por interacción relevante del usuario con la solución.",
      "Incluya flujo principal, alternos y excepciones."])
b.table(["Campo","Contenido (a diligenciar)"],
        [["ID (CU-###)",""],["Nombre",""],["Actor(es)",""],["Precondición",""],
         ["Flujo principal",""],["Flujos alternos / excepciones",""],["Postcondición",""]],
        widths=[2.2,4.4])

b.h2("Escenarios de simulación")
guia(["Escenarios de análisis «qué pasaría si» soportados por la solución.",
      "Variables de entrada y resultado esperado."])
b.table(["ID","Escenario","Supuestos / variables","Resultado esperado"],
        [["SIM-01","Nueva licitación (Transprensa)","Kilos, tarifa y meses de la licitación","Impacto en ingresos, costos y margen"]]+EMPTY(4,2),
        widths=[0.7,1.7,2.3,1.9])

# =====================================================================
# 7. DATOS MAESTROS Y MIGRACIÓN
# =====================================================================
b.h1("Datos maestros y migración")
b.h2("Objetos de datos maestros")
guia(["Maestros necesarios, su fuente, volumen aproximado y responsable.",
      "Indique cuáles se integran desde DataSphere y cuáles se cargan manualmente."])
b.table(["Objeto","Descripción","Volumen aprox.","Fuente","Responsable"],
        [["Cuentas","Plan de cuentas","(n.º)","S/4HANA","Contabilidad"]]+EMPTY(5,2),
        widths=[1.2,2.2,1.0,1.1,1.1])

b.h2("Estrategia de carga inicial y calidad de datos")
guia(["Cómo se realiza la carga inicial (histórico y maestros).",
      "Reglas de depuración, validación y criterios de calidad del dato."])

# =====================================================================
# 8. CRITERIOS DE ACEPTACIÓN Y PRUEBAS
# =====================================================================
b.h1("Criterios de aceptación y pruebas")
b.h2("Criterios de aceptación")
guia(["Condiciones objetivas para dar por aceptada la solución/objeto.",
      "Cada criterio debe ser verificable y trazable a un requerimiento."])

b.h2("Casos de prueba")
guia(["Casos de prueba unitaria, de integración y de aceptación (UAT).",
      "Pasos, datos y resultado esperado por caso."])
b.table(["ID","Escenario","Precondición / datos","Pasos","Resultado esperado","Tipo"],
        [["CP-001","P×Q de ingresos","Histórico N-1 cargado","Ejecutar DA-001","Importe = kilos × tarifa","Unitaria"]]+EMPTY(6,2),
        widths=[0.7,1.4,1.4,1.2,1.2,0.7])

b.h2("Matriz de trazabilidad")
guia(["Enlace requerimiento → objeto → caso de prueba, para demostrar cobertura.",
      "Una fila por requerimiento; su estado de verificación."])
b.table(["Requerimiento (RF)","Objeto (DA/RPT/INT)","Caso de prueba (CP)","Estado"],
        [["RF-001","DA-001","CP-001","Pendiente"]]+EMPTY(4,2),
        widths=[1.8,1.8,1.6,1.4])

# =====================================================================
# 9. SUPUESTOS, RIESGOS Y TEMAS ABIERTOS
# =====================================================================
b.h1("Riesgos y temas abiertos")
b.h2("Riesgos y mitigaciones")
guia(["Riesgos funcionales/técnicos, su probabilidad e impacto.",
      "Acción de mitigación y responsable."])
b.table(["ID","Riesgo","Prob.","Impacto","Mitigación","Responsable"],
        [["R-01","Maestros desalineados con el ERP","Media","Alto","Depuración previa en DataSphere","TI"]]+EMPTY(6,2),
        widths=[0.6,1.9,0.8,0.8,1.7,0.8])

b.h2("Temas abiertos / decisiones pendientes")
guia(["Preguntas o decisiones sin resolver que bloquean el diseño.",
      "Quién lo planteó, fecha, estado y resolución final."])
b.table(["ID","Tema / decisión pendiente","Planteado por","Fecha","Estado","Resolución"],
        EMPTY(6,3), widths=[0.6,2.3,1.1,0.8,0.8,1.0])

# =====================================================================
# 10. ANEXOS
# =====================================================================
b.h1("Anexos")
guia(["Glosario extendido de términos SAP/SAC y del negocio.",
      "Mockups / capturas de plantillas y reportes.",
      "Actas de decisiones y catálogos de referencia (dimensiones, maestros).",
      "Cualquier soporte que respalde el diseño funcional."])

fn=os.path.join(HERE,'Especificacion_Funcional_Plantilla_SAC.docx')
b.save(fn)
print("GUARDADO:", fn)
