import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import matplotlib.patheffects as pe
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy, os

# ─── COLORES CORPORATIVOS ─────────────────────────────────────────────────────
SAP_BLUE    = RGBColor(0x00, 0x61, 0x9A)   # #00619A
SAP_DARK    = RGBColor(0x00, 0x33, 0x69)   # #003369
SAP_GOLD    = RGBColor(0xF0, 0xAB, 0x00)   # #F0AB00
SAP_LIGHT   = RGBColor(0xE8, 0xF4, 0xFD)   # #E8F4FD
SAP_GREEN   = RGBColor(0x10, 0x7E, 0x3E)   # #107E3E
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_DARK   = RGBColor(0x35, 0x35, 0x35)
GRAY_MED    = RGBColor(0x75, 0x75, 0x75)
GRAY_LIGHT  = RGBColor(0xF2, 0xF2, 0xF2)
ORANGE      = RGBColor(0xE8, 0x6A, 0x19)   # accent

def hex_rgb(h):
    h = h.lstrip('#')
    return tuple(int(h[i:i+2],16)/255 for i in (0,2,4))

# ─── HELPERS DOCX ─────────────────────────────────────────────────────────────
def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  '%02X%02X%02X' % (rgb[0], rgb[1], rgb[2]))
    tcPr.append(shd)

def set_cell_borders(cell, sides=('top','bottom','left','right'), color='AAAAAA', sz=4):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for s in sides:
        b = OxmlElement(f'w:{s}')
        b.set(qn('w:val'),   'single')
        b.set(qn('w:sz'),    str(sz))
        b.set(qn('w:color'), color)
        tcBorders.append(b)
    tcPr.append(tcBorders)

def para_font(para, size_pt, bold=False, color=None, align=None, italic=False, space_before=0, space_after=0):
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after  = Pt(space_after)
    if align: para.alignment = align
    for run in para.runs:
        run.font.size   = Pt(size_pt)
        run.font.bold   = bold
        run.font.italic = italic
        if color: run.font.color.rgb = color

def add_heading(doc, text, level=1, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    if level == 1:
        run.font.size  = Pt(16)
        run.font.bold  = True
        run.font.color.rgb = SAP_DARK
        # underline bar via bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bot  = OxmlElement('w:bottom')
        bot.set(qn('w:val'),   'single')
        bot.set(qn('w:sz'),    '6')
        bot.set(qn('w:color'), '003369')
        pBdr.append(bot)
        pPr.append(pBdr)
    elif level == 2:
        run.font.size  = Pt(13)
        run.font.bold  = True
        run.font.color.rgb = SAP_BLUE
    elif level == 3:
        run.font.size  = Pt(11)
        run.font.bold  = True
        run.font.color.rgb = GRAY_DARK
    return p

def add_body(doc, text, size=10, color=None, bold=False, italic=False, space_after=4):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    run.font.color.rgb = color if color else GRAY_DARK
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p

def add_bullet(doc, text, size=10):
    p   = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size      = Pt(size)
    run.font.color.rgb = GRAY_DARK
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    return p

def header_row(table, headers, bg=SAP_DARK, fg=WHITE, widths=None):
    row = table.rows[0]
    for i, h in enumerate(headers):
        cell = row.cells[i]
        cell.text = ''
        set_cell_bg(cell, bg)
        set_cell_borders(cell, color='FFFFFF', sz=6)
        p   = cell.paragraphs[0]
        run = p.add_run(h)
        run.font.bold      = True
        run.font.size      = Pt(10)
        run.font.color.rgb = fg
        p.alignment        = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if widths:
        for i,w in enumerate(widths):
            row.cells[i].width = Inches(w)

def data_row(table, row_idx, values, alt=False, bold_first=False):
    row = table.rows[row_idx]
    bg  = SAP_LIGHT if alt else WHITE
    for i, v in enumerate(values):
        cell = row.cells[i]
        cell.text = ''
        set_cell_bg(cell, bg)
        set_cell_borders(cell, color='C8D8E8', sz=4)
        p   = cell.paragraphs[0]
        run = p.add_run(str(v))
        run.font.size      = Pt(9.5)
        run.font.color.rgb = GRAY_DARK
        run.font.bold      = (bold_first and i == 0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def add_info_box(doc, title, text, bg=SAP_LIGHT, title_color=SAP_DARK):
    tbl = doc.add_table(rows=2, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    # title row
    tc = tbl.rows[0].cells[0]
    set_cell_bg(tc, title_color)
    set_cell_borders(tc, color='003369', sz=6)
    p   = tc.paragraphs[0]
    run = p.add_run(f'  {title}')
    run.font.bold      = True
    run.font.size      = Pt(10)
    run.font.color.rgb = WHITE
    # body row
    bc = tbl.rows[1].cells[0]
    set_cell_bg(bc, bg)
    set_cell_borders(bc, color='90B8D0', sz=4)
    for line in text.split('\n'):
        bp   = bc.add_paragraph()
        brun = bp.add_run(f'  {line}')
        brun.font.size      = Pt(9.5)
        brun.font.color.rgb = GRAY_DARK
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after  = Pt(1)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ═══════════════════════════════════════════════════════════════════════════════
# FLOWCHART — matplotlib
# ═══════════════════════════════════════════════════════════════════════════════
def make_flowchart(path):
    fig, ax = plt.subplots(figsize=(18, 13))
    ax.set_xlim(0, 18); ax.set_ylim(0, 13)
    ax.axis('off')
    fig.patch.set_facecolor('#F7FAFD')

    C_SRC  = '#00619A'   # blue  – fuentes
    C_DIM  = '#003369'   # dark  – dimensiones
    C_ING  = '#107E3E'   # green – ingresos
    C_GAS  = '#E86A19'   # orange– gastos
    C_EFF  = '#8B1A8B'   # purple– EEFF
    C_OUT  = '#F0AB00'   # gold  – salidas
    C_GRAY = '#607080'   # gray  – fondo bandas
    WHITE  = '#FFFFFF'
    LTBLUE = '#E8F4FD'
    LTYELL = '#FFF8E1'

    def box(ax, x, y, w, h, label, sub='', color='#00619A', alpha=0.93, fontsize=11, subsize=9, text_color='white', radius=0.35):
        rect = FancyBboxPatch((x-w/2, y-h/2), w, h,
                              boxstyle=f'round,pad=0.05,rounding_size={radius}',
                              linewidth=1.5, edgecolor='white',
                              facecolor=color, alpha=alpha, zorder=3)
        ax.add_patch(rect)
        ty = y + (0.15 if sub else 0)
        ax.text(x, ty, label, ha='center', va='center', fontsize=fontsize,
                fontweight='bold', color=text_color, zorder=4,
                wrap=True, multialignment='center')
        if sub:
            ax.text(x, y-0.28, sub, ha='center', va='center', fontsize=subsize,
                    color=text_color, alpha=0.88, zorder=4, multialignment='center')

    def arrow(ax, x1,y1, x2,y2, color='#607080'):
        ax.annotate('', xy=(x2,y2), xytext=(x1,y1),
                    arrowprops=dict(arrowstyle='->', color=color,
                                   lw=2.0, connectionstyle='arc3,rad=0.0'), zorder=2)

    def band(ax, y, h, label, color, alpha=0.10):
        ax.add_patch(plt.Rectangle((0.3, y-h/2), 17.4, h,
                                   facecolor=color, alpha=alpha, zorder=1, linewidth=0))
        ax.text(0.55, y, label, ha='left', va='center', fontsize=8,
                color=color, fontweight='bold', rotation=90, alpha=0.7, zorder=2)

    # ── BANDAS DE FONDO ──────────────────────────────────────────────────────
    band(ax, 11.5, 1.8,  'FUENTES DE DATOS',       C_SRC)
    band(ax, 9.2,  1.5,  'DIMENSIONES COMPARTIDAS', C_DIM)
    band(ax, 7.0,  1.6,  'CAPA DE MODELOS',         C_ING)
    band(ax, 4.8,  1.8,  'ESTADOS FINANCIEROS',     C_EFF)
    band(ax, 2.5,  1.8,  'CONSUMO / SALIDAS',       C_OUT)

    # ── CAPA 1: FUENTES ───────────────────────────────────────────────────────
    box(ax, 3.5,  11.5, 2.8, 1.0, 'SAP S/4HANA',    'Datos Reales (ACDOCA)', C_SRC)
    box(ax, 9.0,  11.5, 2.8, 1.0, 'Excel / Plantillas', 'Presupuesto por área', '#2E7D9E')
    box(ax, 14.5, 11.5, 2.8, 1.0, 'Input Manual SAC', 'Ajustes y correcciones', '#1A5F7A')

    # ── CAPA 2: DIMENSIONES ───────────────────────────────────────────────────
    dims = ['SOCIEDAD_CL\n(Organization)', 'CUENTA\n(Account)', 'CEBE_CL\n(Generic)',
            'CECO_CL\n(Generic)', 'AUDITORIA\n(Generic)', 'MONEDA\n(Generic)']
    xs   = [2.0, 5.0, 7.8, 10.6, 13.4, 16.2]
    for i,(d,x) in enumerate(zip(dims, xs)):
        box(ax, x, 9.2, 2.4, 0.85, d, '', C_DIM, fontsize=9)

    # ── CAPA 3: MODELOS ───────────────────────────────────────────────────────
    box(ax, 4.0,  7.0, 3.4, 1.2, 'MODELO INGRESOS',
        'Ciudad Limpia / CGS / RH\nOrdinarios + P×Q', C_ING, fontsize=12)
    box(ax, 9.0,  7.0, 3.4, 1.2, 'MODELO GASTOS/COSTOS',
        'Costos Directos · Flota\nGastos Transversales', C_GAS, fontsize=12)
    box(ax, 14.0, 7.0, 3.4, 1.2, 'PARKING LOT',
        'Modelo Nómina/Personal\n(pendiente de diseño)', C_GRAY, fontsize=11)

    # ── CAPA 4: EEFF ──────────────────────────────────────────────────────────
    box(ax, 9.0,  4.8, 7.0, 1.3, 'MODELO EEFF — ESTADOS FINANCIEROS',
        'P&G  ·  Balance General  ·  Flujo de Caja', C_EFF, fontsize=13)

    # ── CAPA 5: SALIDAS ───────────────────────────────────────────────────────
    box(ax, 3.5,  2.5, 3.0, 1.2, 'Dashboard\nEjecutivo',
        'PDF · PowerPoint', C_OUT, text_color='#1A1A1A', fontsize=11)
    box(ax, 7.5,  2.5, 3.0, 1.2, 'Reporte\nOperativo',
        'Toneladas · km · usuarios', C_OUT, text_color='#1A1A1A', fontsize=11)
    box(ax, 11.5, 2.5, 3.0, 1.2, 'Ejecución\npor Área',
        'Presupuesto vs. Real', C_OUT, text_color='#1A1A1A', fontsize=11)
    box(ax, 15.2, 2.5, 3.0, 1.2, 'Reportes\nExternos *',
        '* Parking Lot', '#B8860B', text_color='white', fontsize=11)

    # ── FLECHAS FUENTES → DIMENSIONES ────────────────────────────────────────
    for fx in [3.5, 9.0, 14.5]:
        arrow(ax, fx, 11.0, 9.0, 9.65, '#90B8D0')

    # ── FLECHAS DIMENSIONES → MODELOS ────────────────────────────────────────
    arrow(ax, 4.0,  8.77, 4.0,  7.6,  C_ING)
    arrow(ax, 9.0,  8.77, 9.0,  7.6,  C_GAS)
    arrow(ax, 14.0, 8.77, 14.0, 7.6,  C_GRAY)

    # ── FLECHAS MODELOS → EEFF ────────────────────────────────────────────────
    arrow(ax, 4.0,  6.4,  6.5,  5.45, C_EFF)
    arrow(ax, 9.0,  6.4,  9.0,  5.45, C_EFF)

    # ── FLECHAS EEFF → SALIDAS ────────────────────────────────────────────────
    for ox in [3.5, 7.5, 11.5, 15.2]:
        arrow(ax, 9.0, 4.15, ox, 3.1, C_OUT)

    # ── TÍTULO ────────────────────────────────────────────────────────────────
    ax.text(9.0, 12.6, 'Arquitectura de Modelos SAP Analytics Cloud — Ciudad Limpia',
            ha='center', va='center', fontsize=15, fontweight='bold',
            color='#003369', zorder=5)
    ax.text(9.0, 12.2, 'Proyecto SAP BUILD  |  Sesión de Diseño N°1  |  Mayo 2026',
            ha='center', va='center', fontsize=10, color='#607080', zorder=5)

    # ── LEYENDA ───────────────────────────────────────────────────────────────
    legend_items = [
        mpatches.Patch(color=C_SRC,  label='Fuentes de Datos'),
        mpatches.Patch(color=C_DIM,  label='Dimensiones Compartidas'),
        mpatches.Patch(color=C_ING,  label='Modelo Ingresos'),
        mpatches.Patch(color=C_GAS,  label='Modelo Gastos/Costos'),
        mpatches.Patch(color=C_EFF,  label='Modelo EEFF'),
        mpatches.Patch(color=C_OUT,  label='Reportes / Dashboards'),
        mpatches.Patch(color=C_GRAY, label='Parking Lot'),
    ]
    ax.legend(handles=legend_items, loc='lower center', ncol=7,
              fontsize=8.5, framealpha=0.9, bbox_to_anchor=(0.5, -0.01),
              frameon=True, edgecolor='#CCCCCC')

    plt.tight_layout(pad=0.3)
    plt.savefig(path, dpi=180, bbox_inches='tight', facecolor='#F7FAFD')
    plt.close()
    print(f'Flowchart guardado: {path}')

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENTO WORD PROFESIONAL
# ═══════════════════════════════════════════════════════════════════════════════
def make_document(flowchart_path, out_path):
    doc = Document()

    # ── MÁRGENES ──────────────────────────────────────────────────────────────
    for sec in doc.sections:
        sec.top_margin    = Cm(2.2)
        sec.bottom_margin = Cm(2.2)
        sec.left_margin   = Cm(2.8)
        sec.right_margin  = Cm(2.2)
        sec.page_width    = Cm(21.59)
        sec.page_height   = Cm(27.94)

    # ── ESTILOS BASE ──────────────────────────────────────────────────────────
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)

    # ══════════════════════════════════════════════════════════════════════════
    # PORTADA
    # ══════════════════════════════════════════════════════════════════════════
    # Banda superior azul
    cover_top = doc.add_table(rows=1, cols=1)
    cover_top.alignment = WD_TABLE_ALIGNMENT.CENTER
    ct = cover_top.rows[0].cells[0]
    set_cell_bg(ct, SAP_DARK)
    p = ct.paragraphs[0]
    run = p.add_run('  SAP Analytics Cloud  |  Documento de Diseño de Solución')
    run.font.size = Pt(11); run.font.color.rgb = WHITE; run.font.bold = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(6)

    doc.add_paragraph().paragraph_format.space_after = Pt(40)

    # Título principal
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Modelos de Planificación\ny Reporte Financiero')
    run.font.size = Pt(32); run.font.bold = True; run.font.color.rgb = SAP_DARK
    p.paragraph_format.space_after = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Ciudad Limpia  ·  Grupo Fanalca')
    run.font.size = Pt(18); run.font.color.rgb = SAP_BLUE
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Proyecto SAP BUILD')
    run.font.size = Pt(14); run.font.italic = True; run.font.color.rgb = GRAY_MED
    p.paragraph_format.space_after = Pt(60)

    # Tabla de metadatos de portada
    meta = doc.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ('Versión',    '2.0'),
        ('Fecha',      'Junio 2026'),
        ('Estado',     'Borrador para Validación'),
        ('Clasificación', 'Confidencial'),
        ('Sesión base', 'Diseño N°1 — 22 de mayo de 2026'),
    ]
    for i, (k, v) in enumerate(meta_data):
        ck = meta.rows[i].cells[0]
        cv = meta.rows[i].cells[1]
        set_cell_bg(ck, SAP_BLUE)
        set_cell_bg(cv, SAP_LIGHT if i%2==0 else WHITE)
        set_cell_borders(ck, color='FFFFFF', sz=4)
        set_cell_borders(cv, color='90B8D0', sz=4)
        pk = ck.paragraphs[0]
        rk = pk.add_run(f'  {k}')
        rk.font.bold = True; rk.font.size = Pt(10); rk.font.color.rgb = WHITE
        pv = cv.paragraphs[0]
        rv = pv.add_run(f'  {v}')
        rv.font.size = Pt(10); rv.font.color.rgb = GRAY_DARK
    for i in range(5):
        meta.rows[i].cells[0].width = Inches(2.0)
        meta.rows[i].cells[1].width = Inches(3.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(50)

    # Banda inferior portada
    cover_bot = doc.add_table(rows=1, cols=1)
    cb = cover_bot.rows[0].cells[0]
    set_cell_bg(cb, SAP_DARK)
    p = cb.paragraphs[0]
    run = p.add_run('  Elaborado por: Equipo SAP Analytics Cloud  |  Confidencial — Solo para uso interno')
    run.font.size = Pt(9); run.font.color.rgb = RGBColor(0xCC,0xDD,0xEE)
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(5)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # SECCIÓN 1 — PROPÓSITO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '1. Propósito del Documento', 1)
    add_body(doc, (
        'Este documento describe el diseño técnico y funcional de los modelos de datos implementados '
        'en SAP Analytics Cloud (SAC) para Ciudad Limpia y las sociedades del Grupo Fanalca. '
        'Incorpora los acuerdos y decisiones tomadas durante la Sesión de Diseño N°1 (22-May-2026), '
        'en la que Planeación Financiera expuso el proceso presupuestal actual.'
    ), space_after=6)

    add_info_box(doc, 'Alcance del documento',
        'Estructura de los 3 modelos SAC: Ingresos, Gastos/Costos y EEFF\n'
        'Decisiones de diseño y métodos de carga acordados en sesión\n'
        'Temas en Parking Lot que requieren sesiones adicionales\n'
        'Flujos de integración: SAP S/4HANA, Excel y entrada manual\n'
        'Criterios de seguridad, acceso y gobierno de datos')

    # ══════════════════════════════════════════════════════════════════════════
    # SECCIÓN 2 — CÓMO USAR
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '2. Cómo Usar el Documento', 1)
    tbl = doc.add_table(rows=5, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    header_row(tbl, ['Perfil', 'Uso recomendado'], widths=[2.0, 4.5])
    perfiles = [
        ('Consultores SAC / Equipo técnico', 'Referencia para construcción y configuración de modelos.'),
        ('Líderes funcionales del cliente',  'Validar que el diseño refleja correctamente los requisitos.'),
        ('Equipo de TI',                     'Comprender integración S/4HANA y flujos de datos.'),
        ('Gerentes de proyecto',             'Seguimiento al alcance y completitud del diseño.'),
    ]
    for i, (p1, p2) in enumerate(perfiles):
        data_row(tbl, i+1, [p1, p2], alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # SECCIÓN 3 — ESCENARIO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '3. Escenario / Proceso de Negocio', 1)
    add_body(doc, (
        'Ciudad Limpia es una empresa de servicios públicos de aseo urbano que opera en Bogotá, Neiva, '
        'Huila y otros municipios. Sus empresas relacionadas CGS y RH gestionan residuos peligrosos y '
        'aprovechamiento. El proceso presupuestal actual está íntegramente basado en macros Excel, '
        'con consolidación manual en Planeación Financiera.'
    ), space_after=6)

    add_heading(doc, '3.1 Objetivos Empresariales y Beneficios Esperados', 2)
    tbl2 = doc.add_table(rows=7, cols=3)
    header_row(tbl2, ['Código', 'Objetivo', 'Beneficio Esperado'], widths=[0.7, 2.8, 3.0])
    obj_data = [
        ('OE-01', 'Centralizar información financiera en SAC',          'Fin de la fragmentación en decenas de Excel'),
        ('OE-02', 'Automatizar consolidación del presupuesto',           'Proceso de semanas → horas/minutos'),
        ('OE-03', 'Habilitar forecast ágil',                             'Actualización rápida ante cambios de tarifa o dotación'),
        ('OE-04', 'Visibilidad en tiempo real para gerentes',            'Cada área consulta SAC sin intermediar a Planeación'),
        ('OE-05', 'Auditoría y trazabilidad completa',                   'Origen claro de cada dato: real, presupuesto, manual'),
        ('OE-06', 'Generar EEFF con mínima intervención manual',         'Reducción de errores en cierre mensual'),
    ]
    for i, row in enumerate(obj_data):
        data_row(tbl2, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading(doc, '3.2 Requisitos Empresariales', 2)
    req_data = [
        ('RE-01', 'Ingresos ordinarios Ciudad Limpia',
         'Cargar valor total por componente desde resumen del jefe de tarifas. Sin replicar cálculo tarifario complejo.'),
        ('RE-02', 'Ingresos CGS y RH (P×Q)',
         'Cálculo tarifa × toneladas/kilos por cliente y tipo de residuo. Se hace seguimiento de cantidades.'),
        ('RE-03', 'Plantilla genérica de gastos',
         'Una plantilla Excel estándar para todas las ~17 áreas. Planeación Financiera consolida y carga.'),
        ('RE-04', 'Costos de flota',
         'Un CECO por vehículo. Rubros: repuestos, llantas, combustibles, lubricantes.'),
        ('RE-05', 'Mano de obra (Parking Lot)',
         '~80% del costo. Sesión dedicada de diseño. Propuesta basada en modelo referencial.'),
        ('RE-06', 'Estados Financieros',
         'P&G fluye automáticamente. Balance y Flujo: cuentas adicionales por carga manual.'),
        ('RE-07', 'Multimoneda',           'Soporte de múltiples monedas por sociedad.'),
        ('RE-08', 'Gestión de versiones',
         'Actual / Presupuesto / Forecast. Regla: siempre el dato final autorizado; sin acumulación.'),
        ('RE-09', 'Auditoría',             'Dimensión AUDITORIA para identificar origen de cada dato.'),
        ('RE-10', 'Opciones de carga',
         'Reemplazar / Acumular / Borrar y recargar. Estándar acordado: Reemplazar con dato definitivo.'),
        ('RE-11', 'Dashboards y KPIs',
         'Filtros por sociedad/área/período. Exportable PDF/PPT. KPIs: toneladas, km, usuarios, árboles.'),
    ]
    tbl3 = doc.add_table(rows=len(req_data)+1, cols=3)
    header_row(tbl3, ['Código', 'Requisito', 'Descripción'], widths=[0.7, 2.0, 3.8])
    for i, row in enumerate(req_data):
        data_row(tbl3, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # SECCIÓN 4 — ESTRUCTURA ORGANIZATIVA
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '4. Estructura Organizativa Relevante', 1)
    add_heading(doc, '4.1 Sociedades', 2)
    soc_data = [
        ('Ciudad Limpia Bogotá',  '~1,900 empleados', 'Principal. Ingresos ordinarios complejos.'),
        ('Ciudad Limpia Neiva',   '~500 empleados',   'Particularidades tarifarias propias.'),
        ('Ciudad Limpia Huila',   '~37 empleados',    'Operación pequeña.'),
        ('CGS',                   '~182 empleados',   'Residuos peligrosos + aseo en municipios. P×Q.'),
        ('RH',                    'En proceso',       'Residuos peligrosos. Ingresando a nómina web.'),
    ]
    tbl4 = doc.add_table(rows=6, cols=3)
    header_row(tbl4, ['Sociedad', 'Empleados', 'Descripción'], widths=[2.0, 1.3, 3.2])
    for i, row in enumerate(soc_data):
        data_row(tbl4, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading(doc, '4.2 Componentes de Servicio (EBEs/ECOs)', 2)
    comp_data = [
        ('Recolección',           'Toneladas recolectadas',          'CEBE específico en S/4HANA'),
        ('Barrido de vías',       'Kilómetros barridos',             'CEBE específico'),
        ('Corte de césped',       'Metros cuadrados cortados',       'CEBE específico'),
        ('Poda de árboles',       'Cantidad de árboles podados',     'CEBE específico'),
        ('Aprovechamiento',       'Toneladas aprovechadas',          'CGS / RH'),
        ('Comercialización',      'Valor facturado',                 'Otros ingresos'),
    ]
    tbl5 = doc.add_table(rows=7, cols=3)
    header_row(tbl5, ['Componente', 'Métrica de seguimiento', 'Observación'], widths=[1.8, 2.2, 2.5])
    for i, row in enumerate(comp_data):
        data_row(tbl5, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # SECCIÓN 5 — DISEÑO Y CONFIGURACIÓN
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '5. Diseño y Configuración de Soluciones', 1)

    # 5.1 Flujograma
    add_heading(doc, '5.1 Diagrama de Arquitectura de Modelos', 2)
    add_body(doc, (
        'El siguiente diagrama muestra la arquitectura completa de la solución: las fuentes de datos, '
        'la capa de dimensiones compartidas, los tres modelos SAC y las salidas de consumo.'
    ), space_after=6)
    doc.add_picture(flowchart_path, width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap = doc.add_paragraph('Figura 1 — Arquitectura de Modelos SAP Analytics Cloud | Ciudad Limpia')
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.runs[0].font.size = Pt(9); p_cap.runs[0].font.italic = True
    p_cap.runs[0].font.color.rgb = GRAY_MED
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 5.2 MODELO INGRESOS
    add_heading(doc, '5.2 Modelo: Ingresos Ciudad Limpia', 2)
    add_info_box(doc, 'Propósito',
        'Captura, planificación y análisis de ingresos operativos de todas las sociedades.\n'
        'Soporta carga de ingresos ordinarios (valor total) y cálculo P×Q para CGS y RH.',
        bg=RGBColor(0xE8,0xF8,0xED), title_color=SAP_GREEN)

    add_heading(doc, 'Medidas', 3)
    tbl_m1 = doc.add_table(rows=4, cols=3)
    header_row(tbl_m1, ['Medida', 'Tipo de Dato', 'Descripción'], widths=[1.8, 1.5, 3.2])
    meds1 = [
        ('Importe',   'Decimal (Moneda $)', 'Valor monetario del ingreso'),
        ('Cantidad',  'Decimal',            'Toneladas / kilos (aplica CGS/RH para P×Q)'),
        ('Tarifa',    'Decimal (Moneda $)', 'Precio unitario por tipo de residuo (CGS/RH)'),
    ]
    for i, row in enumerate(meds1):
        data_row(tbl_m1, i+1, row, alt=(i%2==0), bold_first=True)

    add_heading(doc, 'Dimensiones', 3)
    tbl_d1 = doc.add_table(rows=9, cols=5)
    header_row(tbl_d1, ['Dimensión','Tipo SAC','Jerarquía','Pública','Notas'], widths=[1.3,1.2,0.7,0.6,2.7])
    dims1 = [
        ('Version',     'Version',      '1',  'No',  'Actual / Presupuesto / Forecast'),
        ('CUENTA',      'Account',      '0',  'Sí',  'Ingresos por componente. Jerarquía EEFF y gestión'),
        ('Date',        'Date',         '-',  'No',  'Granularidad mensual. Ejercicio fiscal'),
        ('CEBE_CL',     'Generic',      '1',  'Sí',  'Componente de servicio (Recolección, Barrido…)'),
        ('AUDITORIA',   'Generic',      '1',  'Sí',  'Origen: REAL_S4 / PRESUPUESTO_TARIFAS / PXQ'),
        ('MONEDA',      'Generic',      '1',  'Sí',  'Moneda de la transacción'),
        ('CLIENTE',     'Generic',      '1',  'No',  'Solo CGS/RH — cliente y tipo residuo'),
        ('SOCIEDAD_CL', 'Organization', '1',  'Sí',  'Data Access Control por entidad legal'),
    ]
    for i, row in enumerate(dims1):
        data_row(tbl_d1, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading(doc, 'Lógica de carga por sociedad', 3)
    tbl_lg1 = doc.add_table(rows=3, cols=3)
    header_row(tbl_lg1, ['Sociedad', 'Método', 'Responsable de datos'], widths=[2.0,2.5,2.0])
    data_row(tbl_lg1, 1, ['CL Bogotá / Neiva / Huila','Valor total por componente desde Excel del jefe de tarifas','Jefe Nac. Tarifas → Planeación Fin.'], alt=False, bold_first=True)
    data_row(tbl_lg1, 2, ['CGS / RH','Cálculo P×Q: tarifa × toneladas/kilos por cliente y residuo','Área CGS/RH → Planeación Fin.'], alt=True, bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.3 MODELO GASTOS
    add_heading(doc, '5.3 Modelo: Gastos/Costos Ciudad', 2)
    add_info_box(doc, 'Propósito',
        'Planificación, control y análisis de costos operativos y gastos de todas las sociedades.\n'
        'Permite seguimiento por centro de costo, componente de servicio, cuenta y sociedad.',
        bg=RGBColor(0xFF,0xF0,0xE0), title_color=ORANGE)

    add_heading(doc, 'Dimensiones', 3)
    tbl_d2 = doc.add_table(rows=9, cols=5)
    header_row(tbl_d2, ['Dimensión','Tipo SAC','Jerarquía','Pública','Notas'], widths=[1.3,1.2,0.7,0.6,2.7])
    dims2 = [
        ('Version',     'Version',      '1',  'No',  'Actual / Presupuesto / Forecast'),
        ('CUENTA',      'Account',      '0',  'Sí',  'Cuentas de costos y gastos'),
        ('Date',        'Date',         '-',  'No',  'Mensual'),
        ('CECO_CL',     'Generic',      '1',  'Sí',  'Centro de costo (área / vehículo de flota)'),
        ('CEBE_CL',     'Generic',      '1',  'Sí',  'Componente de servicio / línea de negocio'),
        ('AUDITORIA',   'Generic',      '1',  'Sí',  'Origen: REAL_S4 / PRESUPUESTO_EXCEL / MANUAL'),
        ('MONEDA',      'Generic',      '1',  'Sí',  'Moneda'),
        ('SOCIEDAD_CL', 'Organization', '1',  'Sí',  'Data Access Control por entidad legal'),
    ]
    for i, row in enumerate(dims2):
        data_row(tbl_d2, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading(doc, 'Categorías de gasto y método de carga', 3)
    cat_data = [
        ('Mano de Obra',       '~80% del costo total',          'PARKING LOT — sesión dedicada',        '1,900+500+37+182 empl.'),
        ('Mantenimiento Flota','Repuestos, llantas, combustible','Valor por CECO (1 CECO/vehículo)','Area Mantenimiento → Excel'),
        ('Depreciaciones',     'Activos fijos',                  'Plantilla genérica',                   'Contabilidad'),
        ('TIC y tecnología',   'Equipos, licencias, mant.',      'Plantilla genérica',                   'Área TIC'),
        ('Servicios públicos', 'Agua, energía, telefonía',       'Plantilla genérica',                   'Servicios Generales'),
        ('Arrendamientos',     'Infraestructura y equipos',      'Plantilla genérica',                   'Contabilidad'),
        ('Dotación personal',  'EPP, uniformes, reposiciones',   'Plantilla genérica',                   'Área centralizada'),
        ('Seguros',            'Pólizas corporativas',           'Plantilla genérica',                   'Administración'),
        ('Gastos generales',   'Honorarios, viajes, diversos',   'Plantilla genérica',                   'Cada gerencia'),
    ]
    tbl_cat = doc.add_table(rows=len(cat_data)+1, cols=4)
    header_row(tbl_cat, ['Categoría','Descripción','Método en SAC','Responsable'], widths=[1.5,1.8,1.8,1.4])
    for i, row in enumerate(cat_data):
        data_row(tbl_cat, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_info_box(doc, 'Regla de oro acordada en Sesión N°1 (lección aprendida de Fanalca / BPC)',
        'Cada responsable entrega el DATO FINAL AUTORIZADO — no versiones parciales ni acumulaciones.\n'
        'El sistema usa la opción de carga "Reemplazar": el último dato cargado es el definitivo.\n'
        'Modificaciones puntuales se hacen directamente en la plantilla SAC, sin recargar el archivo completo.',
        bg=RGBColor(0xFF,0xFF,0xE8), title_color=RGBColor(0xB8,0x86,0x0B))

    # 5.4 MODELO EEFF
    add_heading(doc, '5.4 Modelo: EEFF Ciudad Limpia', 2)
    add_info_box(doc, 'Propósito',
        'Consolidar P&G, Balance General y Flujo de Caja. Alimenta reportes internos (gerentes, junta)\n'
        'y externos (Superintendencia, CRA, Unidad de Servicios Públicos).',
        bg=RGBColor(0xF3,0xE8,0xF8), title_color=RGBColor(0x8B,0x1A,0x8B))

    add_heading(doc, 'Lógica de construcción del EEFF', 3)
    eeff_data = [
        ('Estado de Resultados (P&G)', 'Fluye automáticamente desde Modelos de Ingresos y Gastos/Costos via lógica de cálculo SAC.'),
        ('Balance — cuentas de P&G',   'Las cuentas de resultados impactan automáticamente el balance (ej: utilidad del período).'),
        ('Balance — cuentas propias',  'Activos fijos, cartera, deuda, capital: carga manual o integración directa S/4HANA.'),
        ('Flujo de Caja',              'Combinación: parte fluye del P&G, parte del balance y parte por carga manual (impuestos, CAPEX, inversiones).'),
        ('Depreciaciones acumuladas',  'Fluyen desde el rubro de depreciaciones del modelo de gastos.'),
        ('CAPEX / Inversiones flota',  'Carga manual. Proceso muy particular — no fluye directamente del P&G.'),
    ]
    tbl_eeff = doc.add_table(rows=len(eeff_data)+1, cols=2)
    header_row(tbl_eeff, ['Componente EEFF', 'Lógica de alimentación'], widths=[2.2, 4.3])
    for i, row in enumerate(eeff_data):
        data_row(tbl_eeff, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.5 DIMENSIONES COMPARTIDAS
    add_heading(doc, '5.5 Dimensiones Compartidas y Maestros de Datos', 2)
    dim_shared = [
        ('CUENTA',      'Account',      'Sí', 'S/4HANA SKA1/SKB1', 'Plan de cuentas completo. Jerarquías: LR3 (Fanalca) y gestión interna.'),
        ('CEBE_CL',     'Generic',      'Sí', 'S/4HANA CEPC',      'Centros de beneficio / componentes de servicio. Presente en los 3 modelos.'),
        ('CECO_CL',     'Generic',      'Sí', 'S/4HANA CSKS',      'Centros de costo. Incluye un CECO por vehículo de flota. Solo en Gastos.'),
        ('AUDITORIA',   'Generic',      'Sí', 'Maestra SAC',        'Miembros: REAL_S4, PRESUPUESTO_TARIFAS, PRESUPUESTO_EXCEL, PXQ, MANUAL, AJUSTE.'),
        ('MONEDA',      'Generic',      'Sí', 'Maestra SAC',        'Monedas operativas. Evaluar Currency Translation para consolidados.'),
        ('SOCIEDAD_CL', 'Organization', 'Sí', 'S/4HANA T001',      'Habilita Data Access Control. Controla visibilidad por entidad legal.'),
        ('VERSION',     'Version',      'No', 'SAC (privada)',      'Versiones: ACTUAL_YYYY / PRESUPUESTO_YYYY / FORECAST_QX_YYYY.'),
        ('DATE',        'Date',         'No', 'SAC (privada)',      'Granularidad mensual. Rango: 2 años históricos + año en curso + 1 futuro.'),
    ]
    tbl_dims = doc.add_table(rows=len(dim_shared)+1, cols=5)
    header_row(tbl_dims, ['Dimensión','Tipo SAC','Pública','Fuente Maestra','Descripción'], widths=[1.2,1.1,0.6,1.3,2.3])
    for i, row in enumerate(dim_shared):
        data_row(tbl_dims, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.6 INTEGRACIÓN
    add_heading(doc, '5.6 Diseño de Integración y Fuentes de Datos', 2)
    int_data = [
        ('SAP S/4HANA\n(Datos Reales)',   'CDS Views / Conexión nativa SAC-S4',  'Diaria — fuera de horario laboral', 'ACDOCA, SKA1, CEPC, CSKS, T001'),
        ('Excel / Plantillas\n(Presupuesto)', 'Carga de archivo en SAC Data Mgmt', 'Anual + actualizaciones forecast',  'Plantillas estándar por categoría'),
        ('Input Manual SAC\n(Ajustes)',    'Entrada directa en plantilla SAC',    'Bajo demanda',                      'Solo Planeación Financiera'),
    ]
    tbl_int = doc.add_table(rows=4, cols=4)
    header_row(tbl_int, ['Fuente','Mecanismo','Frecuencia','Objetos clave'], widths=[1.6,1.8,1.5,1.6])
    for i, row in enumerate(int_data):
        data_row(tbl_int, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.7 PLANTILLAS
    add_heading(doc, '5.7 Diseño de Plantillas de Carga', 2)
    plant_data = [
        ('Plantilla 1', 'Ingresos Ordinarios CL',
         'Sociedad | CEBE (Componente) | Cuenta | Ene-Dic',
         'Jefe de Tarifas → Planeación Fin.'),
        ('Plantilla 2', 'Ingresos P×Q — CGS/RH',
         'Sociedad | CEBE | Cliente | Tipo Residuo | Cuenta | Tarifa | Cant. | Ene-Dic',
         'Áreas CGS/RH → Planeación Fin.'),
        ('Plantilla 3', 'Gastos y Costos (Genérica)',
         'Sociedad | CECO | CEBE | Cuenta | Ene-Dic',
         'Todas las áreas → Planeación Fin. consolida'),
    ]
    tbl_pl = doc.add_table(rows=4, cols=4)
    header_row(tbl_pl, ['Plantilla','Nombre','Estructura de columnas','Responsable'], widths=[0.8,1.6,2.8,1.3])
    for i, row in enumerate(plant_data):
        data_row(tbl_pl, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.8 SEGURIDAD
    add_heading(doc, '5.8 Consideraciones de Seguridad y Acceso', 2)
    add_info_box(doc, 'Política de licenciamiento acordada en sesión',
        'Las licencias SAC se asignan ÚNICAMENTE al equipo de Planeación Financiera.\n'
        'Las áreas operativas NO tendrán acceso a SAC para carga directa.\n'
        'Lección aprendida Fanalca/BPC: dar licencias a usuarios que acceden 1 vez/año\n'
        'genera más costo de soporte y habilitación que el beneficio obtenido.',
        bg=RGBColor(0xFD,0xF0,0xF0), title_color=RGBColor(0xA0,0x00,0x00))

    rol_data = [
        ('Administrador SAC',    'Total',              'Configura modelos, dimensiones, conexiones, usuarios.'),
        ('Planeación Financiera','Lectura / Escritura', 'Versiones plan. Carga de archivos Excel. Ajustes manuales.'),
        ('Gerente / Ejecutivo',  'Solo Lectura',        'Dashboards y reportes. Filtrado por sociedad asignada.'),
        ('Auditor',              'Solo Lectura',        'Acceso histórico a todas las versiones y sociedades.'),
    ]
    tbl_rol = doc.add_table(rows=5, cols=3)
    header_row(tbl_rol, ['Rol','Nivel de acceso','Descripción'], widths=[1.8,1.5,3.2])
    for i, row in enumerate(rol_data):
        data_row(tbl_rol, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.9 RESUMEN COMPARATIVO
    add_heading(doc, '5.9 Resumen Comparativo de Modelos', 2)
    comp_tbl = doc.add_table(rows=13, cols=4)
    header_row(comp_tbl, ['Elemento','Modelo Gastos/Costos','Modelo EEFF','Modelo Ingresos'], widths=[1.7,1.8,1.7,1.8])
    comp_rows = [
        ('Importe (Medida)',    '✔', '✔', '✔'),
        ('Version',            '✔', '✔', '✔'),
        ('CUENTA',             '✔', '✔', '✔'),
        ('Date',               '✔', '✔', '✔'),
        ('CEBE_CL',            '✔', '✔', '✔'),
        ('AUDITORIA',          '✔', '✔', '✔'),
        ('MONEDA',             '✔', '✔', '✔'),
        ('SOCIEDAD_CL',        '✔', '✔', '✔'),
        ('CECO_CL',            '✔', '—', '—'),
        ('CLIENTE',            '—', '—', '✔ (CGS/RH)'),
        ('Cantidad (P×Q)',      '—', '—', '✔ (CGS/RH)'),
        ('Tarifa (P×Q)',        '—', '—', '✔ (CGS/RH)'),
    ]
    for i, row in enumerate(comp_rows):
        data_row(comp_tbl, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.10 PARKING LOT
    add_heading(doc, '5.10 Parking Lot — Temas Pendientes de Definición', 2)
    park_data = [
        ('PL-01','Modelo de Gasto de Personal / Nómina',
         'Sesión dedicada de diseño. Consultor presenta modelo referencial.',
         'Alta','Próxima semana'),
        ('PL-02','Reportes para Entes Externos (Supersociedades, CRA, SUI, Unidad)',
         'Confirmar si está en alcance del proyecto.',
         'Media','Por definir'),
        ('PL-03','Distribuciones Intercompañía (centros de costo entre sociedades)',
         'Sesión con equipo DataSphere/COPA. Pendiente con Alejandra.',
         'Media','Por definir'),
        ('PL-04','Reporte Balance por APS (municipio vigilado vs. no vigilado)',
         'Evaluar si CEBE_CL actual permite separación por municipio.',
         'Media','Por definir'),
        ('PL-05','Cálculo P×Q detallado para CL Bogotá (ordinarios)',
         'Validar con jefe de tarifas si es viable simplificar a P×Q en el futuro.',
         'Baja','Largo plazo'),
    ]
    tbl_park = doc.add_table(rows=len(park_data)+1, cols=5)
    header_row(tbl_park, ['Código','Tema','Próximo paso','Prioridad','Fecha est.'], widths=[0.6,2.0,2.2,0.7,0.9])
    for i, row in enumerate(park_data):
        data_row(tbl_park, i+1, row, alt=(i%2==0), bold_first=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # SECCIÓN 6 — PRÓXIMOS PASOS
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '6. Próximos Pasos', 1)
    steps = [
        ('1', 'Validar este documento con Laidys y equipo de Planeación Financiera.',       'Equipo cliente + Consultor'),
        ('2', 'Sesión dedicada: diseño Modelo de Gasto de Personal (PL-01).',               'Consultor + TH Ciudad Limpia'),
        ('3', 'Confirmar con Carlos Otálora si ingresos CL pueden simplificarse a P×Q.',   'Planeación Fin. + Tarifas'),
        ('4', 'Diseñar y validar las 3 plantillas Excel estándar.',                         'Consultor + Planeación Fin.'),
        ('5', 'Confirmar plan de cuentas completo y jerarquías EEFF (LR3 y gerencial).',   'Contabilidad + Consultor'),
        ('6', 'Confirmar maestros CECO_CL incluyendo codificación de vehículos.',          'Mantenimiento + TI'),
        ('7', 'Confirmar alcance de reportes a entes externos (PL-02).',                    'Gerencia + Equipo proyecto'),
        ('8', 'Diseñar conexión SAP S/4HANA y CDS Views disponibles.',                     'Consultor + TI S4'),
        ('9', 'Definir usuarios, roles y matriz de accesos por sociedad.',                  'TI + Planeación Fin.'),
        ('10','Programar Sesión de Diseño N°2: EEFF y dashboards.',                        'Equipo proyecto'),
    ]
    tbl_steps = doc.add_table(rows=len(steps)+1, cols=3)
    header_row(tbl_steps, ['#','Acción','Responsable'], widths=[0.4, 4.5, 1.8])
    for i, row in enumerate(steps):
        data_row(tbl_steps, i+1, row, alt=(i%2==0), bold_first=False)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ══════════════════════════════════════════════════════════════════════════
    # CONTROL DE VERSIONES
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '7. Control de Versiones del Documento', 1)
    tbl_ver = doc.add_table(rows=3, cols=4)
    header_row(tbl_ver, ['Versión','Fecha','Descripción','Autor'], widths=[0.8,1.2,3.8,0.8])
    ver_data = [
        ('1.0', 'Jun-2026', 'Versión inicial basada en análisis de modelos SAC', 'Equipo SAC'),
        ('2.0', 'Jun-2026', 'Enriquecida con Sesión Diseño N°1: proceso actual, decisiones, parking lot y próximos pasos', 'Equipo SAC'),
    ]
    for i, row in enumerate(ver_data):
        data_row(tbl_ver, i+1, row, alt=(i%2==0), bold_first=True)

    # ── PIE DE PÁGINA ─────────────────────────────────────────────────────────
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run('Documento de Diseño SAP Analytics Cloud  |  Ciudad Limpia / Grupo Fanalca  |  Versión 2.0  |  Confidencial')
        run.font.size = Pt(8); run.font.color.rgb = GRAY_MED

    doc.save(out_path)
    print(f'Documento guardado: {out_path}')

# ─── MAIN ─────────────────────────────────────────────────────────────────────
if __name__ == '__main__':
    FC  = '/home/user/Johnmeg/arquitectura_sac.png'
    OUT = '/home/user/Johnmeg/Diseno_Modelos_SAC_CiudadLimpia_v2.docx'
    make_flowchart(FC)
    make_document(FC, OUT)
    print('¡Proceso completado!')
