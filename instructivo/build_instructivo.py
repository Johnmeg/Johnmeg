# -*- coding: utf-8 -*-
"""Instructivo: Conexión de importación de datos a Microsoft SharePoint en
SAP Analytics Cloud (SAC) vía SAP Integration Suite — Open Connectors.
Basado en la Nota/KBA SAP 3446524 y documentación pública de SAP."""
import sys, os
HERE=os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'generador_diseno'))
from docx_helpers import Builder, NAVY, NAVY2, BLUE_D, RED, GREEN, GREY
from docx.shared import Pt, RGBColor

TPL=os.path.join(HERE,'..','plan_fs','template_funcional.docx')
LOGO_BUILD=os.path.join(HERE,'..','plan_fs','logos','build.png')
LOGO_FAN=os.path.join(HERE,'..','plan_fs','logos','fanalca.png')
DIAG=os.path.join(HERE,'img','sharepoint_sac_flujo.png')
IMG=os.path.join(HERE,'img')
def fig(name, n, desc):
    b.figure(os.path.join(IMG,name),
             f"Figura {n}. Ilustración esquemática — {desc} (los números corresponden a los pasos).")

b=Builder(TPL); b.setup_headers(LOGO_FAN)
hp=b.doc.sections[0].header.paragraphs[0]
for r in hp.runs:
    if 'Documento de Diseño' in r.text:
        r.text='Instructivo · Conexión de importación SharePoint → SAP Analytics Cloud'
    elif 'Programa SAP BUILD' in r.text:
        r.text='\tNota SAP 3446524'

def fm(text):
    p=b.doc.add_paragraph()
    p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5)
    r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(13); r.font.bold=True
    r.font.color.rgb=RGBColor.from_string(BLUE_D); b._bottom_border(p); return p

def steps(rows):
    b.table(["#","Acción","Detalle / valor"], rows, widths=[0.5,2.6,3.5], size=9.0)

# =====================================================================
# PORTADA
# =====================================================================
b.spacer(30)
b.cover_logo(LOGO_BUILD, width=3.0)
b.spacer(16)
b.cover_title([
    ("INSTRUCTIVO", 24, BLUE_D, True, 4),
    ("Conexión de importación de datos a Microsoft SharePoint", 14, NAVY, False, 2),
    ("en SAP Analytics Cloud (SAC) — vía SAP Integration Suite · Open Connectors", 12, NAVY, False, 2),
    ("Basado en la Nota SAP 3446524", 11, GREY, False, 20),
])
b.meta_table([
    ("Tema", "Crear una conexión de importación a Microsoft SharePoint en SAC"),
    ("Nota SAP", "3446524 (KBA) · Referencia complementaria: KBA 3435156"),
    ("Producto", "SAP Analytics Cloud · SAP Integration Suite (Open Connectors) · SAP BTP"),
    ("Aplica a", "Inquilinos de SAC en data center no-SAP (entorno Cloud Foundry)"),
    ("Versión", "1.0"),
    ("Fecha", "(por confirmar)"),
    ("Clasificación", "Uso interno"),
])
b.spacer(24)
b.cover_logo(LOGO_FAN, width=2.9)
b.page_break()

# =====================================================================
# FRONT MATTER
# =====================================================================
b.callout("Origen de la información", [
    "Este instructivo se elaboró a partir de la Nota SAP 3446524 «How to create an import data connection "
    "to Microsoft SharePoint in SAP Analytics Cloud (SAC)» y de la documentación pública de SAP sobre "
    "Open Connectors.",
    "El portal SAP for Me (me.sap.com) requiere autenticación S-user; se recomienda verificar los valores "
    "y actualizaciones exactos directamente en la nota antes de ejecutar el procedimiento en productivo.",
], accent=NAVY2)
fm("Control de versiones")
b.table(["Versión","Fecha","Autor","Descripción"],
        [["1.0","(fecha)","Equipo SAP BUILD","Versión inicial del instructivo"]],
        widths=[0.8,1.1,1.6,3.1])
fm("Contenido")
b.toc()
b.page_break()

# =====================================================================
# 1. OBJETIVO Y ALCANCE
# =====================================================================
b.h1("Objetivo y alcance")
b.para("Este instructivo guía, paso a paso, la creación de una conexión de importación de datos desde "
       "Microsoft SharePoint hacia SAP Analytics Cloud (SAC), empleando SAP Integration Suite — Open "
       "Connectors, conforme a la Nota SAP 3446524.")
b.para("Alcance. Aplica a inquilinos de SAC alojados en un data center no-SAP (entorno Cloud Foundry). "
       "Permite importar archivos y datos de SharePoint Online para construir modelos e historias. No "
       "cubre conexiones en vivo (live); el acceso a SharePoint mediante Open Connectors es de "
       "importación (adquisición de datos).", size=9.5)

# =====================================================================
# 2. CONCEPTOS Y ARQUITECTURA
# =====================================================================
b.h1("Conceptos y arquitectura")
b.para("SAP Integration Suite — Open Connectors es un servicio de SAP BTP que expone conectores "
       "normalizados hacia aplicaciones de terceros (SharePoint, OneDrive, Dropbox, entre otras) mediante "
       "APIs unificadas. Para conectar SAC con SharePoint se crea una instancia del conector de SharePoint "
       "en Open Connectors (autenticada por OAuth) y, sobre ella, una conexión de importación en SAC.")
b.figure(DIAG, "Figura 1. Flujo de la solución: SharePoint → Open Connectors → SAP Analytics Cloud.")

# =====================================================================
# 3. PRERREQUISITOS
# =====================================================================
b.h1("Prerrequisitos")
b.table(["#","Prerrequisito","Detalle"],
        [["1","Inquilino de SAC en data center no-SAP","Debe ejecutarse en entorno Cloud Foundry."],
         ["2","Cuenta de SAP BTP con Open Connectors","El servicio Open Connectors debe estar habilitado."],
         ["3","Derechos de administrador","En SAC y en la subcuenta de SAP BTP."],
         ["4","Open Connectors integrado en SAC","Configuración única (ver Fase 1 · KBA 3435156)."],
         ["5","Cuenta de Microsoft SharePoint Online","Con permisos para registrar una aplicación OAuth."]],
        widths=[0.5,2.5,3.6], size=9.0)
b.callout("Importante — deprecación de la autenticación ACS", [
    "La autenticación de SharePoint basada en «client secret» de Azure ACS se deprecia el 2 de abril de 2026.",
    "Para integraciones nuevas, registre la aplicación en Azure AD (Microsoft Entra ID) y prefiera la "
    "autenticación basada en certificado; migre las integraciones existentes con «client secret».",
], accent=RED, fill="FDECEA")

# =====================================================================
# 4. PROCEDIMIENTO
# =====================================================================
b.h1("Procedimiento")
b.callout("Sobre las figuras de este apartado",
          "Las imágenes son ilustraciones esquemáticas de las pantallas (wireframes), no capturas reales. "
          "Los rótulos numerados corresponden a los pasos de la tabla de cada fase. La interfaz real de SAP "
          "puede variar según la versión.", accent=NAVY2)

b.h2("Fase 1 — Habilitar Open Connectors en SAC (una sola vez)")
b.para("Si Open Connectors aún no está integrado en su inquilino de SAC, realice esta configuración una "
       "sola vez (requiere administrador). Referencia: KBA SAP 3435156.", size=9.5)
steps([["1","Abrir la configuración","En SAC: menú lateral → Administración del sistema → Configuración de origen de datos."],
       ["2","Iniciar la integración","En el área «Open Connectors», seleccionar «Let's integrate your Open Connectors Account»."],
       ["3","Seleccionar la región","Elegir la Región de la subcuenta de SAP BTP donde está Open Connectors."],
       ["4","Ingresar el User Secret","Introducir el «Open Connectors User Secret»."],
       ["5","Ingresar el Organization Secret","Introducir el «Open Connectors Organization Secret»."],
       ["6","Guardar","Confirmar. Los secretos se obtienen en Open Connectors (perfil / API Keys)."]])
fig("paso1_sac_openconnectors.png", 2, "integración de Open Connectors en SAP Analytics Cloud")

b.h2("Fase 2 — Registrar la aplicación OAuth en Microsoft")
b.para("Registre una aplicación que autorice a Open Connectors a acceder a SharePoint y obtenga el Client "
       "ID y el Client Secret.", size=9.5)
steps([["1","Acceder al registro de apps","Azure Portal → «App registrations» (o, en el modelo ACS, {sitio}/_layouts/15/appregnew.aspx)."],
       ["2","Crear la aplicación","«New registration». Anotar Application (Client) ID y Directory (Tenant) ID."],
       ["3","Generar el Client Secret","Crear un secreto y copiar su valor de inmediato."],
       ["4","Configurar el Redirect URI","Usar la Callback URL de Open Connectors como URI de redirección."],
       ["5","Otorgar permisos","Conceder los permisos de SharePoint requeridos (p. ej., alcance AllSites.Manage)."]])
fig("paso2a_azure_appreg.png", 3, "registro de la aplicación en Azure (Microsoft Entra ID)")
b.callout("Guarde el Client Secret", "El valor del Client Secret solo se muestra una vez y no puede "
          "recuperarse después. Guárdelo en un lugar seguro; si se pierde, deberá generar uno nuevo.",
          accent=NAVY2, fill="FFF7E6", icon="⚠")
fig("paso2b_azure_secret.png", 4, "generación del Client Secret en Azure")

b.h2("Fase 3 — Crear la instancia del conector SharePoint en Open Connectors")
b.para("En Open Connectors, cree y autentique una instancia del conector de SharePoint con las "
       "credenciales OAuth obtenidas.", size=9.5)
steps([["1","Buscar el conector","En Open Connectors, abrir el catálogo y buscar «SharePoint»."],
       ["2","Crear/autenticar la instancia","Seleccionar «Authenticate» / «Create Instance»."],
       ["3","Nombre de la instancia","Asignar un nombre (Name) descriptivo."],
       ["4","Dirección del sitio","SharePoint Site Address: {su_dominio}.sharepoint.com."],
       ["5","API Key","Ingresar el Client ID en «API Key» (oauth.api.key)."],
       ["6","API Secret","Ingresar el Client Secret en «API Secret» (oauth.api.secret)."],
       ["7","Crear e iniciar sesión","«Create Instance» → autenticarse en SharePoint y autorizar. Se genera la instancia y su token."]])
b.callout("Correspondencia de credenciales", [
    "Client ID  →  API Key (oauth.api.key)",
    "Client Secret  →  API Secret (oauth.api.secret)",
    "Redirect URI  →  Callback URL (oauth.callback.url)",
], accent=NAVY2)
fig("paso3_openconnectors_instance.png", 5, "instancia del conector SharePoint en Open Connectors")

b.h2("Fase 4 — Crear la conexión de importación a SharePoint en SAC")
b.para("Con Open Connectors integrado (Fase 1) y la instancia creada (Fase 3), cree la conexión de "
       "importación en SAC.", size=9.5)
steps([["1","Abrir Conexiones","En SAC: Conexiones (o «Adquirir datos» dentro de un modelo)."],
       ["2","Añadir conexión","Seleccionar «Añadir conexión» y filtrar por la categoría «Open Connectors»."],
       ["3","Tipo de conexión","Elegir una «query-based data connection» bajo «Adquirir datos» (Acquire Data)."],
       ["4","Connection Name","Asignar un nombre único a la conexión."],
       ["5","SharePoint Site Address","Ingresar la dirección del sitio SIN el prefijo «https://»."],
       ["6","OAuth API Key","Ingresar el Client ID."],
       ["7","OAuth API Secret","Ingresar el Client Secret."],
       ["8","MS OAuth Scope","Dejar el valor por defecto: AllSites.Manage."],
       ["9","Use Scope","Dejar el valor por defecto: true."],
       ["10","Iniciar sesión y guardar","Autenticarse en SharePoint con la cuenta y guardar la conexión."]])
b.callout("Elimine el prefijo «https://»", "En el campo «SharePoint Site Address» la dirección NO debe "
          "incluir «https://»; escríbala como dominio.sharepoint.com/sites/… De lo contrario, la conexión fallará.",
          accent=RED, fill="FDECEA", icon="⚠")
fig("paso4_sac_conexion.png", 6, "conexión de importación a SharePoint en SAP Analytics Cloud")

b.h2("Fase 5 — Verificación y uso")
steps([["1","Crear un modelo","Nuevo modelo → «Adquirir datos» → seleccionar la conexión de SharePoint."],
       ["2","Elegir el recurso","Seleccionar la biblioteca/consulta o el archivo a importar."],
       ["3","Previsualizar y cargar","Revisar los datos, aplicar transformaciones e importar."],
       ["4","Construir","Usar los datos en modelos e historias de SAC."]])
fig("paso5_sac_importar.png", 7, "adquisición e importación de datos en SAP Analytics Cloud")

# =====================================================================
# 5. REFERENCIA DE CAMPOS
# =====================================================================
b.h1("Referencia de campos de la conexión (SAC)")
b.table(["Campo","Descripción","Valor / ejemplo"],
        [["Connection Name","Nombre único de la conexión","SP_Finanzas"],
         ["SharePoint Site Address","Dirección del sitio, sin «https://»","contoso.sharepoint.com/sites/finanzas"],
         ["OAuth API Key","Client ID de la aplicación registrada","(GUID de la app)"],
         ["OAuth API Secret","Client Secret de la aplicación","(secreto de la app)"],
         ["MS OAuth Scope","Alcance de permisos de SharePoint","AllSites.Manage (por defecto)"],
         ["Use Scope","Aplicar el alcance indicado","true (por defecto)"]],
        widths=[1.7,2.5,2.4], size=8.8)

# =====================================================================
# 6. CONSIDERACIONES Y SOLUCIÓN DE PROBLEMAS
# =====================================================================
b.h1("Consideraciones y solución de problemas")
b.table(["Situación","Recomendación"],
        [["La conexión falla al guardar","Verifique que la dirección del sitio no incluya «https://» y que el sitio sea accesible."],
         ["No se dispone del Client Secret","No es recuperable: genere uno nuevo en la app y actualice la instancia del conector."],
         ["Autenticación ACS","Se deprecia el 02/04/2026: migre a registro de app en Azure AD / certificado."],
         ["No aparece la categoría Open Connectors","Complete primero la Fase 1 (integración de Open Connectors en SAC)."],
         ["Permisos insuficientes","El alcance (scope) debe permitir leer los sitios/bibliotecas (p. ej. AllSites.Read/Manage)."],
         ["Región incorrecta","Al integrar Open Connectors, seleccione la Región correcta de la subcuenta de BTP."]],
        widths=[2.3,4.3], size=8.8)

# =====================================================================
# 7. REFERENCIAS
# =====================================================================
b.h1("Referencias")
b.bullets([
    "Nota / KBA SAP 3446524 — «How to create an import data connection to Microsoft SharePoint in SAP Analytics Cloud (SAC)».",
    "KBA SAP 3435156 — Configurar SAP Integration Suite Open Connectors en SAP Analytics Cloud.",
    "Documentación SAC — «Use SAP Integration Suite Open Connectors» y «Import Data Connection to SAP Open Connectors Query-Based Data Sources».",
    "Documentación SAP Open Connectors — «SharePoint API Provider Setup» y «SharePoint — Authenticate a Connector».",
])

fn=os.path.join(HERE,'Instructivo_Conexion_SharePoint_SAC.docx')
b.save(fn)
print("GUARDADO:", fn)
