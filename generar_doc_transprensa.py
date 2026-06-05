import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# ─── COLORES CORPORATIVOS ─────────────────────────────────────────────────────
SAP_BLUE   = RGBColor(0x00, 0x61, 0x9A)
SAP_DARK   = RGBColor(0x00, 0x33, 0x69)
SAP_GOLD   = RGBColor(0xF0, 0xAB, 0x00)
SAP_LIGHT  = RGBColor(0xE8, 0xF4, 0xFD)
SAP_GREEN  = RGBColor(0x10, 0x7E, 0x3E)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_DARK  = RGBColor(0x35, 0x35, 0x35)
GRAY_MED   = RGBColor(0x75, 0x75, 0x75)
GRAY_LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
ORANGE     = RGBColor(0xE8, 0x6A, 0x19)

def hex_rgb(h):
    h = h.lstrip('#')
    return tuple(int(h[i:i+2], 16) / 255 for i in (0, 2, 4))

# ─── HELPERS DOCX ─────────────────────────────────────────────────────────────
def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  '%02X%02X%02X' % (rgb[0], rgb[1], rgb[2]))
    tcPr.append(shd)

def set_cell_borders(cell, sides=('top','bottom','left','right'), color='B0C4D8', sz=4):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for s in sides:
        b = OxmlElement(f'w:{s}')
        b.set(qn('w:val'),   'single')
        b.set(qn('w:sz'),    str(sz))
        b.set(qn('w:color'), color)
        tcBorders.append(b)
    tcPr.append(tcBorders)

def para_font(para, size_pt, bold=False, color=None, align=None, italic=False,
              space_before=0, space_after=0):
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after  = Pt(space_after)
    if align:
        para.alignment = align
    for run in para.runs:
        run.font.size   = Pt(size_pt)
        run.font.bold   = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = color

def add_heading(doc, text, level=1, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    if level == 1:
        run.font.size  = Pt(16)
        run.font.bold  = True
        run.font.color.rgb = SAP_DARK
        pPr  = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bot  = OxmlElement('w:bottom')
        bot.set(qn('w:val'),   'single')
        bot.set(qn('w:sz'),    '6')
        bot.set(qn('w:color'), '003369')
        pBdr.append(bot)
        pPr.append(pBdr)
    elif level == 2:
        run.font.size  = Pt(13)
        run.font.bold  = True
        run.font.color.rgb = SAP_BLUE
    elif level == 3:
        run.font.size  = Pt(11)
        run.font.bold  = True
        run.font.color.rgb = SAP_BLUE

def add_body(doc, text, space_before=2, space_after=4, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.size   = Pt(10)
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def add_bullet(doc, text, level=0, space_before=1, space_after=1):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before    = Pt(space_before)
    p.paragraph_format.space_after     = Pt(space_after)
    p.paragraph_format.left_indent     = Cm(0.5 + level * 0.5)
    run = p.add_run(text)
    run.font.size = Pt(10)

def header_row(table, cols, bg=None):
    bg = bg or SAP_DARK
    row = table.rows[0]
    for i, txt in enumerate(cols):
        cell = row.cells[i]
        set_cell_bg(cell, bg)
        set_cell_borders(cell, color='003369', sz=6)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(txt)
        run.font.size  = Pt(10)
        run.font.bold  = True
        run.font.color.rgb = WHITE

def data_row(table, row_idx, values, alt=False, center_cols=None, bold_col=None):
    center_cols = center_cols or []
    row = table.rows[row_idx]
    bg  = SAP_LIGHT if alt else WHITE
    for i, val in enumerate(values):
        cell = row.cells[i]
        set_cell_bg(cell, bg)
        set_cell_borders(cell, color='C5D8EA', sz=2)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p   = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(val))
        run.font.size = Pt(9)
        if bold_col is not None and i == bold_col:
            run.font.bold  = True
            run.font.color.rgb = SAP_DARK

def add_info_box(doc, title, lines, bg_title=None, bg_body=None):
    bg_title = bg_title or SAP_BLUE
    bg_body  = bg_body  or SAP_LIGHT
    tbl = doc.add_table(rows=1 + len(lines), cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Title row
    tc = tbl.rows[0].cells[0]
    set_cell_bg(tc, bg_title)
    set_cell_borders(tc, color='003369', sz=6)
    p = tc.paragraphs[0]
    run = p.add_run(title)
    run.font.size  = Pt(10)
    run.font.bold  = True
    run.font.color.rgb = WHITE
    # Body rows
    for i, line in enumerate(lines, 1):
        tc = tbl.rows[i].cells[0]
        set_cell_bg(tc, bg_body)
        set_cell_borders(tc, color='B0C4D8', sz=2)
        p = tc.paragraphs[0]
        run = p.add_run(line)
        run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_footer(doc):
    section = doc.sections[0]
    footer  = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('SAP BUILD | Grupo Fanalca — Sociedad Transprensa | Documento Confidencial | 2026')
    run.font.size   = Pt(8)
    run.font.color.rgb = GRAY_MED
    run.font.italic = True

# ─── FLOWCHART ────────────────────────────────────────────────────────────────
def generar_flowchart(path):
    fig, ax = plt.subplots(figsize=(20, 14))
    ax.set_xlim(0, 20)
    ax.set_ylim(0, 14)
    ax.axis('off')
    fig.patch.set_facecolor('#F7FAFD')

    def band(y, h, color, label):
        ax.add_patch(plt.Rectangle((0, y), 20, h, color=color, zorder=0, alpha=0.25))
        ax.text(0.18, y + h/2, label, va='center', ha='left',
                fontsize=8.5, color='#333333', fontweight='bold',
                rotation=90, zorder=2)

    def box(x, y, w, h, label, sub='', fc='#00619A', tc='white', fs=9, subfs=7.5, radius=0.25):
        ax.add_patch(FancyBboxPatch((x, y), w, h,
                     boxstyle=f'round,pad=0.05,rounding_size={radius}',
                     fc=fc, ec='white', lw=1.5, zorder=3))
        cy = y + h/2 + (0.15 if sub else 0)
        ax.text(x + w/2, cy, label, ha='center', va='center',
                fontsize=fs, color=tc, fontweight='bold', zorder=4,
                wrap=True, multialignment='center')
        if sub:
            ax.text(x + w/2, y + h/2 - 0.22, sub, ha='center', va='center',
                    fontsize=subfs, color=tc, style='italic', zorder=4)

    def arrow(x1, y1, x2, y2, color='#003369', lw=1.5, style='->'):
        ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle=style, color=color,
                                   lw=lw, connectionstyle='arc3,rad=0.0'))

    # ── bands ──
    band(0.3,  1.6,  '#E3EDF7', 'FUENTES DE DATOS')
    band(2.1,  1.5,  '#FFF3D6', 'INTEGRACIÓN')
    band(3.8,  3.8,  '#EAF6EC', 'MODELOS SAC – PLANEACIÓN')
    band(7.8,  1.5,  '#EDE8F5', 'ESTADOS FINANCIEROS')
    band(9.5,  2.0,  '#FDE8E8', 'SALIDAS / DASHBOARDS')
    band(11.7, 1.8,  '#E0F0FF', 'MODELOS FUTUROS')

    # ── Fuentes de datos ──
    box(0.6,  0.5,  2.8, 1.1, 'Silotrans /\nSilotrack', 'Operaciones Logísticas',
        fc='#005B99', fs=8.5, subfs=7)
    box(4.0,  0.5,  2.8, 1.1, 'SAP S/4HANA', 'Contabilidad Financiera',
        fc='#005B99', fs=8.5, subfs=7)
    box(7.4,  0.5,  2.8, 1.1, 'SIESA Payroll', 'Nómina y RRHH',
        fc='#005B99', fs=8.5, subfs=7)
    box(10.8, 0.5,  2.8, 1.1, 'Futura:\nFuente TH', 'Talento Humano / Legado',
        fc='#4A7FA5', fs=8, subfs=7)
    box(14.2, 0.5,  2.8, 1.1, 'Manual /\nExcel Budget', 'Datos de Planificación',
        fc='#4A7FA5', fs=8, subfs=7)
    box(17.4, 0.5,  2.0, 1.1, 'Macros\nEconómicos', 'IPC, TRM',
        fc='#4A7FA5', fs=8, subfs=7)

    # ── Integración ──
    box(4.0,  2.3,  12.0, 1.0, 'SAP Datasphere   —   Capa de Integración y Maestros',
        fc='#F0AB00', tc='#003369', fs=10)
    # arrows up → datasphere
    for x in [2.0, 5.4, 8.8]:
        arrow(x, 1.6, x, 2.3)
    # manual arrow
    arrow(15.6, 1.6, 12.0, 2.3)
    arrow(18.4, 1.6, 16.0, 2.3)
    arrow(12.2, 1.6, 12.5, 2.3)

    # ── Modelos SAC ──
    # Ingresos
    box(0.8,  4.1,  5.2, 3.2,
        'Modelo Ingresos\nTransprensa',
        'Measure: Importe (Decimal)\nDims: CUENTA · Date · MONEDA\nCLIENTE · SERVICIO · AUDITORIA\nCEBE_TRANS · SOCIEDAD_TRANS\nVersion',
        fc='#107E3E', fs=10, subfs=7.5)

    # Gastos
    box(7.0,  4.1,  5.2, 3.2,
        'Modelo Gastos y Costos\nTransprensa',
        'Measure: Importe (Decimal)\nDims: CUENTA · Date · MONEDA\nCECOS_TRANS · AUDITORIA\nCEBE_TRANS · SOCIEDAD_TRANS\nVersion',
        fc='#E86A19', fs=10, subfs=7.5)

    # EEFF placeholder box (feeds into next band)
    box(13.2, 4.1,  5.8, 3.2,
        'Modelo EEFF\nTransprensa',
        'Measure: Importe (Decimal)\nDims: CUENTA · Date · MONEDA\nAUDITORIA · CEBE_TRANS\nSOCIEDAD_TRANS · Version',
        fc='#7B3FA0', fs=10, subfs=7.5)

    # arrow datasphere → models
    for x in [3.4, 9.6, 16.1]:
        arrow(x, 3.3, x, 4.1)

    # ── EEFF band ──
    box(5.5,  8.0,  9.0, 1.1,
        'Estado de Resultados Consolidado (P&G)  |  Balance General  |  Flujo de Caja Indirecto',
        fc='#4B2E83', fs=9)
    # arrows models → EEFF
    arrow(3.4, 7.3, 7.0, 8.0)
    arrow(9.6, 7.3, 10.0, 8.0)
    arrow(16.1, 7.3, 13.0, 8.0)

    # ── Dashboards ──
    dash_items = [
        (1.0,  'Dashboard\nEjecución\nPresupuestal'),
        (4.5,  'Análisis P×Q\nKg × Tarifa\nPor Regional'),
        (8.0,  'Rentabilidad\nPor Regional\n/ Negocio'),
        (11.5, 'Seguimiento\nKPIs\nFinancieros'),
        (15.0, 'Reportes\nJunta\nDirectiva'),
    ]
    for x, lbl in dash_items:
        box(x, 9.7, 3.2, 1.5, lbl, fc='#C0392B', fs=8.5)
    arrow(10.0, 9.1, 10.0, 9.7)

    # ── Modelos futuros ──
    box(2.5,  12.0, 5.5, 1.2, 'Modelo Talento Humano', 'Por empleado → CECOS (en construcción)',
        fc='#2980B9', fs=9.5, subfs=7.5)
    box(9.5,  12.0, 5.5, 1.2, 'Modelo Flujo de Caja Diario', 'Directo / Indirecto (en construcción)',
        fc='#2980B9', fs=9.5, subfs=7.5)
    box(16.0, 12.0, 3.5, 1.2, 'Business\nData Cloud',
        fc='#1A5276', fs=9)

    # ── Title ──
    ax.text(10, 13.55, 'Arquitectura de Modelos SAC — Sociedad Transprensa | Grupo Fanalca',
            ha='center', va='center', fontsize=13, fontweight='bold', color='#003369')
    ax.text(10, 13.15, 'Proyecto SAP BUILD  ·  SAP Analytics Cloud Planning  ·  2026',
            ha='center', va='center', fontsize=9, color='#555555', style='italic')

    # ── Legend ──
    legend_items = [
        (mpatches.Patch(color='#005B99'), 'Fuentes de Datos'),
        (mpatches.Patch(color='#F0AB00'), 'Integración DataSphere'),
        (mpatches.Patch(color='#107E3E'), 'Modelo Ingresos'),
        (mpatches.Patch(color='#E86A19'), 'Modelo Gastos/Costos'),
        (mpatches.Patch(color='#7B3FA0'), 'Modelo EEFF'),
        (mpatches.Patch(color='#4B2E83'), 'Estados Financieros'),
        (mpatches.Patch(color='#C0392B'), 'Salidas / Dashboards'),
        (mpatches.Patch(color='#2980B9'), 'Modelos Futuros'),
    ]
    fig.legend([h for h, _ in legend_items], [l for _, l in legend_items],
               loc='lower center', ncol=4, fontsize=8,
               framealpha=0.85, edgecolor='#CCCCCC',
               bbox_to_anchor=(0.5, -0.01))

    plt.tight_layout(rect=[0, 0.04, 1, 1])
    plt.savefig(path, dpi=150, bbox_inches='tight', facecolor='#F7FAFD')
    plt.close()
    print(f'Flowchart guardado: {path}')

# ─── DOCUMENTO WORD ───────────────────────────────────────────────────────────
def generar_doc(flowchart_path, doc_path):
    doc = Document()

    # ── Márgenes ──
    for sec in doc.sections:
        sec.top_margin    = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    # ── Fuente base ──
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10)

    add_footer(doc)

    # ══════════════════════════════════════════════════════════════════════════
    # PORTADA
    # ══════════════════════════════════════════════════════════════════════════
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, SAP_DARK)
    cell.width = Cm(16)
    for para in cell.paragraphs:
        para.clear()

    def cover_line(text, sz, bold=True, color=WHITE, space_b=4, space_a=4, italic=False):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_b)
        p.paragraph_format.space_after  = Pt(space_a)
        r = p.add_run(text)
        r.font.size   = Pt(sz)
        r.font.bold   = bold
        r.font.color.rgb = color
        r.font.italic = italic

    cover_line('', 6, space_b=18, space_a=0)
    cover_line('SAP BUILD', 13, color=SAP_GOLD, italic=True, space_a=2)
    cover_line('GRUPO FANALCA', 11, color=RGBColor(0xCC, 0xDD, 0xEE), bold=False)
    cover_line('', 6, space_b=6, space_a=6)
    cover_line('Documento de Diseño', 14, color=WHITE)
    cover_line('Modelos de Planeación SAC', 22, color=WHITE, space_a=2)
    cover_line('Sociedad Transprensa', 16, color=SAP_GOLD, space_b=2)
    cover_line('', 6, space_b=10, space_a=10)

    # Tabla metadata en portada
    meta_tbl = doc.add_table(rows=5, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ('Sociedad',    'Transprensa S.A.S.'),
        ('Proyecto',    'SAP BUILD – Implementación SAP S/4HANA + SAC'),
        ('Versión',     'v1.0'),
        ('Fecha',       '04 de junio de 2026'),
        ('Estado',      'Borrador para revisión'),
    ]
    for i, (k, v) in enumerate(meta_data):
        r = meta_tbl.rows[i]
        set_cell_bg(r.cells[0], SAP_BLUE)
        set_cell_bg(r.cells[1], SAP_LIGHT)
        set_cell_borders(r.cells[0], color='003369', sz=4)
        set_cell_borders(r.cells[1], color='B0C4D8', sz=2)
        pk = r.cells[0].paragraphs[0]
        pk.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rk = pk.add_run(k)
        rk.font.size = Pt(9); rk.font.bold = True; rk.font.color.rgb = WHITE
        pv = r.cells[1].paragraphs[0]
        rv = pv.add_run(v)
        rv.font.size = Pt(9); rv.font.color.rgb = SAP_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # TABLA DE CONTENIDO (manual)
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, 'Tabla de Contenido', level=1, space_before=6)
    toc_items = [
        ('1.', 'Propósito del documento', '4'),
        ('2.', 'Cómo usar el documento', '4'),
        ('3.', 'Escenario / Proceso de negocio', '4'),
        ('3.1.', 'Objetivos empresariales y beneficios esperados', '5'),
        ('3.2.', 'Descripción a alto nivel de Requisitos empresariales', '5'),
        ('4.', 'Estructura organizativa relevante', '6'),
        ('5.', 'Diseño y configuración de soluciones', '7'),
        ('5.1.', 'Alcance del proceso de solución', '7'),
        ('5.2.', 'Valor de Actividad de Configuración', '9'),
        ('5.3.', 'Roles', '10'),
        ('5.4.', 'Inventario de datos maestros susceptibles de migración', '10'),
        ('5.5.', 'Componente de Solución Técnica Relacionada', '11'),
    ]
    toc_tbl = doc.add_table(rows=len(toc_items), cols=3)
    for i, (num, title, pg) in enumerate(toc_items):
        r = toc_tbl.rows[i]
        bg = SAP_LIGHT if i % 2 == 0 else WHITE
        for c in r.cells:
            set_cell_bg(c, bg)
        indent = 1 if '.' in num[1:] else 0
        r.cells[0].paragraphs[0].add_run(num).font.size = Pt(9)
        p1 = r.cells[1].paragraphs[0]
        p1.paragraph_format.left_indent = Cm(indent * 0.5)
        run = p1.add_run(title)
        run.font.size = Pt(9)
        run.font.bold = ('.' not in num[1:])
        r.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r.cells[2].paragraphs[0].add_run(pg).font.size = Pt(9)
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # 1. PROPÓSITO DEL DOCUMENTO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '1. Propósito del documento', level=1)
    add_body(doc,
        'El presente documento describe el diseño de los modelos de planeación financiera '
        'implementados en SAP Analytics Cloud (SAC) para la Sociedad Transprensa S.A.S., '
        'en el marco del proyecto SAP BUILD de Grupo Fanalca. Su objetivo es formalizar la '
        'arquitectura de los modelos de Ingresos, Gastos y Costos, y Estados Financieros, '
        'documentar las reglas de negocio, las fuentes de datos, los datos maestros y la '
        'configuración técnica requerida para soportar el proceso de presupuestación y '
        'seguimiento de ejecución presupuestal de Transprensa.')
    add_body(doc,
        'Transprensa es una empresa de logística y transporte del Grupo Fanalca con operaciones '
        'en 21 regionales a lo largo de Colombia, 46 Centros de Recibo de Mercancía (CRM) y '
        'servicios de paqueteo, masivo, vehículos dedicados, almacenamiento e ingredientes. '
        'El proceso de presupuestación es centralizado desde el área financiera, con insumos '
        'provenientes de la gerencia general y la gerencia comercial.')

    add_info_box(doc, 'Alcance del documento',
        ['✔  Diseño de los 3 modelos SAC: Ingresos, Gastos y Costos, EEFF',
         '✔  Definición de dimensiones, medidas y capacidades de planeación',
         '✔  Reglas de negocio: cálculo P×Q, estacionalidad, versiones presupuestales',
         '✔  Integración con SAP S/4HANA, Silotrans y DataSphere',
         '✔  Inventario de datos maestros para migración',
         '✗  Fuera de alcance: modelos de Talento Humano y Flujo de Caja Diario (en diseño)'],
        bg_title=SAP_DARK)

    # ══════════════════════════════════════════════════════════════════════════
    # 2. CÓMO USAR EL DOCUMENTO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '2. Cómo usar el documento', level=1)
    add_body(doc,
        'Este documento está estructurado para ser consultado en secuencia. Se recomienda '
        'iniciar por la sección 3 para comprender el contexto de negocio de Transprensa, '
        'continuar con la sección 4 para entender la estructura organizativa, y revisar '
        'la sección 5 para el detalle técnico de configuración en SAC.')

    uso_tbl = doc.add_table(rows=6, cols=3)
    header_row(uso_tbl, ['Sección', 'Contenido', 'Audiencia'])
    rows_uso = [
        ('1 – 2', 'Propósito y uso del documento',             'Todos los interesados'),
        ('3',     'Escenario de negocio y objetivos',           'Gerencia, Sponsors'),
        ('4',     'Estructura organizativa',                    'Equipo funcional SAP'),
        ('5.1',   'Alcance y modelos SAC',                      'Consultores SAC / SAP'),
        ('5.2 – 5.5', 'Configuración, roles y datos maestros', 'Consultores técnicos'),
    ]
    for i, (s, c, a) in enumerate(rows_uso, 1):
        data_row(uso_tbl, i, [s, c, a], alt=(i % 2 == 0), center_cols=[0])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # 3. ESCENARIO / PROCESO DE NEGOCIO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '3. Escenario / Proceso de negocio', level=1)
    add_body(doc,
        'Transprensa S.A.S. es una empresa logística del Grupo Fanalca dedicada al transporte '
        'de carga en Colombia. Sus líneas de negocio principales son el paqueteo (B2B y B2C), '
        'el servicio masivo, vehículos dedicados, almacenamiento e ingredientes/mercancías '
        'peligrosas. La empresa cuenta con 21 regionales en el país y 46 CRM '
        '(Centros de Recibo de Mercancía) propios y externos. Su proceso de presupuestación '
        'se realiza centralizadamente desde el área financiera usando Excel y Power BI, '
        'con datos provenientes del software logístico Silotrans y del sistema contable SIESA. '
        'El proyecto SAP BUILD migra la planificación financiera a SAP Analytics Cloud.')

    # Diagrama
    add_heading(doc, 'Arquitectura de Modelos — Transprensa', level=2, space_before=10)
    doc.add_picture(flowchart_path, width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph('Figura 1. Arquitectura de modelos SAC Planning — Sociedad Transprensa')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(8)
    cap.runs[0].font.italic = True
    cap.runs[0].font.color.rgb = GRAY_MED
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3.1
    add_heading(doc, '3.1. Objetivos empresariales y beneficios esperados', level=2)
    add_body(doc,
        'La implementación de SAC Planning en Transprensa busca reemplazar el proceso manual '
        'de presupuestación en Excel y consolidar la planeación financiera en una plataforma '
        'corporativa integrada con SAP S/4HANA y Silotrans. Los objetivos estratégicos son:')

    obj_tbl = doc.add_table(rows=7, cols=3)
    header_row(obj_tbl, ['#', 'Objetivo Empresarial', 'Beneficio Esperado'])
    obj_rows = [
        ('1', 'Centralizar la planeación financiera en SAC',
              'Eliminar silos de información en Excel; control de versiones'),
        ('2', 'Automatizar el cálculo P×Q por regional y servicio',
              'Reducción de errores; mayor velocidad en ciclos de presupuesto'),
        ('3', 'Habilitar multi-versionamiento presupuestal',
              'Versión comercial vs. financiera; escenarios y simulaciones'),
        ('4', 'Integrar ejecución real desde Silotrans y S/4HANA',
              'Comparación presupuesto vs. real en tiempo real'),
        ('5', 'Visibilidad de P&G y EBITDA por regional',
              'Toma de decisiones a nivel de regional y centro de beneficio'),
        ('6', 'Estacionalidad y ajuste dinámico de mano de obra',
              'Optimización de costos de personal temporal ligado a kg'),
    ]
    for i, (n, obj, ben) in enumerate(obj_rows, 1):
        data_row(obj_tbl, i, [n, obj, ben], alt=(i % 2 == 0),
                 center_cols=[0], bold_col=1)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3.2
    add_heading(doc, '3.2. Descripción a alto nivel de Requisitos empresariales', level=2)
    add_body(doc,
        'Con base en la sesión de diseño realizada el 22 de mayo de 2026, se identificaron '
        'los siguientes requisitos de negocio para el proceso de presupuestación de Transprensa:')

    req_tbl = doc.add_table(rows=12, cols=4)
    header_row(req_tbl, ['ID', 'Área', 'Requisito', 'Modelo SAC'])
    req_rows = [
        ('REQ-01', 'Ingresos',
         'Presupuestar ingresos por regional y tipo de servicio (paqueteo, masivo, almacenamiento)',
         'Ingresos'),
        ('REQ-02', 'Ingresos',
         'Cálculo P×Q: Kilogramos proyectados × Tarifa (pesos/kg) por regional',
         'Ingresos'),
        ('REQ-03', 'Ingresos',
         'Discriminación de ingresos por tipo: Crédito (B2B) vs. Contado/contraentrega (CRM)',
         'Ingresos'),
        ('REQ-04', 'Ingresos',
         'Presupuestar clientes especiales por separado (Ingreso, T1, Soy Más)',
         'Ingresos'),
        ('REQ-05', 'Ingresos',
         'Definir estacionalidad mensual con base en histórico N-1',
         'Ingresos'),
        ('REQ-06', 'Ingresos',
         'Soporte de multi-versión: Presupuesto Comercial y Presupuesto Financiero',
         'Ingresos'),
        ('REQ-07', 'Costos',
         'Presupuesto de fletes de transporte por regional y destino (desde Silotrans)',
         'Gastos y Costos'),
        ('REQ-08', 'Costos',
         'Personal temporal variable ligado a kg: índice de personal por kg movilizado',
         'Gastos y Costos'),
        ('REQ-09', 'Costos',
         'Arrendamientos por contrato: IPC + factor, ponderado por tipo de servicio',
         'Gastos y Costos'),
        ('REQ-10', 'Costos',
         'Gastos fijos (histórico N-1 + IPC): honorarios, servicios públicos, ICA por ciudad',
         'Gastos y Costos'),
        ('REQ-11', 'EEFF',
         'Estado de Resultados (P&G administrativo – PCGA) con apertura por regional',
         'EEFF'),
    ]
    for i, row in enumerate(req_rows, 1):
        data_row(req_tbl, i, list(row), alt=(i % 2 == 0), center_cols=[0, 3])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Proceso de ingresos P×Q
    add_heading(doc, 'Metodología de Cálculo de Ingresos (P × Q)', level=3, space_before=8)
    add_body(doc,
        'El proceso de presupuestación de ingresos de Transprensa sigue la metodología '
        'P×Q (Precio × Cantidad), donde el precio corresponde a la tarifa por kilogramo '
        'y la cantidad a los kilogramos movilizados por regional y tipo de servicio. '
        'El flujo de cálculo es el siguiente:')
    steps = [
        'Base histórica (N-1): Ejecución real 2025 de kg y pesos por regional desde Silotrans',
        'Crecimiento en volumen: % de crecimiento en kg definido por gerencia por regional (ej: +21% global)',
        'Crecimiento en tarifa: Índice de incremento tarifario (ej: +15%), basado en IPC + análisis de mercado',
        'Distribución CRM vs. Crédito: % del total asignado a CRM propios/externos (contado) y clientes B2B (crédito)',
        'Estacionalidad mensual: Distribución del presupuesto anual en 12 meses con base en histórico N-1',
        'Clientes especiales: Ingreso, T1 y Soy Más presupuestados individualmente (licitaciones)',
        'Resultado: Presupuesto de ingresos por regional × tipo de servicio × mes',
    ]
    for s in steps:
        add_bullet(doc, s)

    # Proceso de costos
    add_heading(doc, 'Metodología de Cálculo de Costos', level=3, space_before=8)
    cost_tbl = doc.add_table(rows=8, cols=4)
    header_row(cost_tbl, ['Tipo Costo', 'Descripción', 'Driver', 'Fuente Dato'])
    cost_rows = [
        ('Fletes expedición', 'Costo de flete por kg despachado por regional', 'Kg × tarifa flete por destino', 'Silotrans'),
        ('Mano de obra temporal', 'Personal plataforma y ruta variable según kg', 'Índice persona/kg por regional', 'Nómina SIESA'),
        ('Fletes reexpedición', '% del flete a destinos sin flota propia', '% histórico de kg a terceros', 'Silotrans'),
        ('Comisiones CRM ext.', 'Comisión pagada a CRM externos por kg', 'Kg CRM ext. × tarifa comisión', 'Silotrans'),
        ('Comisiones fuerza ventas', 'Parte variable del salario comercial por kg', 'Kg crédito × índice', 'SIESA'),
        ('Arrendamientos', 'Contrato por sede, ponderado por servicio', 'IPC + factor contractual', 'Contratos'),
        ('Gastos fijos (IPC)', 'Honorarios, serv. públicos, asistencia técnica', 'Histórico N-1 × IPC', 'SIESA'),
    ]
    for i, row in enumerate(cost_rows, 1):
        data_row(cost_tbl, i, list(row), alt=(i % 2 == 0), center_cols=[2, 3])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # 4. ESTRUCTURA ORGANIZATIVA RELEVANTE
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '4. Estructura organizativa relevante', level=1)
    add_body(doc,
        'Transprensa opera bajo una estructura organizativa matricial: una dirección nacional '
        'centralizada en Bogotá con 21 regionales operativas distribuidas en Colombia. '
        'La siguiente tabla describe los elementos organizativos relevantes para la '
        'configuración de los modelos SAC:')

    org_tbl = doc.add_table(rows=8, cols=4)
    header_row(org_tbl, ['Entidad Org.', 'Descripción', 'Dimensión SAC', 'Ejemplos'])
    org_rows = [
        ('Sociedad',           'Unidad legal de reporte financiero',
         'SOCIEDAD_TRANS',     'Transprensa S.A.S.'),
        ('Regional / Sede',    '21 regionales operativas en Colombia',
         'CEBE_TRANS',         'Antioquia, Valle del Cauca, Cundinamarca, Eje Cafetero'),
        ('CRM',                '46 Centros de Recibo de Mercancía (propios y externos)',
         'CLIENTE / CEBE',     'CRM Antioquia (4), CRM Valle (10), CRM Santa Marta'),
        ('Tipo de Servicio',   'Línea de negocio / Canal de venta',
         'SERVICIO',           'Paqueteo, Masivo, Almacenamiento, Vehículos Dedicados, Ingredientes'),
        ('Centro de Costo',    'Unidad funcional para asignación de costos',
         'CECOS_TRANS',        'Plataforma Regional, Ruta, Fuerza de Ventas, Nacional'),
        ('Centro de Beneficio','Agrupación por tipo de servicio / negocio',
         'CEBE_TRANS',         'Paqueteo, Masivo, Almacenamiento, Ingreso, T1, Soy Más'),
        ('Cliente Especial',   'Clientes de alta representatividad (~30% ingresos)',
         'CLIENTE',            'Ingreso, T1, Soy Más (licitaciones)'),
    ]
    for i, row in enumerate(org_rows, 1):
        data_row(org_tbl, i, list(row), alt=(i % 2 == 0), bold_col=0, center_cols=[2])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_info_box(doc, 'Regiones de Operación de Transprensa',
        ['Antioquia · Valle del Cauca · Cundinamarca · Eje Cafetero (Armenia, Manizales, Pereira)',
         'Santander · Nariño · Bolívar · Atlántico · Santa Marta · Tolima · Meta',
         '21 regionales en total — Cobertura nacional con 46 CRM activos',
         'Sede operativa Funza (almacén Sodimac) para distribución Soy Más'],
        bg_title=SAP_GREEN)

    # ══════════════════════════════════════════════════════════════════════════
    # 5. DISEÑO Y CONFIGURACIÓN DE SOLUCIONES
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '5. Diseño y configuración de soluciones', level=1)

    # 5.1
    add_heading(doc, '5.1. Alcance del proceso de solución', level=2)
    add_body(doc,
        'El alcance de SAC Planning para Transprensa contempla 3 modelos de planeación '
        'interconectados. Los modelos de Ingresos y Gastos/Costos alimentan el modelo '
        'de Estados Financieros, generando el P&G administrativo por regional y el consolidado.')

    # Tabla modelos
    mod_tbl = doc.add_table(rows=4, cols=5)
    header_row(mod_tbl, ['Modelo SAC', 'Tipo', 'Medida', 'Dimensiones', 'Capacidades'])
    mod_rows = [
        ('Ingresos Transprensa',
         'Planning (Measure)',
         'Importe\n(Decimal)',
         'Version · CUENTA · Date · MONEDA\nCLIENTE · SERVICIO · AUDITORIA\nCEBE_TRANS · SOCIEDAD_TRANS',
         'Planning Capabilities ✔\nP×Q por regional\nMulti-versión'),
        ('Gastos y Costos\nTransprensa',
         'Planning (Measure)',
         'Importe\n(Decimal)',
         'Version · CUENTA · Date · MONEDA\nCECOS_TRANS · AUDITORIA\nCEBE_TRANS · SOCIEDAD_TRANS',
         'Planning Capabilities ✔\nCostos variables y fijos\nCentros de costo'),
        ('EEFF Transprensa',
         'Planning (Measure)',
         'Importe\n(Decimal)',
         'Version · CUENTA · Date · MONEDA\nAUDITORIA · CEBE_TRANS\nSOCIEDAD_TRANS',
         'Planning Capabilities ✔\nEstado de Resultados\nBalance · Flujo de Caja'),
    ]
    for i, row in enumerate(mod_rows, 1):
        data_row(mod_tbl, i, list(row), alt=(i % 2 == 0), bold_col=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Dimensiones detalle
    add_heading(doc, 'Detalle de Dimensiones por Modelo', level=3, space_before=8)
    dim_tbl = doc.add_table(rows=11, cols=5)
    header_row(dim_tbl, ['Dimensión', 'Tipo SAC', 'Ingresos', 'Gastos/Costos', 'EEFF'])
    dim_rows = [
        ('Version',        'Version',          '✔', '✔', '✔'),
        ('CUENTA',         'Account',           '✔', '✔', '✔'),
        ('Date',           'Date',              '✔', '✔', '✔'),
        ('MONEDA',         'Generic',           '✔', '✔', '✔'),
        ('AUDITORIA',      'Generic',           '✔', '✔', '✔'),
        ('CEBE_TRANS',     'Organization',      '✔', '✔', '✔'),
        ('SOCIEDAD_TRANS', 'Generic',           '✔', '✔', '✔'),
        ('CLIENTE',        'Generic',           '✔', '—', '—'),
        ('SERVICIO',       'Generic',           '✔', '—', '—'),
        ('CECOS_TRANS',    'Generic',           '—', '✔', '—'),
    ]
    for i, row in enumerate(dim_rows, 1):
        data_row(dim_tbl, i, list(row), alt=(i % 2 == 0),
                 center_cols=[1, 2, 3, 4], bold_col=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Versiones
    add_heading(doc, 'Configuración de Versiones Presupuestales', level=3, space_before=8)
    add_body(doc,
        'Transprensa maneja múltiples versiones presupuestales que deben ser configuradas '
        'en la dimensión Version de SAC:')
    ver_tbl = doc.add_table(rows=5, cols=3)
    header_row(ver_tbl, ['Versión', 'Descripción', 'Uso'])
    ver_rows = [
        ('PRESUPUESTO_FIN', 'Presupuesto Financiero oficial',
         'Base para P&G y seguimiento financiero'),
        ('PRESUPUESTO_COM', 'Presupuesto Comercial',
         'Meta para fuerza de ventas y regionales'),
        ('REAL',            'Ejecución real (actuals)',
         'Datos de Silotrans / SAP S/4HANA vía DataSphere'),
        ('FORECAST',        'Proyección intra-año',
         'Actualización de estimados durante el año'),
    ]
    for i, row in enumerate(ver_rows, 1):
        data_row(ver_tbl, i, list(row), alt=(i % 2 == 0), bold_col=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.2
    add_heading(doc, '5.2. Valor de Actividad de Configuración', level=2)
    cfg_tbl = doc.add_table(rows=11, cols=4)
    header_row(cfg_tbl, ['Actividad', 'Descripción', 'Prioridad', 'Estado'])
    cfg_rows = [
        ('Creación modelos SAC',
         'Crear los 3 modelos Planning en SAC con sus dimensiones y medidas',
         'Alta', 'Completado'),
        ('Carga datos maestros',
         'Cargar jerarquías de CUENTA, CEBE_TRANS, CECOS_TRANS, CLIENTE, SERVICIO',
         'Alta', 'Pendiente'),
        ('Configurar versiones',
         'Definir y configurar versiones: PRESUPUESTO_FIN, PRESUPUESTO_COM, REAL, FORECAST',
         'Alta', 'Pendiente'),
        ('Plantillas de entrada de datos',
         'Diseñar input forms para carga del presupuesto por región y servicio',
         'Alta', 'Pendiente'),
        ('Data Actions P×Q',
         'Programar data actions para cálculo automático de ingresos (kg × tarifa)',
         'Alta', 'Pendiente'),
        ('Data Action personal temporal',
         'Algoritmo para ajuste de personal temporal según variación de kg',
         'Media', 'Pendiente'),
        ('Conectores DataSphere',
         'Configurar extracción de Silotrans y S/4HANA hacia DataSphere',
         'Alta', 'En definición'),
        ('Estacionalidad',
         'Cargar factores de estacionalidad histórica por regional y servicio',
         'Media', 'Pendiente'),
        ('Historiales N-1',
         'Cargar ejecución real 2025 para cálculo de base y crecimiento',
         'Alta', 'Pendiente'),
        ('Reportes y dashboards',
         'Crear stories SAC: ejecución presupuestal, P×Q, EBITDA por regional',
         'Media', 'Pendiente'),
    ]
    for i, row in enumerate(cfg_rows, 1):
        color = 'Completado' == row[3]
        data_row(cfg_tbl, i, list(row), alt=(i % 2 == 0), center_cols=[2, 3])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.3
    add_heading(doc, '5.3. Roles', level=2)
    rol_tbl = doc.add_table(rows=6, cols=4)
    header_row(rol_tbl, ['Rol', 'Responsable', 'Acceso SAC', 'Responsabilidades'])
    rol_rows = [
        ('Administrador SAC',      'Consultor SAP BUILD',
         'Full admin',             'Configuración modelos, dimensiones, data actions, versiones'),
        ('Planeador Financiero',   'Liliana (Finanzas Transprensa)',
         'Lectura / Escritura',    'Carga y validación del presupuesto por regional y servicio'),
        ('Analista Comercial',     'Andrés (Comercial Transprensa)',
         'Lectura / Escritura',    'Definición de crecimientos por regional; revisión Presupuesto Comercial'),
        ('Gerencia Regional',      'Gerentes de Regional',
         'Solo lectura',           'Consulta de presupuesto y ejecución de su regional'),
        ('Gerencia General / CFO', 'Dirección Transprensa',
         'Solo lectura',           'Reportes consolidados, P&G, EBITDA, Junta Directiva'),
    ]
    for i, row in enumerate(rol_rows, 1):
        data_row(rol_tbl, i, list(row), alt=(i % 2 == 0), bold_col=0, center_cols=[2])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.4
    add_heading(doc, '5.4. Inventario de datos maestros susceptibles de migración', level=2)
    add_body(doc,
        'Los siguientes datos maestros deben ser cargados en SAC y alineados con las '
        'definiciones realizadas en los módulos SD, FI y CO de SAP S/4HANA. La fuente '
        'de verdad es SAP S/4HANA; Silotrans es fuente complementaria para los datos '
        'operativos (clientes, servicios, regionales):')

    dm_tbl = doc.add_table(rows=9, cols=5)
    header_row(dm_tbl, ['Dato Maestro', 'Dimensión SAC', 'Fuente', 'Jerarquía / Niveles', 'Volumen Est.'])
    dm_rows = [
        ('Plan de Cuentas',          'CUENTA',         'SAP S/4HANA (SKA1)',
         'Clase → Grupo → Cuenta', '~500 cuentas'),
        ('Centros de Beneficio',     'CEBE_TRANS',     'SAP S/4HANA (CEPC)',
         'Sociedad → Tipo Servicio → Regional', '~30 CEBE'),
        ('Centros de Costo',         'CECOS_TRANS',    'SAP S/4HANA (CSKS)',
         'Sociedad → Área → Centro', '~80 CECOS'),
        ('Clientes / Regionales',    'CLIENTE',        'Silotrans / SD',
         'Regional → CRM / Ejecutivo', '~50 + clientes especiales'),
        ('Tipos de Servicio',        'SERVICIO',       'Silotrans',
         'Categoría → Subcategoría (CV/Subservicio)', '~15 servicios'),
        ('Monedas',                  'MONEDA',         'SAP S/4HANA',
         'COP (principal) / USD', '2 monedas'),
        ('Sociedad',                 'SOCIEDAD_TRANS', 'SAP S/4HANA (T001)',
         'Grupo Fanalca → Transprensa S.A.S.', '1 sociedad'),
        ('Versiones',                'Version',        'SAC (configuración)',
         'Real · Presupuesto Fin · Presupuesto Com · Forecast', '4 versiones'),
    ]
    for i, row in enumerate(dm_rows, 1):
        data_row(dm_tbl, i, list(row), alt=(i % 2 == 0), bold_col=0, center_cols=[2, 4])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_info_box(doc, 'Pendiente — Alineación con otros módulos SAP',
        ['Coordinar con equipos SD / FI / CO las jerarquías definitivas antes de cargar en SAC',
         'Definir granularidad de CECOS_TRANS: este año sin CECOS en presupuesto; objetivo: incluir en 2027',
         'Confirmar fuente de datos reales: SAP S/4HANA vs. Silotrans (en definición)',
         'Reunión de integración pendiente: John + Sergio + equipo analítica (DataSphere + SAC)'],
        bg_title=ORANGE)

    # 5.5
    add_heading(doc, '5.5. Componente de Solución Técnica Relacionada', level=2)
    tech_tbl = doc.add_table(rows=8, cols=4)
    header_row(tech_tbl, ['Componente', 'Tipo', 'Rol en la solución', 'Estado'])
    tech_rows = [
        ('SAP Analytics Cloud (SAC)',      'SaaS – Cloud',
         'Plataforma de planeación: modelos, input forms, data actions, dashboards',
         'Activo'),
        ('SAP Datasphere',                 'SaaS – Cloud',
         'Capa de integración, virtualización y maestros entre SAC y sistemas fuente',
         'En configuración'),
        ('SAP S/4HANA',                    'On-premise / Cloud',
         'Fuente de datos contables: SKA1 (cuentas), CEPC (CBen), CSKS (CCOS), ACDOCA',
         'En implementación'),
        ('Silotrans / Silotrack',          'Legado On-premise',
         'Fuente operativa: kg, remesas, clientes, origen-destino, tarifas',
         'Activo (continúa)'),
        ('SIESA (Contabilidad)',            'Legado On-premise',
         'Fuente histórica de P&G y nómina hasta Go-Live S/4HANA',
         'Activo (transitorio)'),
        ('SIESA Payroll (Nómina)',          'Legado On-premise',
         'Fuente de datos de personal para modelo de nómina temporal',
         'Activo'),
        ('Power BI / Excel (actual)',       'Herramienta actual',
         'Proceso actual de presupuestación; a reemplazar por SAC',
         'A reemplazar'),
    ]
    for i, row in enumerate(tech_rows, 1):
        data_row(tech_tbl, i, list(row), alt=(i % 2 == 0), bold_col=0, center_cols=[3])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Modelos futuros
    add_heading(doc, 'Modelos SAC en Construcción (Fuera de alcance actual)', level=3, space_before=8)
    fut_tbl = doc.add_table(rows=3, cols=4)
    header_row(fut_tbl, ['Modelo', 'Descripción', 'Dimensiones Clave', 'Estado'])
    fut_rows = [
        ('Talento Humano',
         'Presupuesto de nómina a nivel de empleado; integración con SIESA Payroll',
         'Empleado · CECOS · CEBE · Cargo · Contrato',
         'En diseño'),
        ('Flujo de Caja Diario',
         'Gestión de liquidez: método directo; proyección semanal / diaria',
         'Fecha · Cuenta · Tipo de Flujo · Sociedad',
         'En diseño'),
    ]
    for i, row in enumerate(fut_rows, 1):
        data_row(fut_tbl, i, list(row), alt=(i % 2 == 0), bold_col=0, center_cols=[3])

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Nota final
    add_info_box(doc, 'Próximos pasos — Proyecto SAP BUILD',
        ['1. Revisión y aprobación de este documento de diseño (antes del 05/06/2026)',
         '2. Reunión de integración: SAC + DataSphere + módulos SD/FI/CO para alinear maestros',
         '3. Sesión de diseño: Modelo de Talento Humano (fuente: SIESA Payroll)',
         '4. Sesión de diseño: Modelo de Flujo de Caja Diario',
         '5. Inicio fase de realización: configuración, carga de maestros y plantillas (jul-ago 2026)',
         '6. Reunión arquitectura analítica: SAC + DataSphere + Business Data Cloud (Carolina / Kevin)'],
        bg_title=SAP_DARK)

    doc.save(doc_path)
    print(f'Documento guardado: {doc_path}')

# ─── MAIN ─────────────────────────────────────────────────────────────────────
if __name__ == '__main__':
    BASE = '/home/user/Johnmeg'
    flowchart_path = os.path.join(BASE, 'arquitectura_transprensa.png')
    doc_path       = os.path.join(BASE, 'Diseno_Modelos_SAC_Transprensa_v1.docx')

    print('Generando flowchart...')
    generar_flowchart(flowchart_path)

    print('Generando documento Word...')
    generar_doc(flowchart_path, doc_path)

    print('Listo!')
