# Uso: python3 scripts/make-addin-excel.py samples/fanalca/Plantillas_AddIn_SAC_y_Reportes_Fanalca.xlsx
# Requiere openpyxl. Abra el archivo en Excel (o recalcule) para ver los valores de las fórmulas.
# Genera el libro de ejemplo: plantillas de carga para el SAP Analytics Cloud add-in
# for Microsoft Excel (modelos de Fanalca) y reportes P&L, Balance y Flujo de caja.
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule

OUT = sys.argv[1]

NAVY, BLUE, SKY, CYAN = '071D49', '005EB8', '0085CF', '00AFE7'
GREY, LGREY, INFILL, SOFT = '5B6578', 'F2F4F7', 'FFF2CC', 'E8F2FB'
FONT = 'Arial'

def f(bold=False, color='1D2433', size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)

F_IN = f(color='0000FF')            # valores digitados
F_LINK = f(color='008000')          # vínculos a otra hoja
F_CALC = f()                        # fórmulas
FILL_IN = PatternFill('solid', fgColor=INFILL)
FILL_HEAD = PatternFill('solid', fgColor=NAVY)
FILL_SEC = PatternFill('solid', fgColor=SOFT)
FILL_DIM = PatternFill('solid', fgColor=LGREY)
FILL_TOT = PatternFill('solid', fgColor='DCE6F2')
thin = Side(style='thin', color='D9E2EE')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
TOPLINE = Border(top=Side(style='thin', color=NAVY))
DBL = Border(top=Side(style='thin', color=NAVY), bottom=Side(style='double', color=NAVY))

NUM = '#,##0;(#,##0);"-"'
NUM_COP = '#,##0;-#,##0;"-"'
PCT = '0.0%;(0.0%);"-"'

MESES = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
PERIODOS = [f'{m} 2026' for m in MESES]

wb = Workbook()

def title(ws, text, sub=None, width_to=16):
    ws['A1'] = text
    ws['A1'].font = f(True, 'FFFFFF', 14)
    for c in range(1, width_to + 1):
        ws.cell(1, c).fill = FILL_HEAD
    ws.row_dimensions[1].height = 26
    ws['A1'].alignment = Alignment(vertical='center')
    if sub:
        ws['A2'] = sub
        ws['A2'].font = f(italic=True, color=GREY, size=9)
    ws.sheet_view.showGridLines = False

def header_row(ws, row, labels, start=1):
    for i, h in enumerate(labels):
        c = ws.cell(row, start + i, h)
        c.font = f(True, 'FFFFFF')
        c.fill = PatternFill('solid', fgColor=BLUE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 30

def addin_header(ws, modelo, model_id, filtros, row=3):
    """Bloque superior que imita el panel de la tabla del add-in."""
    ws.cell(row, 1, 'Fuente de datos (add-in)').font = f(True, BLUE)
    ws.cell(row + 1, 1, 'Modelo').font = f(color=GREY)
    ws.cell(row + 1, 2, f'{modelo}  ({model_id})').font = f(True)
    r = row + 2
    for k, v in filtros:
        ws.cell(r, 1, k).font = f(color=GREY)
        c = ws.cell(r, 2, v)
        c.font = F_IN
        c.fill = FILL_IN
        r += 1
    return r + 1

# ----------------------------------------------------------------------------- Listas
li = wb.active
li.title = 'Listas'
SOCIEDADES = ['FN_MOTOS', 'FN_AUTOS', 'FN_TUBOS']
CEBES = ['CB_MOTOS', 'CB_AUTOS', 'CB_TUBERIA', 'CB_CORPORATIVO']
CECOS = ['CC_PLANTA', 'CC_ADMIN', 'CC_VENTAS', 'CC_LOGISTICA']
RATIOS = ['UNIDADES', 'PRECIO', 'ING_VENTAS_NAL']
AUDITORIA = ['PRESUPUESTO_EXCEL', 'MANUAL', 'CALCULADO_DA']
VERSIONES = ['public.Presupuesto_2026', 'public.Forecast_2026', 'public.PRUEBAS']
CUENTAS_EG = [
    ('5105_SALARIOS', 'Salarios y prestaciones', 'Gastos de administración'),
    ('5110_HONORARIOS', 'Honorarios', 'Gastos de administración'),
    ('5135_SERVICIOS', 'Servicios públicos', 'Gastos de administración'),
    ('5145_MANTENIMIENTO', 'Mantenimiento y reparaciones', 'Gastos de administración'),
    ('5205_SALARIOS_VENTAS', 'Salarios fuerza de ventas', 'Gastos de ventas'),
    ('5235_PUBLICIDAD', 'Publicidad y mercadeo', 'Gastos de ventas'),
    ('5250_FLETES', 'Fletes y transporte', 'Gastos de ventas'),
]
CUENTAS_EF = [
    ('1105_CAJA', 'Efectivo y equivalentes', 'Activo'),
    ('1305_CLIENTES', 'Cuentas por cobrar clientes', 'Activo'),
    ('1435_INVENTARIOS', 'Inventarios', 'Activo'),
    ('1520_PPE', 'Propiedad, planta y equipo (bruto)', 'Activo'),
    ('1592_DEP_ACUMULADA', 'Depreciación acumulada', 'Activo'),
    ('2105_OBLIGACIONES_FIN', 'Obligaciones financieras', 'Pasivo'),
    ('2205_PROVEEDORES', 'Proveedores', 'Pasivo'),
    ('2404_IMPUESTO_RENTA', 'Impuesto de renta por pagar', 'Pasivo'),
    ('3105_CAPITAL', 'Capital suscrito y pagado', 'Patrimonio'),
    ('3705_UTILIDADES_ACUM', 'Utilidades acumuladas', 'Patrimonio'),
    ('3605_UTILIDAD_EJERCICIO', 'Utilidad del ejercicio', 'Patrimonio'),
]
title(li, 'Listas de miembros (ejemplo)', 'Reemplace por los IDs reales de las dimensiones de sus modelos en SAC. Alimentan las listas desplegables de las plantillas.', 12)
cols = [('Sociedades', SOCIEDADES), ('Cebes', CEBES), ('Cecos', CECOS), ('Ratio', RATIOS),
        ('Auditoria', AUDITORIA), ('Versión', VERSIONES)]
LIST_RANGES = {}
for j, (name, vals) in enumerate(cols):
    col = j + 1
    header_row(li, 4, [name], col)
    for i, v in enumerate(vals):
        c = li.cell(5 + i, col, v)
        c.font = F_IN
        c.border = BORDER
    LIST_RANGES[name] = f"Listas!${L(col)}$5:${L(col)}${4 + len(vals)}"
    li.column_dimensions[L(col)].width = 24
# Cuentas de egresos con grupo de P&L
header_row(li, 4, ['Cuentas_Egresos', 'Descripción', 'Grupo P&L (propiedad)'], 8)
for i, (a, b, c3) in enumerate(CUENTAS_EG):
    for k, v in enumerate((a, b, c3)):
        c = li.cell(5 + i, 8 + k, v)
        c.font = F_IN
        c.border = BORDER
EG_FIRST, EG_LAST = 5, 4 + len(CUENTAS_EG)
LIST_RANGES['Cuentas_Egresos'] = f'Listas!$H${EG_FIRST}:$H${EG_LAST}'
header_row(li, 4, ['Cuentas_EF', 'Descripción', 'Tipo'], 12)
for i, (a, b, c3) in enumerate(CUENTAS_EF):
    for k, v in enumerate((a, b, c3)):
        c = li.cell(5 + i, 12 + k, v)
        c.font = F_IN
        c.border = BORDER
for col, w in zip('HIJKLMN', (24, 30, 26, 3, 26, 34, 12)):
    li.column_dimensions[col].width = w

def dv_list(ws, name, ref):
    dv = DataValidation(type='list', formula1=LIST_RANGES[name], allow_blank=True,
                        showErrorMessage=True, errorTitle='Miembro no válido',
                        error=f'Elija un miembro de la lista {name} (hoja Listas).')
    ws.add_data_validation(dv)
    dv.add(ref)

# ----------------------------------------------------------------------------- Supuestos
su = wb.create_sheet('Supuestos', 1)
title(su, 'Supuestos de los reportes', 'Celdas amarillas con texto azul: valores que usted puede cambiar. Todo lo demás se calcula.', 6)
SUP = {}
rows = [
    ('Escala de los reportes (COP por unidad)', 1_000_000, '#,##0', 'ESCALA', 'Los reportes se muestran en millones de COP; las plantillas de carga van en COP.'),
    ('Costo de ventas (% de los ingresos)', 0.68, '0.0%', 'PCT_COSTO', 'Supuesto de ejemplo. En SAC suele venir del modelo de costos o de un data action.'),
    ('Tasa de impuesto de renta', 0.35, '0.0%', 'TASA_IMP', 'Tarifa general de renta en Colombia (35 %). Ajuste según su caso.'),
    ('Días de cartera (sobre ingresos)', 45, '0', 'DIAS_CXC', None),
    ('Días de inventario (sobre costo de ventas)', 60, '0', 'DIAS_INV', None),
    ('Días de proveedores (sobre costo de ventas)', 50, '0', 'DIAS_PROV', None),
    ('Inversión en PP&E mensual (millones COP)', 3_000, '#,##0', 'CAPEX', None),
    ('Vida útil de la PP&E (años)', 10, '0', 'VIDA', 'Depreciación en línea recta sobre el saldo bruto.'),
    ('Tasa de interés anual de la deuda', 0.12, '0.0%', 'TASA_INT', None),
    ('Amortización mensual de deuda (millones COP)', 1_000, '#,##0', 'AMORT', None),
]
header_row(su, 4, ['Supuesto', 'Valor', 'Nota'])
r = 5
for label, val, fmt, key, note in rows:
    su.cell(r, 1, label).font = f()
    c = su.cell(r, 2, val)
    c.font, c.fill, c.number_format, c.border = F_IN, FILL_IN, fmt, BORDER
    if note:
        su.cell(r, 3, note).font = f(italic=True, color=GREY, size=9)
    SUP[key] = f'Supuestos!$B${r}'
    r += 1
r += 1
su.cell(r, 1, 'Balance de apertura al 31 dic 2025 (millones COP)').font = f(True, BLUE)
r += 1
header_row(su, r, ['Cuenta', 'Saldo', 'Nota'])
r += 1
APERTURA = [
    ('Efectivo y equivalentes', 25_000, 'AP_CAJA'),
    ('Cuentas por cobrar clientes', 98_000, 'AP_CXC'),
    ('Inventarios', 76_000, 'AP_INV'),
    ('Propiedad, planta y equipo (bruto)', 420_000, 'AP_PPE'),
    ('Depreciación acumulada (negativo)', -150_000, 'AP_DEP'),
    ('Obligaciones financieras', 180_000, 'AP_DEUDA'),
    ('Proveedores', 63_000, 'AP_PROV'),
    ('Impuesto de renta por pagar', 0, 'AP_IMP'),
    ('Capital suscrito y pagado', 200_000, 'AP_CAPITAL'),
]
for label, val, key in APERTURA:
    su.cell(r, 1, label).font = f()
    c = su.cell(r, 2, val)
    c.font, c.fill, c.number_format, c.border = F_IN, FILL_IN, NUM, BORDER
    SUP[key] = f'Supuestos!$B${r}'
    r += 1
su.cell(r, 1, 'Utilidades acumuladas (cuadra la apertura)').font = f()
c = su.cell(r, 2, f"={SUP['AP_CAJA']}+{SUP['AP_CXC']}+{SUP['AP_INV']}+{SUP['AP_PPE']}+{SUP['AP_DEP']}"
                  f"-{SUP['AP_DEUDA']}-{SUP['AP_PROV']}-{SUP['AP_IMP']}-{SUP['AP_CAPITAL']}")
c.number_format, c.border = NUM, BORDER
su.cell(r, 3, 'Fórmula: activos − pasivos − capital.').font = f(italic=True, color=GREY, size=9)
SUP['AP_UTIL'] = f'Supuestos!$B${r}'
su.column_dimensions['A'].width = 46
su.column_dimensions['B'].width = 16
su.column_dimensions['C'].width = 70

# ----------------------------------------------------------------------------- Carga_Ingresos
ci = wb.create_sheet('Carga_Ingresos', 0)
title(ci, 'Plantilla de carga: Ingresos FANALCA (add-in de SAC para Excel)',
      'Filas = Sociedades × Cebes × Clientes × Referencias × Ratio. Columnas = Date (mes). Digite UNIDADES y PRECIO; el ingreso lo calcula el data action PxQ en SAC.', 20)
r0 = addin_header(ci, 'Modelo Ingresos FANALCA', 'Couh7ojg5e54rh2d4udijq6o83k',
                  [('Versión', 'public.Presupuesto_2026'), ('Auditoria', 'PRESUPUESTO_EXCEL'), ('Moneda', 'COP')])
CI_DIMS = ['Sociedades', 'Cebes', 'Clientes', 'Referencias', 'Ratio']
CI_M0 = len(CI_DIMS) + 1                      # primera columna de meses
header_row(ci, r0, CI_DIMS + PERIODOS + ['Total 2026'])
CI_HEAD = r0
seas = [0.82, 0.86, 0.97, 0.95, 1.00, 1.03, 0.98, 1.02, 1.05, 1.08, 1.06, 1.18]
LINEAS = [  # sociedad, cebes, cliente, referencia, unidades base, precio base (COP)
    ('FN_MOTOS', 'CB_MOTOS', 'CLI_001', 'REF_CB125', 2000, 9_890_000),
    ('FN_MOTOS', 'CB_MOTOS', 'CLI_003', 'REF_XR190', 900, 14_500_000),
    ('FN_AUTOS', 'CB_AUTOS', 'CLI_004', 'REF_CIVIC', 60, 145_000_000),
    ('FN_AUTOS', 'CB_AUTOS', 'CLI_005', 'REF_HRV', 80, 135_000_000),
    ('FN_TUBOS', 'CB_TUBERIA', 'CLI_002', 'REF_TUBO_2P', 400_000, 18_000),
    ('FN_TUBOS', 'CB_TUBERIA', 'CLI_006', 'REF_TUBO_4P', 150_000, 42_000),
]
r = CI_HEAD + 1
CI_FIRST = r
for soc, ceb, cli, ref, u0, p0 in LINEAS:
    ru, rp, ri = r, r + 1, r + 2
    for rr, ratio in ((ru, 'UNIDADES'), (rp, 'PRECIO'), (ri, 'ING_VENTAS_NAL')):
        for k, v in enumerate((soc, ceb, cli, ref, ratio)):
            c = ci.cell(rr, k + 1, v)
            c.fill, c.border = FILL_DIM, BORDER
            c.font = f(True) if k == 0 else f()
    for m in range(12):
        col = CI_M0 + m
        u = round(u0 * seas[m])
        p = round(p0 * (1.04 if m >= 6 else 1.0), -3 if p0 > 100_000 else 0)
        for rr, v in ((ru, u), (rp, p)):
            c = ci.cell(rr, col, v)
            c.font, c.fill, c.number_format, c.border = F_IN, FILL_IN, NUM_COP, BORDER
        c = ci.cell(ri, col, f'={L(col)}{ru}*{L(col)}{rp}')
        c.font, c.number_format, c.border = f(italic=True, color=GREY), NUM_COP, BORDER
    tcol = CI_M0 + 12
    first, last = L(CI_M0), L(CI_M0 + 11)
    ci.cell(ru, tcol, f'=SUM({first}{ru}:{last}{ru})').number_format = NUM_COP
    ci.cell(rp, tcol, f'=IFERROR(SUMPRODUCT({first}{ru}:{last}{ru},{first}{rp}:{last}{rp})/{L(tcol)}{ru},0)').number_format = NUM_COP
    ci.cell(ri, tcol, f'=SUM({first}{ri}:{last}{ri})').number_format = NUM_COP
    for rr in (ru, rp, ri):
        ci.cell(rr, tcol).font = f(True)
        ci.cell(rr, tcol).border = BORDER
    r += 3
CI_LAST = r - 1
ci.cell(r + 1, 1, 'Notas').font = f(True, BLUE)
notas_ci = [
    'Celdas amarillas (texto azul): valores para digitar o pegar en la tabla del add-in. Columnas grises: miembros de las dimensiones.',
    'Las filas ING_VENTAS_NAL (gris cursiva) muestran el PxQ como referencia: no las publique, el data action PxQ las calcula en SAC.',
    '"Total 2026" lo calcula SAC; en el add-in no se escribe sobre totales. PRECIO total = precio promedio ponderado.',
    'Si su tabla del add-in tiene otro orden de filas o columnas, pegue con Pegado especial → Valores respetando ese orden.',
]
for i, t in enumerate(notas_ci):
    ci.cell(r + 2 + i, 1, f'{i + 1}. {t}').font = f(size=9, color=GREY)
for k, w in enumerate((13, 15, 11, 15, 17)):
    ci.column_dimensions[L(k + 1)].width = w
for m in range(13):
    ci.column_dimensions[L(CI_M0 + m)].width = 15
ci.column_dimensions[L(CI_M0 + 12)].width = 18
ci.freeze_panes = ci.cell(CI_HEAD + 1, CI_M0)
dv_list(ci, 'Sociedades', f'A{CI_FIRST}:A{CI_LAST}')
dv_list(ci, 'Cebes', f'B{CI_FIRST}:B{CI_LAST}')
dv_list(ci, 'Ratio', f'E{CI_FIRST}:E{CI_LAST}')
dv_list(ci, 'Versión', 'B5')
dv_list(ci, 'Auditoria', 'B6')

# ----------------------------------------------------------------------------- Carga_Gastos
cg = wb.create_sheet('Carga_Gastos', 1)
title(cg, 'Plantilla de carga: Gastos FANALCA (add-in de SAC para Excel)',
      'Filas = Sociedades × Cecos × Cebes × Cuentas_Egresos. Columnas = Date (mes). Valores en COP.', 20)
r0 = addin_header(cg, 'Modelo Gastos FANALCA', 'Cunejii8r2d7vr4ldq2ousgfg7f',
                  [('Versión', 'public.Presupuesto_2026'), ('Auditoria', 'PRESUPUESTO_EXCEL'), ('Moneda', 'COP')])
CG_DIMS = ['Sociedades', 'Cecos', 'Cebes', 'Cuentas_Egresos']
CG_M0 = len(CG_DIMS) + 1
header_row(cg, r0, CG_DIMS + PERIODOS + ['Total 2026', 'Grupo P&L (ayuda, no se publica)'])
CG_HEAD = r0
GASTOS = [
    ('FN_MOTOS', 'CC_PLANTA', 'CB_MOTOS', '5105_SALARIOS', 2_400_000_000, 0.0),
    ('FN_MOTOS', 'CC_PLANTA', 'CB_MOTOS', '5145_MANTENIMIENTO', 380_000_000, 0.02),
    ('FN_MOTOS', 'CC_VENTAS', 'CB_MOTOS', '5205_SALARIOS_VENTAS', 950_000_000, 0.0),
    ('FN_MOTOS', 'CC_VENTAS', 'CB_MOTOS', '5235_PUBLICIDAD', 620_000_000, 0.08),
    ('FN_MOTOS', 'CC_LOGISTICA', 'CB_MOTOS', '5250_FLETES', 540_000_000, 0.03),
    ('FN_AUTOS', 'CC_ADMIN', 'CB_AUTOS', '5105_SALARIOS', 1_150_000_000, 0.0),
    ('FN_AUTOS', 'CC_ADMIN', 'CB_AUTOS', '5110_HONORARIOS', 210_000_000, 0.0),
    ('FN_AUTOS', 'CC_VENTAS', 'CB_AUTOS', '5205_SALARIOS_VENTAS', 620_000_000, 0.0),
    ('FN_AUTOS', 'CC_VENTAS', 'CB_AUTOS', '5235_PUBLICIDAD', 480_000_000, 0.10),
    ('FN_TUBOS', 'CC_PLANTA', 'CB_TUBERIA', '5105_SALARIOS', 1_300_000_000, 0.0),
    ('FN_TUBOS', 'CC_PLANTA', 'CB_TUBERIA', '5135_SERVICIOS', 430_000_000, 0.04),
    ('FN_TUBOS', 'CC_PLANTA', 'CB_TUBERIA', '5145_MANTENIMIENTO', 260_000_000, 0.02),
    ('FN_TUBOS', 'CC_LOGISTICA', 'CB_TUBERIA', '5250_FLETES', 410_000_000, 0.03),
    ('FN_TUBOS', 'CC_ADMIN', 'CB_CORPORATIVO', '5110_HONORARIOS', 150_000_000, 0.0),
]
r = CG_HEAD + 1
CG_FIRST = r
TCOL_G = CG_M0 + 12
GRP_COL = TCOL_G + 1
for soc, cc, ceb, cta, base, var in GASTOS:
    for k, v in enumerate((soc, cc, ceb, cta)):
        c = cg.cell(r, k + 1, v)
        c.fill, c.border = FILL_DIM, BORDER
        c.font = f(True) if k == 0 else f()
    for m in range(12):
        factor = 1 + var * ((m % 3) - 1) + (0.05 if (cta.startswith('5105') or cta.startswith('5205')) and m >= 6 else 0)
        if cta.startswith('5105') and m == 11:
            factor += 0.5  # prima de diciembre
        c = cg.cell(r, CG_M0 + m, round(base * factor, -5))
        c.font, c.fill, c.number_format, c.border = F_IN, FILL_IN, NUM_COP, BORDER
    c = cg.cell(r, TCOL_G, f'=SUM({L(CG_M0)}{r}:{L(CG_M0 + 11)}{r})')
    c.font, c.number_format, c.border = f(True), NUM_COP, BORDER
    c = cg.cell(r, GRP_COL, f'=IFERROR(INDEX(Listas!$J${EG_FIRST}:$J${EG_LAST},MATCH(D{r},Listas!$H${EG_FIRST}:$H${EG_LAST},0)),"Cuenta no encontrada")')
    c.font, c.border = f(italic=True, color=GREY), BORDER
    r += 1
CG_LAST = r - 1
c = cg.cell(r, 1, 'Total gastos')
c.font = f(True)
for col in range(CG_M0, TCOL_G + 1):
    c = cg.cell(r, col, f'=SUM({L(col)}{CG_FIRST}:{L(col)}{CG_LAST})')
    c.font, c.number_format, c.fill, c.border = f(True), NUM_COP, FILL_TOT, DBL
cg.cell(r + 2, 1, 'Notas').font = f(True, BLUE)
for i, t in enumerate([
    'La columna "Grupo P&L" sale de la propiedad de la cuenta (hoja Listas) y solo sirve para armar el P&L de este libro; no existe en la tabla del add-in.',
    'La fila "Total gastos" es de control: en SAC los totales los calcula el modelo.',
    'Diciembre incluye la prima de servicios en 5105_SALARIOS (+50 %).',
]):
    cg.cell(r + 3 + i, 1, f'{i + 1}. {t}').font = f(size=9, color=GREY)
for k, w in enumerate((13, 14, 16, 24)):
    cg.column_dimensions[L(k + 1)].width = w
for m in range(13):
    cg.column_dimensions[L(CG_M0 + m)].width = 15
cg.column_dimensions[L(TCOL_G)].width = 18
cg.column_dimensions[L(GRP_COL)].width = 26
cg.freeze_panes = cg.cell(CG_HEAD + 1, CG_M0)
dv_list(cg, 'Sociedades', f'A{CG_FIRST}:A{CG_LAST}')
dv_list(cg, 'Cecos', f'B{CG_FIRST}:B{CG_LAST}')
dv_list(cg, 'Cebes', f'C{CG_FIRST}:C{CG_LAST}')
dv_list(cg, 'Cuentas_Egresos', f'D{CG_FIRST}:D{CG_LAST}')
dv_list(cg, 'Versión', 'B5')
dv_list(cg, 'Auditoria', 'B6')

# ----------------------------------------------------------------------------- reportes (layout común)
# Columna A = código, B = concepto, C = saldo inicial (dic 2025), D..O = meses 2026, P = total
RC0 = 4  # D
RCOLS = [L(RC0 + m) for m in range(12)]
TOT = L(RC0 + 12)
ES = SUP['ESCALA']

def report_sheet(name, ttl, sub, idx, inicial=False):
    ws = wb.create_sheet(name, idx)
    title(ws, ttl, sub, 16)
    labels = ['Código', 'Concepto', 'Dic 2025' if inicial else ''] + PERIODOS + ['Total 2026']
    header_row(ws, 4, labels)
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 13 if inicial else 3
    for m in range(13):
        ws.column_dimensions[L(RC0 + m)].width = 12
    ws.column_dimensions[TOT].width = 14
    ws.freeze_panes = 'D5'
    return ws

def line(ws, r, code, label, formulas, total='sum', style=None, fmt=NUM, inicial=None):
    ws.cell(r, 1, code).font = f(color=GREY, size=9)
    lc = ws.cell(r, 2, label)
    lc.font = f(bold=style in ('sub', 'tot'))
    if style == 'sec':
        for col in range(1, RC0 + 13):
            ws.cell(r, col).fill = FILL_SEC
        lc.font = f(True, BLUE)
        return
    if inicial is not None:
        c = ws.cell(r, 3, inicial)
        c.number_format, c.font = fmt, F_LINK if 'Supuestos' in str(inicial) else F_CALC
    for m in range(12):
        c = ws.cell(r, RC0 + m, formulas(m))
        c.number_format = fmt
        txt = str(c.value)
        c.font = f(bold=style in ('sub', 'tot'), color='008000' if ('Carga_' in txt or ('!' in txt and 'Supuestos' not in txt)) else '1D2433')
    if total == 'sum':
        tv = f'=SUM({RCOLS[0]}{r}:{RCOLS[-1]}{r})'
    elif total == 'last':
        tv = f'={RCOLS[-1]}{r}'
    elif total == 'first':
        tv = f'={RCOLS[0]}{r}'
    else:
        tv = total
    if tv:
        c = ws.cell(r, RC0 + 12, tv)
        c.number_format = fmt
        c.font = f(True)
    if style in ('sub', 'tot'):
        for col in range(2, RC0 + 13):
            ws.cell(r, col).border = DBL if style == 'tot' else TOPLINE
            if style == 'tot':
                ws.cell(r, col).fill = FILL_TOT

# ----------------------------------------------------------------------------- P&L
pl = report_sheet('PyG', 'Estado de resultados (P&L) 2026 – FANALCA consolidado',
                  'Millones de COP. Ingresos y gastos vienen de las plantillas de carga; costo, depreciación, intereses e impuesto de la hoja Supuestos.', 3)
CIR = (CI_FIRST, CI_LAST)
CGR = (CG_FIRST, CG_LAST)

def ci_sum(m, soc):
    col = L(CI_M0 + m)
    return (f"=SUMIFS(Carga_Ingresos!{col}${CIR[0]}:{col}${CIR[1]},Carga_Ingresos!$A${CIR[0]}:$A${CIR[1]},\"{soc}\","
            f"Carga_Ingresos!$E${CIR[0]}:$E${CIR[1]},\"ING_VENTAS_NAL\")/{ES}")

def cg_sum(m, grupo):
    col = L(CG_M0 + m)
    g = L(GRP_COL)
    return (f"=-SUMIFS(Carga_Gastos!{col}${CGR[0]}:{col}${CGR[1]},Carga_Gastos!${g}${CGR[0]}:${g}${CGR[1]},\"{grupo}\")/{ES}")

P = {}
r = 5
line(pl, r, '', 'Ingresos operacionales', None, style='sec'); r += 1
for soc in SOCIEDADES:
    line(pl, r, soc, f'Ventas {soc.replace("FN_", "").title()}', lambda m, s=soc: ci_sum(m, s)); P[soc] = r; r += 1
line(pl, r, '41', 'Total ingresos', lambda m: f'=SUM({RCOLS[m]}{P[SOCIEDADES[0]]}:{RCOLS[m]}{P[SOCIEDADES[-1]]})', style='sub'); P['ING'] = r; r += 1
line(pl, r, '61', 'Costo de ventas', lambda m: f"=-{RCOLS[m]}{P['ING']}*{SUP['PCT_COSTO']}"); P['COSTO'] = r; r += 1
line(pl, r, '', 'Utilidad bruta', lambda m: f"={RCOLS[m]}{P['ING']}+{RCOLS[m]}{P['COSTO']}", style='sub'); P['UB'] = r; r += 1
line(pl, r, '', 'Margen bruto', lambda m: f"=IFERROR({RCOLS[m]}{P['UB']}/{RCOLS[m]}{P['ING']},0)", total=f"=IFERROR({TOT}{P['UB']}/{TOT}{P['ING']},0)", fmt=PCT); r += 1
r += 1
line(pl, r, '', 'Gastos operacionales', None, style='sec'); r += 1
line(pl, r, '51', 'Gastos de administración', lambda m: cg_sum(m, 'Gastos de administración')); P['ADM'] = r; r += 1
line(pl, r, '52', 'Gastos de ventas', lambda m: cg_sum(m, 'Gastos de ventas')); P['VTA'] = r; r += 1
line(pl, r, '', 'EBITDA', lambda m: f"={RCOLS[m]}{P['UB']}+{RCOLS[m]}{P['ADM']}+{RCOLS[m]}{P['VTA']}", style='sub'); P['EBITDA'] = r; r += 1
line(pl, r, '', 'Margen EBITDA', lambda m: f"=IFERROR({RCOLS[m]}{P['EBITDA']}/{RCOLS[m]}{P['ING']},0)", total=f"=IFERROR({TOT}{P['EBITDA']}/{TOT}{P['ING']},0)", fmt=PCT); r += 1
line(pl, r, '5160', 'Depreciación', lambda m: f"=-Balance!{RCOLS[m]}{{PPE}}/({SUP['VIDA']}*12)"); P['DEP'] = r; r += 1
line(pl, r, '', 'Utilidad operacional (EBIT)', lambda m: f"={RCOLS[m]}{P['EBITDA']}+{RCOLS[m]}{P['DEP']}", style='sub'); P['EBIT'] = r; r += 1
line(pl, r, '5305', 'Gastos financieros (intereses)', lambda m: f"=-{('Balance!C' if m == 0 else 'Balance!' + RCOLS[m - 1])}{{DEUDA}}*{SUP['TASA_INT']}/12"); P['INT'] = r; r += 1
line(pl, r, '', 'Utilidad antes de impuestos', lambda m: f"={RCOLS[m]}{P['EBIT']}+{RCOLS[m]}{P['INT']}", style='sub'); P['UAI'] = r; r += 1
line(pl, r, '5405', 'Impuesto de renta', lambda m: f"=-MAX(0,{RCOLS[m]}{P['UAI']})*{SUP['TASA_IMP']}"); P['IMP'] = r; r += 1
line(pl, r, '', 'Utilidad neta', lambda m: f"={RCOLS[m]}{P['UAI']}+{RCOLS[m]}{P['IMP']}", style='tot'); P['UN'] = r; r += 1
line(pl, r, '', 'Margen neto', lambda m: f"=IFERROR({RCOLS[m]}{P['UN']}/{RCOLS[m]}{P['ING']},0)", total=f"=IFERROR({TOT}{P['UN']}/{TOT}{P['ING']},0)", fmt=PCT); r += 1

# ----------------------------------------------------------------------------- Balance
ba = report_sheet('Balance', 'Balance general 2026 – FANALCA consolidado',
                  'Millones de COP, saldos al cierre de cada mes. La caja viene del flujo de caja; las demás cuentas de los supuestos de capital de trabajo, inversión y deuda.', 4, inicial=True)
B = {}
prev = lambda m: 'C' if m == 0 else RCOLS[m - 1]
r = 5
line(ba, r, '', 'Activo', None, style='sec'); r += 1
B['CAJA'] = r; line(ba, r, '1105_CAJA', 'Efectivo y equivalentes', lambda m: f"=Flujo_Caja!{RCOLS[m]}{{CAJA_FIN}}", total='last', inicial=f"={SUP['AP_CAJA']}"); r += 1
B['CXC'] = r; line(ba, r, '1305_CLIENTES', 'Cuentas por cobrar clientes', lambda m: f"=PyG!{RCOLS[m]}{P['ING']}*{SUP['DIAS_CXC']}/30", total='last', inicial=f"={SUP['AP_CXC']}"); r += 1
B['INV'] = r; line(ba, r, '1435_INVENTARIOS', 'Inventarios', lambda m: f"=-PyG!{RCOLS[m]}{P['COSTO']}*{SUP['DIAS_INV']}/30", total='last', inicial=f"={SUP['AP_INV']}"); r += 1
B['PPE'] = r; line(ba, r, '1520_PPE', 'Propiedad, planta y equipo (bruto)', lambda m: f"={prev(m)}{B['PPE']}+{SUP['CAPEX']}", total='last', inicial=f"={SUP['AP_PPE']}"); r += 1
B['DEP'] = r; line(ba, r, '1592_DEP_ACUMULADA', 'Depreciación acumulada', lambda m: f"={prev(m)}{B['DEP']}+PyG!{RCOLS[m]}{P['DEP']}", total='last', inicial=f"={SUP['AP_DEP']}"); r += 1
B['ACT'] = r; line(ba, r, '', 'Total activo', lambda m: f"=SUM({RCOLS[m]}{B['CAJA']}:{RCOLS[m]}{B['DEP']})", total='last', style='tot', inicial=f"=SUM(C{B['CAJA']}:C{B['DEP']})"); r += 1
r += 1
line(ba, r, '', 'Pasivo', None, style='sec'); r += 1
B['DEUDA'] = r; line(ba, r, '2105_OBLIGACIONES_FIN', 'Obligaciones financieras', lambda m: f"=MAX(0,{prev(m)}{B['DEUDA']}-{SUP['AMORT']})", total='last', inicial=f"={SUP['AP_DEUDA']}"); r += 1
B['PROV'] = r; line(ba, r, '2205_PROVEEDORES', 'Proveedores', lambda m: f"=-PyG!{RCOLS[m]}{P['COSTO']}*{SUP['DIAS_PROV']}/30", total='last', inicial=f"={SUP['AP_PROV']}"); r += 1
B['IMP'] = r; line(ba, r, '2404_IMPUESTO_RENTA', 'Impuesto de renta por pagar', lambda m: f"={prev(m)}{B['IMP']}-PyG!{RCOLS[m]}{P['IMP']}", total='last', inicial=f"={SUP['AP_IMP']}"); r += 1
B['PAS'] = r; line(ba, r, '', 'Total pasivo', lambda m: f"=SUM({RCOLS[m]}{B['DEUDA']}:{RCOLS[m]}{B['IMP']})", total='last', style='sub', inicial=f"=SUM(C{B['DEUDA']}:C{B['IMP']})"); r += 1
r += 1
line(ba, r, '', 'Patrimonio', None, style='sec'); r += 1
B['CAPITAL'] = r; line(ba, r, '3105_CAPITAL', 'Capital suscrito y pagado', lambda m: f"={prev(m)}{B['CAPITAL']}", total='last', inicial=f"={SUP['AP_CAPITAL']}"); r += 1
B['UACUM'] = r; line(ba, r, '3705_UTILIDADES_ACUM', 'Utilidades acumuladas', lambda m: f"={prev(m)}{B['UACUM']}", total='last', inicial=f"={SUP['AP_UTIL']}"); r += 1
B['UEJ'] = r; line(ba, r, '3605_UTILIDAD_EJERCICIO', 'Utilidad del ejercicio (acumulada en el año)', lambda m: f"=SUM(PyG!$D${P['UN']}:{RCOLS[m]}${P['UN']})", total='last', inicial=0); r += 1
B['PAT'] = r; line(ba, r, '', 'Total patrimonio', lambda m: f"=SUM({RCOLS[m]}{B['CAPITAL']}:{RCOLS[m]}{B['UEJ']})", total='last', style='sub', inicial=f"=SUM(C{B['CAPITAL']}:C{B['UEJ']})"); r += 1
B['PP'] = r; line(ba, r, '', 'Total pasivo y patrimonio', lambda m: f"={RCOLS[m]}{B['PAS']}+{RCOLS[m]}{B['PAT']}", total='last', style='tot', inicial=f"=C{B['PAS']}+C{B['PAT']}"); r += 1
r += 1
B['CHK'] = r; line(ba, r, '', 'Control: activo − pasivo − patrimonio (debe ser 0)', lambda m: f"=ROUND({RCOLS[m]}{B['ACT']}-{RCOLS[m]}{B['PP']},2)", total='last', inicial=f"=ROUND(C{B['ACT']}-C{B['PP']},2)"); r += 1
ba.cell(B['CHK'], 2).font = f(italic=True, color=GREY)
ba.conditional_formatting.add(f'C{B["CHK"]}:{TOT}{B["CHK"]}',
    CellIsRule(operator='notEqual', formula=['0'], fill=PatternFill('solid', fgColor='F8D7DA'), font=Font(name=FONT, color='9C0006', bold=True)))

# ----------------------------------------------------------------------------- Flujo de caja (método indirecto)
fc = report_sheet('Flujo_Caja', 'Flujo de caja 2026 (método indirecto) – FANALCA consolidado',
                  'Millones de COP. Parte de la utilidad neta del P&L y de las variaciones del balance; la caja final alimenta el balance.', 5)
FC = {}
var = lambda key, m, sign=1: (f"{'-' if sign < 0 else ''}(Balance!{RCOLS[m]}{B[key]}-Balance!{prev(m)}{B[key]})")
r = 5
line(fc, r, '', 'Actividades de operación', None, style='sec'); r += 1
FC['UN'] = r; line(fc, r, '', 'Utilidad neta', lambda m: f"=PyG!{RCOLS[m]}{P['UN']}"); r += 1
FC['DEP'] = r; line(fc, r, '', '(+) Depreciación', lambda m: f"=-PyG!{RCOLS[m]}{P['DEP']}"); r += 1
FC['CXC'] = r; line(fc, r, '', '(−) Aumento de cuentas por cobrar', lambda m: '=' + var('CXC', m, -1)); r += 1
FC['INV'] = r; line(fc, r, '', '(−) Aumento de inventarios', lambda m: '=' + var('INV', m, -1)); r += 1
FC['PROV'] = r; line(fc, r, '', '(+) Aumento de proveedores', lambda m: '=' + var('PROV', m)); r += 1
FC['IMP'] = r; line(fc, r, '', '(+) Aumento de impuestos por pagar', lambda m: '=' + var('IMP', m)); r += 1
FC['FCO'] = r; line(fc, r, '', 'Flujo de caja de operación', lambda m: f"=SUM({RCOLS[m]}{FC['UN']}:{RCOLS[m]}{FC['IMP']})", style='sub'); r += 1
r += 1
line(fc, r, '', 'Actividades de inversión', None, style='sec'); r += 1
FC['CAPEX'] = r; line(fc, r, '', '(−) Compras de propiedad, planta y equipo', lambda m: '=' + var('PPE', m, -1)); r += 1
FC['FCI'] = r; line(fc, r, '', 'Flujo de caja de inversión', lambda m: f"={RCOLS[m]}{FC['CAPEX']}", style='sub'); r += 1
r += 1
line(fc, r, '', 'Actividades de financiación', None, style='sec'); r += 1
FC['DEUDA'] = r; line(fc, r, '', '(+/−) Variación de obligaciones financieras', lambda m: '=' + var('DEUDA', m)); r += 1
FC['FCF'] = r; line(fc, r, '', 'Flujo de caja de financiación', lambda m: f"={RCOLS[m]}{FC['DEUDA']}", style='sub'); r += 1
r += 1
FC['NETO'] = r; line(fc, r, '', 'Variación neta del efectivo', lambda m: f"={RCOLS[m]}{FC['FCO']}+{RCOLS[m]}{FC['FCI']}+{RCOLS[m]}{FC['FCF']}", style='sub'); r += 1
FC['INI'] = r; line(fc, r, '', 'Efectivo al inicio del mes', lambda m: (f"={SUP['AP_CAJA']}" if m == 0 else f"={RCOLS[m - 1]}{{CAJA_FIN}}"), total='first'); r += 1
FC['FIN'] = r; line(fc, r, '', 'Efectivo al cierre del mes', lambda m: f"={RCOLS[m]}{FC['INI']}+{RCOLS[m]}{FC['NETO']}", total='last', style='tot'); r += 1

# Resolver referencias diferidas ({PPE}, {DEUDA}, {CAJA_FIN})
REPL = {'{PPE}': str(B['PPE']), '{DEUDA}': str(B['DEUDA']), '{CAJA_FIN}': str(FC['FIN'])}
for ws in (pl, ba, fc):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and '{' in c.value:
                v = c.value
                for k, val in REPL.items():
                    v = v.replace(k, val)
                c.value = v

# ----------------------------------------------------------------------------- Carga_EEFF (balance proyectado en formato del modelo)
ce = wb.create_sheet('Carga_EEFF', 2)
title(ce, 'Plantilla de carga: EEFF FANALCA (add-in de SAC para Excel)',
      'Balance proyectado en el formato del modelo EEFF: Cuentas_EF en filas, meses en columnas, COP. Pasivo y patrimonio con signo negativo (crédito), como en el modelo.', 20)
r0 = addin_header(ce, 'Modelo EEFF FANALCA', 'C2kgen7fqc13vleoibl06hk4j7o',
                  [('Versión', 'public.Presupuesto_2026'), ('Sociedades', 'FN_MOTOS'), ('Cebes', 'CB_CORPORATIVO'),
                   ('Auditoria', 'CALCULADO_DA'), ('Moneda', 'COP')])
header_row(ce, r0, ['Cuentas_EF', 'Descripción', 'Tipo'] + PERIODOS)
CE_M0 = 4
EF_MAP = [('1105_CAJA', 'CAJA', 1), ('1305_CLIENTES', 'CXC', 1), ('1435_INVENTARIOS', 'INV', 1), ('1520_PPE', 'PPE', 1),
          ('1592_DEP_ACUMULADA', 'DEP', 1), ('2105_OBLIGACIONES_FIN', 'DEUDA', -1), ('2205_PROVEEDORES', 'PROV', -1),
          ('2404_IMPUESTO_RENTA', 'IMP', -1), ('3105_CAPITAL', 'CAPITAL', -1), ('3705_UTILIDADES_ACUM', 'UACUM', -1),
          ('3605_UTILIDAD_EJERCICIO', 'UEJ', -1)]
desc = {a: (b, c3) for a, b, c3 in CUENTAS_EF}
r = r0 + 1
CE_FIRST = r
for cta, key, sign in EF_MAP:
    for k, v in enumerate((cta, desc[cta][0], desc[cta][1])):
        c = ce.cell(r, k + 1, v)
        c.fill, c.border = FILL_DIM, BORDER
    for m in range(12):
        c = ce.cell(r, CE_M0 + m, f"={'-' if sign < 0 else ''}Balance!{RCOLS[m]}{B[key]}*{ES}")
        c.font, c.number_format, c.border = F_LINK, NUM_COP, BORDER
    r += 1
CE_LAST = r - 1
c = ce.cell(r, 1, 'Control: suma debe ser 0')
c.font = f(True)
for m in range(12):
    col = L(CE_M0 + m)
    c = ce.cell(r, CE_M0 + m, f'=ROUND(SUM({col}{CE_FIRST}:{col}{CE_LAST}),0)')
    c.font, c.number_format, c.fill, c.border = f(True), NUM_COP, FILL_TOT, DBL
ce.conditional_formatting.add(f'{L(CE_M0)}{r}:{L(CE_M0 + 11)}{r}',
    CellIsRule(operator='notEqual', formula=['0'], fill=PatternFill('solid', fgColor='F8D7DA'), font=Font(name=FONT, color='9C0006', bold=True)))
ce.cell(r + 2, 1, 'Notas').font = f(True, BLUE)
for i, t in enumerate([
    'Los valores (texto verde) vienen de la hoja Balance: copie y pegue como valores en la tabla del add-in conectada al modelo EEFF.',
    'Ajuste Sociedades y Cebes del encabezado al miembro donde se registra el consolidado en su modelo.',
    'Si el modelo EEFF calcula la utilidad del ejercicio con un data action, no publique la fila 3605.',
]):
    ce.cell(r + 3 + i, 1, f'{i + 1}. {t}').font = f(size=9, color=GREY)
ce.column_dimensions['A'].width = 26
ce.column_dimensions['B'].width = 38
ce.column_dimensions['C'].width = 12
for m in range(12):
    ce.column_dimensions[L(CE_M0 + m)].width = 17
ce.freeze_panes = ce.cell(r0 + 1, CE_M0)
dv_list(ce, 'Versión', 'B5')

# ----------------------------------------------------------------------------- Carga_Plana
cp = wb.create_sheet('Carga_Plana', 3)
title(cp, 'Formato plano: botón de carga de archivos (data upload starter) o Cargador de datos a SAC',
      'Una fila por combinación y mes, con la medida en una sola columna. Los encabezados deben coincidir con los de la plantilla del trabajo de carga en SAC.', 10)
heads = ['Date', 'Sociedades', 'Cebes', 'Clientes', 'Referencias', 'Ratio', 'Auditoria', 'Moneda', 'Importe']
header_row(cp, 4, heads)
r = 5
CP_FIRST = r
for li_idx, (soc, ceb, cli, ref, u0, p0) in enumerate(LINEAS[:3]):
    base_row = CI_FIRST + li_idx * 3
    for m in range(2):
        for ratio, off in (('UNIDADES', 0), ('PRECIO', 1)):
            vals = [f'2026{m + 1:02d}', soc, ceb, cli, ref, ratio, 'PRESUPUESTO_EXCEL', 'COP']
            for k, v in enumerate(vals):
                c = cp.cell(r, k + 1, v)
                c.border = BORDER
                c.font = f()
            c = cp.cell(r, 9, f'=Carga_Ingresos!{L(CI_M0 + m)}{base_row + off}')
            c.font, c.number_format, c.border = F_LINK, NUM_COP, BORDER
            r += 1
CP_LAST = r - 1
cp.cell(r + 1, 1, 'Notas').font = f(True, BLUE)
for i, t in enumerate([
    'Date en formato AAAAMM como texto (202601). La versión no va en el archivo: se elige al cargar.',
    'Los importes (verde) se toman de Carga_Ingresos para el ejemplo; en un archivo real van como valores.',
    'El botón de carga de SAC exige las mismas columnas y encabezados que el archivo de muestra con que el modelador creó el trabajo de carga.',
    'Este mismo formato lo acepta el Cargador de datos a SAC (encabezados = nombres de las dimensiones).',
]):
    cp.cell(r + 2 + i, 1, f'{i + 1}. {t}').font = f(size=9, color=GREY)
for k, w in enumerate((10, 13, 14, 11, 15, 17, 20, 9, 18)):
    cp.column_dimensions[L(k + 1)].width = w
cp.freeze_panes = 'A5'
dv_list(cp, 'Sociedades', f'B{CP_FIRST}:B{CP_LAST}')
dv_list(cp, 'Ratio', f'F{CP_FIRST}:F{CP_LAST}')

# ----------------------------------------------------------------------------- Instrucciones
ins = wb.create_sheet('Instrucciones', 0)
title(ins, 'Plantillas de carga para el add-in de SAC y reportes financieros – FANALCA',
      'Libro de ejemplo. Los miembros, IDs de modelo y cifras son ilustrativos: reemplácelos por los de sus modelos.', 4)
ins.column_dimensions['A'].width = 24
ins.column_dimensions['B'].width = 110
row = 4
def sec(t):
    global row
    row += 1
    ins.cell(row, 1, t).font = f(True, BLUE, 12)
    row += 1
def item(a, b, color=None):
    global row
    c = ins.cell(row, 1, a)
    c.font = f(True)
    c.alignment = Alignment(vertical='top')
    d = ins.cell(row, 2, b)
    d.font = f()
    d.alignment = Alignment(wrap_text=True, vertical='top')
    if color:
        c.fill = color
    row += 1

sec('Contenido del libro')
for a, b in [
    ('Carga_Ingresos', 'Plantilla del modelo Ingresos FANALCA: UNIDADES y PRECIO por sociedad, Cebe, cliente y referencia, mes a mes.'),
    ('Carga_Gastos', 'Plantilla del modelo Gastos FANALCA: gasto por sociedad, centro de costo, Cebe y cuenta.'),
    ('Carga_EEFF', 'Balance proyectado en el formato del modelo EEFF FANALCA, listo para publicar.'),
    ('Carga_Plana', 'Formato plano (una fila por mes) para el botón de carga de archivos de una historia o el Cargador de datos a SAC.'),
    ('PyG', 'Estado de resultados mensual 2026 con márgenes.'),
    ('Balance', 'Balance general mensual con control de cuadre.'),
    ('Flujo_Caja', 'Flujo de caja por el método indirecto, conciliado con la caja del balance.'),
    ('Supuestos', 'Parámetros de costo, capital de trabajo, inversión, deuda, impuestos y balance de apertura.'),
    ('Listas', 'Miembros de las dimensiones que alimentan las listas desplegables.'),
]:
    item(a, b)

sec('Convenciones de color')
c = ins.cell(row, 1, '123.456'); c.font, c.fill = F_IN, FILL_IN
ins.cell(row, 2, 'Amarillo con texto azul: valor que usted digita o pega en la tabla del add-in.').font = f(); row += 1
c = ins.cell(row, 1, 'FN_MOTOS'); c.fill, c.font = FILL_DIM, f(True)
ins.cell(row, 2, 'Gris: miembros de las dimensiones (filas de la tabla del add-in).').font = f(); row += 1
c = ins.cell(row, 1, 'Vínculo'); c.font = F_LINK
ins.cell(row, 2, 'Texto verde: valor que viene de otra hoja del libro.').font = f(); row += 1
c = ins.cell(row, 1, 'Fórmula'); c.font = F_CALC
ins.cell(row, 2, 'Texto negro: fórmula. Gris cursiva: cálculo de referencia que SAC hace por su cuenta (no se publica).').font = f(); row += 1

sec('Cargar datos con el add-in de SAC para Excel')
for a, b in [
    ('1. Conectar', 'En Excel: pestaña SAP Analytics Cloud → Iniciar sesión con la URL de su tenant (por ejemplo https://fanalca.us10.hcs.cloud.sap).'),
    ('2. Insertar tabla', 'Insertar → Tabla → elija el modelo (Ingresos, Gastos o EEFF FANALCA). Ponga en filas las dimensiones de la plantilla y en columnas Date (meses de 2026).'),
    ('3. Filtrar', 'En el panel del diseñador fije la versión (por ejemplo public.Presupuesto_2026), Auditoria y Moneda como en el encabezado de la plantilla.'),
    ('4. Habilitar planificación', 'Menú contextual del modelo en el panel → Planificación habilitada. Sin esto las celdas no aceptan valores.'),
    ('5. Pegar valores', 'Copie las celdas amarillas y péguelas con Pegado especial → Valores en las mismas posiciones de la tabla. Si una combinación no existe aún, agréguela con "Agregar miembro" en la tabla (datos no reservados).'),
    ('6. Publicar', 'SAP Analytics Cloud → Publicar datos (o Guardar datos en la versión). Revise que no queden celdas rechazadas por bloqueos o control de acceso.'),
    ('7. Calcular', 'Si el modelo tiene data actions (por ejemplo el PxQ de ingresos), ejecútelas desde la historia o la multi action para calcular los derivados.'),
]:
    item(a, b)

sec('Permisos que necesita el usuario (licencia Planning Standard)')
for a, b in [
    ('Rol', 'Copia de Planner Reporter. Planning Model: Leer + Mantener + Ejecutar. No marcar Crear, Actualizar ni Eliminar (son de Planning Professional).'),
    ('Add-in Workbook', 'Leer para usar libros guardados; Crear + Leer + Actualizar para crear y guardar libros en el repositorio de SAC.'),
    ('Modelo', 'Compartido con el equipo con acceso de edición (Leer, Actualizar, Mantener).'),
    ('Datos', 'Permiso de escritura en el control de acceso a datos de Sociedades / Cebes, celdas sin bloqueo y versión editable.'),
    ('Carga de archivos', 'Para usar Carga_Plana con el botón de carga de una historia: Personal Data Acquisition = Ejecutar y un trabajo de carga creado por el modelador.'),
]:
    item(a, b)

sec('Cómo se conectan los reportes')
for a, b in [
    ('PyG', 'Ingresos = suma de ING_VENTAS_NAL (unidades × precio) por sociedad. Gastos por grupo de cuenta (51 administración, 52 ventas). Costo, depreciación, intereses e impuesto con los supuestos.'),
    ('Balance', 'Cartera, inventarios y proveedores por días de rotación; PP&E con la inversión mensual; deuda con la amortización; patrimonio con la utilidad del P&L. La caja sale del flujo.'),
    ('Flujo_Caja', 'Utilidad neta + depreciación − variaciones de capital de trabajo − inversión ± deuda. La caja final debe igualar la del balance; la fila de control del balance debe quedar en 0.'),
]:
    item(a, b)

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {'Instrucciones': NAVY, 'Supuestos': 'C9A227', 'Listas': '8A94A6'}.get(
        ws.title, CYAN if ws.title.startswith('Carga') else BLUE)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

order = ['Instrucciones', 'Carga_Ingresos', 'Carga_Gastos', 'Carga_EEFF', 'Carga_Plana', 'PyG', 'Balance', 'Flujo_Caja', 'Supuestos', 'Listas']
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(OUT)
print('OK', OUT)
