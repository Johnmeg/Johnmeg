# -*- coding: utf-8 -*-
"""
Generador de presentación profesional:
SAP Analytics Cloud enfocado en Planning.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
import copy

# ----------------------------------------------------------------------------
# PALETA Y TIPOGRAFÍA
# ----------------------------------------------------------------------------
NAVY      = RGBColor(0x0A, 0x2A, 0x43)   # Azul profundo (fondo oscuro)
BLUE      = RGBColor(0x00, 0x6E, 0xC7)   # Azul SAP
TEAL      = RGBColor(0x0F, 0xA9, 0xB0)   # Verde-azulado acento
GOLD      = RGBColor(0xF5, 0xB2, 0x41)   # Ámbar acento
LIGHT     = RGBColor(0xF4, 0xF7, 0xFA)   # Fondo claro
GREY      = RGBColor(0x5B, 0x6B, 0x7B)   # Texto secundario
DARKTXT   = RGBColor(0x13, 0x2A, 0x3E)   # Texto principal
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CARD      = RGBColor(0xFF, 0xFF, 0xFF)
SOFT      = RGBColor(0xE6, 0xED, 0xF3)   # Bordes suaves

FONT_H = "Segoe UI Semibold"
FONT_T = "Segoe UI"
FONT_L = "Segoe UI Light"

# 16:9
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ----------------------------------------------------------------------------
# HELPERS
# ----------------------------------------------------------------------------
def slide():
    return prs.slides.add_slide(BLANK)

def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None, line_w=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp

def no_fill_rect(s, x, y, w, h, line, line_w=Pt(1.25), shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.background()
    sp.line.color.rgb = line
    sp.line.width = line_w
    sp.shadow.inherit = False
    return sp

def soft_shadow(sp):
    """Sombra sutil para tarjetas."""
    spPr = sp._element.spPr
    el = spPr.makeelement(qn('a:effectLst'), {})
    sh = el.makeelement(qn('a:outerShdw'),
        {'blurRad':'90000','dist':'40000','dir':'5400000','rotWithShape':'0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val':'0A2A43'})
    alpha = clr.makeelement(qn('a:alpha'), {'val':'22000'})
    clr.append(alpha); sh.append(clr); el.append(sh); spPr.append(el)

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6, line_spacing=1.0, wrap=True):
    """runs: lista de párrafos; cada párrafo lista de (texto, size, color, font, bold)"""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, sz, col, fn, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col
            r.font.name = fn; r.font.bold = bold
    return tb

def one(t, sz, col, fn=FONT_T, bold=False):
    return [(t, sz, col, fn, bold)]

def icon_circle(s, x, y, d, color, glyph, gsize=18, gcolor=WHITE):
    c = rect(s, x, y, d, d, color, shape=MSO_SHAPE.OVAL)
    txt(s, x, y, d, d, [one(glyph, gsize, gcolor, FONT_H, True)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return c

def page_header(s, kicker, title, dark=False):
    tc = WHITE if dark else DARKTXT
    kc = TEAL if dark else BLUE
    rect(s, Inches(0.55), Inches(0.62), Inches(0.14), Inches(0.62), GOLD)
    txt(s, Inches(0.85), Inches(0.5), Inches(11.6), Inches(0.4),
        [one(kicker.upper(), 12.5, kc, FONT_H, True)], space_after=0)
    txt(s, Inches(0.83), Inches(0.82), Inches(11.8), Inches(0.7),
        [one(title, 30, tc, FONT_H, True)], space_after=0)

def footer(s, n, dark=False):
    col = RGBColor(0x8A,0x9C,0xAD) if not dark else RGBColor(0x5E,0x7A,0x92)
    txt(s, Inches(0.85), Inches(7.02), Inches(8), Inches(0.3),
        [one("SAP Analytics Cloud  ·  Planning", 9, col, FONT_T, False)], space_after=0)
    txt(s, Inches(11.4), Inches(7.02), Inches(1.1), Inches(0.3),
        [one(f"{n:02d}", 9, col, FONT_H, True)], align=PP_ALIGN.RIGHT, space_after=0)

_page = {"n": 0}
def pnum():
    _page["n"] += 1
    return _page["n"]


# ============================================================================
# SLIDE 1 — PORTADA
# ============================================================================
s = slide(); bg(s, NAVY)
# banda diagonal decorativa
band = rect(s, Inches(-1), Inches(4.7), Inches(16), Inches(4), BLUE, shape=MSO_SHAPE.PARALLELOGRAM)
band.rotation = 0
band.fill.fore_color.rgb = RGBColor(0x0C,0x3A,0x5C)
# círculos acento
rect(s, Inches(10.4), Inches(-1.3), Inches(3.6), Inches(3.6), RGBColor(0x0F,0x3E,0x60), shape=MSO_SHAPE.OVAL)
rect(s, Inches(11.6), Inches(4.9), Inches(2.4), Inches(2.4), RGBColor(0x11,0x4A,0x70), shape=MSO_SHAPE.OVAL)
rect(s, Inches(0.9), Inches(1.05), Inches(0.16), Inches(0.9), GOLD)

txt(s, Inches(1.25), Inches(1.0), Inches(9), Inches(0.5),
    [one("SOLUCIÓN DE PLANIFICACIÓN EMPRESARIAL", 13, TEAL, FONT_H, True)], space_after=0)
txt(s, Inches(1.2), Inches(1.7), Inches(10.7), Inches(2.2),
    [ one("SAP Analytics Cloud", 54, WHITE, FONT_H, True),
      one("Enfoque en Planning", 40, GOLD, FONT_L, False) ],
    space_after=6, line_spacing=1.02)
txt(s, Inches(1.24), Inches(4.05), Inches(9.6), Inches(1.0),
    [one("Planificación conectada, colaborativa e inteligente para Finanzas, "
         "Ventas y Recursos Humanos en una única plataforma en la nube.",
         16, RGBColor(0xC7,0xD6,0xE3), FONT_T, False)],
    line_spacing=1.25, space_after=0)

# chips inferiores
chips = ["Planificación", "Análisis", "Predicción"]
cx = Inches(1.24)
for c in chips:
    w = Inches(2.0)
    ch = no_fill_rect(s, cx, Inches(5.7), w, Inches(0.52), TEAL, line_w=Pt(1.25))
    txt(s, cx, Inches(5.7), w, Inches(0.52), [one(c, 13, WHITE, FONT_H, True)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cx += Inches(2.2)

txt(s, Inches(1.24), Inches(6.65), Inches(9), Inches(0.4),
    [one("Presentación profesional  ·  2026", 11, RGBColor(0x8A,0x9C,0xAD), FONT_T, False)],
    space_after=0)


# ============================================================================
# SLIDE 2 — AGENDA
# ============================================================================
s = slide(); bg(s, LIGHT)
rect(s, 0, 0, Inches(4.5), SH, NAVY)
rect(s, Inches(4.5), 0, Inches(0.06), SH, GOLD)
txt(s, Inches(0.6), Inches(0.9), Inches(0.16), Inches(0.9), [], )
rect(s, Inches(0.6), Inches(0.95), Inches(0.14), Inches(0.7), GOLD)
txt(s, Inches(0.85), Inches(0.85), Inches(3.4), Inches(0.5),
    [one("CONTENIDO", 13, TEAL, FONT_H, True)], space_after=0)
txt(s, Inches(0.83), Inches(1.25), Inches(3.5), Inches(1.4),
    [one("Agenda de la sesión", 30, WHITE, FONT_H, True)], line_spacing=1.02, space_after=0)
txt(s, Inches(0.85), Inches(5.9), Inches(3.4), Inches(1),
    [one("Un recorrido desde la visión general de la plataforma hasta las "
         "capacidades específicas de Planning.", 12, RGBColor(0xB9,0xCB,0xDA), FONT_T, False)],
    line_spacing=1.25, space_after=0)

agenda = [
    ("01", "¿Qué es SAP Analytics Cloud?", "Plataforma y propuesta de valor"),
    ("02", "Los tres pilares de la plataforma", "BI, Planning y Predictive"),
    ("03", "¿Qué es SAC Planning?", "Concepto y alcance"),
    ("04", "Capacidades clave de Planning", "Funcionalidades diferenciadoras"),
    ("05", "El ciclo de planificación", "Del dato a la decisión"),
    ("06", "Casos de uso y beneficios", "FP&A, ventas y RR. HH."),
]
ay = Inches(0.95)
for num, t, sub in agenda:
    card = rect(s, Inches(4.95), ay, Inches(7.9), Inches(0.88), CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    soft_shadow(card)
    rect(s, Inches(5.12), ay+Inches(0.14), Inches(0.6), Inches(0.6), BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, Inches(5.12), ay+Inches(0.14), Inches(0.6), Inches(0.6),
        [one(num, 16, WHITE, FONT_H, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(5.95), ay+Inches(0.13), Inches(6.7), Inches(0.4),
        [one(t, 16, DARKTXT, FONT_H, True)], space_after=0)
    txt(s, Inches(5.95), ay+Inches(0.5), Inches(6.7), Inches(0.3),
        [one(sub, 11.5, GREY, FONT_T, False)], space_after=0)
    ay += Inches(1.0)
footer(s, pnum())


# ============================================================================
# SLIDE 3 — ¿QUÉ ES SAC?
# ============================================================================
s = slide(); bg(s, LIGHT)
page_header(s, "Visión general", "¿Qué es SAP Analytics Cloud?")

txt(s, Inches(0.85), Inches(1.85), Inches(6.2), Inches(2.4),
    [ one("SAP Analytics Cloud (SAC) es la solución SaaS de análisis de SAP "
          "que reúne, en un único entorno en la nube, las capacidades de "
          "inteligencia de negocio, planificación empresarial y análisis "
          "predictivo.", 15.5, DARKTXT, FONT_T, False),
      [("", 6, DARKTXT, FONT_T, False)],
      one("Su propuesta central es eliminar los silos entre el reporting del "
          "pasado y la planificación del futuro: los mismos datos, modelos y "
          "usuarios trabajan sobre una sola plataforma de “confianza única”.",
          14, GREY, FONT_T, False) ],
    line_spacing=1.28, space_after=8)

# tarjeta lateral con datos clave
card = rect(s, Inches(7.55), Inches(1.85), Inches(5.0), Inches(4.55), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
soft_shadow(card)
rect(s, Inches(7.55), Inches(1.85), Inches(5.0), Inches(0.12), GOLD, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
txt(s, Inches(7.95), Inches(2.15), Inches(4.3), Inches(0.5),
    [one("EN POCAS PALABRAS", 12, TEAL, FONT_H, True)], space_after=0)
facts = [
    ("100% nube", "Solución SaaS, sin infraestructura que mantener."),
    ("Todo en uno", "Analítica, planificación y predicción integradas."),
    ("Colaborativa", "Múltiples usuarios y áreas sobre un único modelo."),
    ("Conectada", "Datos en vivo o importados de SAP y de terceros."),
]
fy = Inches(2.65)
for h, d in facts:
    icon_circle(s, Inches(7.95), fy, Inches(0.4), TEAL, "✔", 14)
    txt(s, Inches(8.5), fy-Inches(0.02), Inches(3.8), Inches(0.35),
        [one(h, 14.5, WHITE, FONT_H, True)], space_after=0)
    txt(s, Inches(8.5), fy+Inches(0.32), Inches(3.8), Inches(0.5),
        [one(d, 11.5, RGBColor(0xB9,0xCB,0xDA), FONT_T, False)], line_spacing=1.15, space_after=0)
    fy += Inches(0.95)

# banda inferior "una sola verdad"
b = rect(s, Inches(0.85), Inches(5.35), Inches(6.2), Inches(1.05), WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
soft_shadow(b)
rect(s, Inches(0.85), Inches(5.35), Inches(0.12), Inches(1.05), BLUE)
txt(s, Inches(1.15), Inches(5.5), Inches(5.7), Inches(0.8),
    [ one("Una única fuente de la verdad", 14.5, BLUE, FONT_H, True),
      one("Reporting, planes y previsiones comparten datos y definiciones.",
          12, GREY, FONT_T, False) ], line_spacing=1.15, space_after=3)
footer(s, pnum())


# ============================================================================
# SLIDE 4 — LOS TRES PILARES
# ============================================================================
s = slide(); bg(s, LIGHT)
page_header(s, "Arquitectura funcional", "Los tres pilares de la plataforma")
txt(s, Inches(0.85), Inches(1.6), Inches(11.6), Inches(0.5),
    [one("SAC integra tres disciplinas analíticas que tradicionalmente vivían en herramientas separadas.",
         14, GREY, FONT_T, False)], space_after=0)

pillars = [
    ("BI", "Business Intelligence", BLUE,
     ["Cuadros de mando y reporting", "Exploración visual de datos", "Historias interactivas"]),
    ("PL", "Planning", TEAL,
     ["Presupuestos y forecasts", "Escenarios y simulaciones", "Planificación colaborativa"]),
    ("AI", "Predictive & IA", GOLD,
     ["Smart Predict y forecast", "Insights automáticos", "Asistente 'Just Ask'"]),
]
px = Inches(0.85)
cw = Inches(3.83)
for i,(tag, name, col, items) in enumerate(pillars):
    highlight = (name == "Planning")
    top = Inches(2.35) if not highlight else Inches(2.15)
    hgt = Inches(4.05) if not highlight else Inches(4.35)
    base = NAVY if highlight else CARD
    card = rect(s, px, top, cw, hgt, base, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    soft_shadow(card)
    rect(s, px, top, cw, Inches(0.14), col, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    icon_circle(s, px+Inches(0.35), top+Inches(0.42), Inches(0.85), col, tag, 20)
    tcol = WHITE if highlight else DARKTXT
    scol = RGBColor(0xB9,0xCB,0xDA) if highlight else GREY
    txt(s, px+Inches(0.35), top+Inches(1.5), cw-Inches(0.7), Inches(0.55),
        [one(name, 20, tcol, FONT_H, True)], space_after=0)
    if highlight:
        txt(s, px+Inches(0.35), top+Inches(2.02), cw-Inches(0.7), Inches(0.35),
            [one("● NUESTRO FOCO", 11, GOLD, FONT_H, True)], space_after=0)
    iy = top + (Inches(2.5) if highlight else Inches(2.15))
    for it in items:
        rect(s, px+Inches(0.4), iy+Inches(0.08), Inches(0.13), Inches(0.13), col, shape=MSO_SHAPE.OVAL)
        txt(s, px+Inches(0.68), iy, cw-Inches(1.0), Inches(0.5),
            [one(it, 12.5, tcol if not highlight else RGBColor(0xDD,0xE7,0xF0), FONT_T, False)],
            line_spacing=1.1, space_after=0)
        iy += Inches(0.52)
    px += Inches(4.08)
footer(s, pnum())


# ============================================================================
# SLIDE 5 — ¿QUÉ ES SAC PLANNING?
# ============================================================================
s = slide(); bg(s, NAVY)
# panel derecho decorativo
rect(s, Inches(8.9), 0, Inches(4.43), SH, RGBColor(0x0C,0x3A,0x5C))
rect(s, Inches(8.9), 0, Inches(0.06), SH, GOLD)
page_header(s, "El foco de hoy", "¿Qué es SAP Analytics Cloud Planning?", dark=True)

txt(s, Inches(0.85), Inches(1.9), Inches(7.6), Inches(2.0),
    [ one("Es el módulo de planificación empresarial de SAC. Permite crear, "
          "gestionar y colaborar en presupuestos, previsiones y planes "
          "operativos, trabajando directamente sobre los datos analíticos, "
          "sin exportar a hojas de cálculo.", 15.5, RGBColor(0xDD,0xE7,0xF0), FONT_T, False) ],
    line_spacing=1.3, space_after=0)

txt(s, Inches(0.85), Inches(3.7), Inches(7.6), Inches(0.4),
    [one("LO QUE HACE DIFERENTE A SAC PLANNING", 12.5, TEAL, FONT_H, True)], space_after=0)

diff = [
    ("Datos y plan, unidos", "El plan se construye sobre el modelo analítico real."),
    ("Fin de las hojas sueltas", "Reemplaza los Excel dispersos por un modelo gobernado."),
    ("Colaboración en vivo", "Tareas, comentarios y flujos de aprobación integrados."),
    ("Del análisis a la acción", "Se analiza el pasado y se planifica el futuro en el mismo sitio."),
]
dy = Inches(4.15)
for h, d in diff:
    icon_circle(s, Inches(0.9), dy, Inches(0.42), TEAL, "→", 16)
    txt(s, Inches(1.5), dy-Inches(0.02), Inches(6.9), Inches(0.35),
        [one(h, 14.5, WHITE, FONT_H, True)], space_after=0)
    txt(s, Inches(1.5), dy+Inches(0.31), Inches(6.9), Inches(0.35),
        [one(d, 12, RGBColor(0xB9,0xCB,0xDA), FONT_T, False)], space_after=0)
    dy += Inches(0.72)

# panel derecho: cifras/proceso
txt(s, Inches(9.25), Inches(1.95), Inches(3.8), Inches(0.4),
    [one("EN LA PRÁCTICA", 12, GOLD, FONT_H, True)], space_after=0)
steps = [("Recopila", "datos reales y supuestos"),
         ("Modela", "escenarios y drivers"),
         ("Colabora", "entre áreas y responsables"),
         ("Decide", "con previsiones fiables")]
sy = Inches(2.5)
for i,(h,d) in enumerate(steps):
    icon_circle(s, Inches(9.25), sy, Inches(0.5), GOLD, str(i+1), 16, NAVY)
    txt(s, Inches(9.95), sy-Inches(0.03), Inches(3.1), Inches(0.35),
        [one(h, 15, WHITE, FONT_H, True)], space_after=0)
    txt(s, Inches(9.95), sy+Inches(0.32), Inches(3.1), Inches(0.4),
        [one(d, 11.5, RGBColor(0xB9,0xCB,0xDA), FONT_T, False)], line_spacing=1.1, space_after=0)
    if i < len(steps)-1:
        rect(s, Inches(9.48), sy+Inches(0.52), Inches(0.035), Inches(0.5), RGBColor(0x2A,0x5A,0x7E))
    sy += Inches(1.02)
footer(s, pnum(), dark=True)


# ============================================================================
# SLIDE 6 — CAPACIDADES CLAVE
# ============================================================================
s = slide(); bg(s, LIGHT)
page_header(s, "Funcionalidad", "Capacidades clave de Planning")
txt(s, Inches(0.85), Inches(1.6), Inches(11.6), Inches(0.5),
    [one("Un conjunto de herramientas diseñadas para automatizar y enriquecer el proceso de planificación.",
         14, GREY, FONT_T, False)], space_after=0)

caps = [
    ("▦", "Modelos multidimensionales",
     "Cuentas, tiempo, versiones y dimensiones de negocio en un único modelo.", BLUE),
    ("↻", "Data Actions",
     "Automatizan cálculos, copias de versiones y distribuciones complejas.", TEAL),
    ("⌥", "Value Driver Trees",
     "Simulación visual del impacto de los inductores clave del negocio.", GOLD),
    ("◑", "Versiones y escenarios",
     "Compara Actual, Budget y Forecast; simula hipótesis 'what-if'.", BLUE),
    ("☷", "Allocations",
     "Reparte costes e ingresos según reglas y criterios de asignación.", TEAL),
    ("⌘", "Predictive Planning",
     "Genera previsiones automáticas con machine learning integrado.", GOLD),
]
gx0, gy0 = Inches(0.85), Inches(2.35)
cw, chh = Inches(3.83), Inches(1.95)
gapx, gapy = Inches(0.25), Inches(0.22)
for i,(ic, h, d, col) in enumerate(caps):
    r, c = divmod(i, 3)
    x = gx0 + c*(cw+gapx)
    y = gy0 + r*(chh+gapy)
    card = rect(s, x, y, cw, chh, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    soft_shadow(card)
    rect(s, x, y, Inches(0.12), chh, col)
    icon_circle(s, x+Inches(0.35), y+Inches(0.32), Inches(0.62), col, ic, 20)
    txt(s, x+Inches(1.15), y+Inches(0.34), cw-Inches(1.35), Inches(0.55),
        [one(h, 15, DARKTXT, FONT_H, True)], line_spacing=1.0, space_after=0)
    txt(s, x+Inches(0.35), y+Inches(1.15), cw-Inches(0.65), Inches(0.7),
        [one(d, 12, GREY, FONT_T, False)], line_spacing=1.18, space_after=0)
footer(s, pnum())


# ============================================================================
# SLIDE 7 — EL CICLO DE PLANIFICACIÓN
# ============================================================================
s = slide(); bg(s, LIGHT)
page_header(s, "Proceso", "El ciclo de planificación en SAC")
txt(s, Inches(0.85), Inches(1.6), Inches(11.6), Inches(0.5),
    [one("Un flujo continuo y cerrado que conecta el análisis del pasado con la decisión sobre el futuro.",
         14, GREY, FONT_T, False)], space_after=0)

flow = [
    ("1", "Integrar", "Conectar datos reales de SAP y de terceros.", BLUE),
    ("2", "Modelar", "Definir dimensiones, versiones y reglas de cálculo.", TEAL),
    ("3", "Planificar", "Presupuestar y proyectar de forma colaborativa.", GOLD),
    ("4", "Simular", "Evaluar escenarios y análisis 'what-if'.", BLUE),
    ("5", "Analizar", "Comparar plan vs. real y medir desviaciones.", TEAL),
]
n = len(flow)
x0 = Inches(0.85)
total = Inches(11.63)
cw = Inches(2.05)
gap = (total - cw*n) / (n-1)
y = Inches(2.7)
for i,(num, h, d, col) in enumerate(flow):
    x = x0 + i*(cw+gap)
    card = rect(s, x, y, cw, Inches(2.9), CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    soft_shadow(card)
    rect(s, x, y, cw, Inches(0.7), col, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    icon_circle(s, x+cw/2-Inches(0.5), y+Inches(0.28), Inches(1.0), NAVY, num, 26, WHITE)
    txt(s, x+Inches(0.1), y+Inches(1.4), cw-Inches(0.2), Inches(0.45),
        [one(h, 16, DARKTXT, FONT_H, True)], align=PP_ALIGN.CENTER, space_after=0)
    txt(s, x+Inches(0.18), y+Inches(1.9), cw-Inches(0.36), Inches(0.9),
        [one(d, 11.5, GREY, FONT_T, False)], align=PP_ALIGN.CENTER, line_spacing=1.18, space_after=0)
    if i < n-1:
        ar = rect(s, x+cw+Inches(0.02), y+Inches(1.15), gap-Inches(0.04), Inches(0.4),
                  col, shape=MSO_SHAPE.CHEVRON)
# ciclo cerrado
band = rect(s, Inches(0.85), Inches(5.95), Inches(11.63), Inches(0.72), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, Inches(1.1), Inches(6.06), Inches(11.2), Inches(0.5),
    [[("↻  ", 15, GOLD, FONT_H, True),
      ("Proceso cíclico y continuo: ", 13.5, WHITE, FONT_H, True),
      ("los resultados del análisis retroalimentan el siguiente ciclo de planificación.",
       13.5, RGBColor(0xC7,0xD6,0xE3), FONT_T, False)]],
    anchor=MSO_ANCHOR.MIDDLE, space_after=0)
footer(s, pnum())


# ============================================================================
# SLIDE 8 — CASOS DE USO
# ============================================================================
s = slide(); bg(s, LIGHT)
page_header(s, "Aplicación", "Casos de uso de Planning")
txt(s, Inches(0.85), Inches(1.6), Inches(11.6), Inches(0.5),
    [one("La planificación se extiende a todas las áreas de la organización sobre un mismo modelo integrado.",
         14, GREY, FONT_T, False)], space_after=0)

cases = [
    ("Finanzas (FP&A)", BLUE, [
        "Presupuesto anual y rolling forecast",
        "Estados financieros integrados (P&L, balance, cash-flow)",
        "Consolidación y análisis de desviaciones"]),
    ("Ventas e Ingresos", TEAL, [
        "Planificación comercial y de cuotas",
        "Previsión de demanda por producto y región",
        "Simulación de precios y márgenes"]),
    ("Recursos Humanos", GOLD, [
        "Planificación de plantilla y costes laborales",
        "Escenarios de contratación y retribución",
        "Alineación de capacidad con la demanda"]),
]
cx = Inches(0.85)
cw = Inches(3.83)
for name, col, items in cases:
    card = rect(s, cx, Inches(2.35), cw, Inches(4.05), CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    soft_shadow(card)
    head = rect(s, cx, Inches(2.35), cw, Inches(0.95), col, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    txt(s, cx+Inches(0.35), Inches(2.35), cw-Inches(0.7), Inches(0.95),
        [one(name, 17, WHITE, FONT_H, True)], anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    iy = Inches(3.55)
    for it in items:
        rect(s, cx+Inches(0.38), iy+Inches(0.09), Inches(0.14), Inches(0.14), col, shape=MSO_SHAPE.OVAL)
        txt(s, cx+Inches(0.68), iy, cw-Inches(1.0), Inches(0.8),
            [one(it, 12.5, DARKTXT, FONT_T, False)], line_spacing=1.18, space_after=0)
        iy += Inches(0.92)
    cx += Inches(4.08)
footer(s, pnum())


# ============================================================================
# SLIDE 9 — BENEFICIOS
# ============================================================================
s = slide(); bg(s, LIGHT)
page_header(s, "Valor de negocio", "Beneficios clave")
txt(s, Inches(0.85), Inches(1.6), Inches(11.6), Inches(0.5),
    [one("Por qué las organizaciones eligen SAP Analytics Cloud para su planificación.",
         14, GREY, FONT_T, False)], space_after=0)

# métricas destacadas
metrics = [("1", "plataforma", "Analítica y planificación unificadas", BLUE),
           ("+", "colaboración", "Áreas y responsables trabajando juntos", TEAL),
           ("<>", "agilidad", "Del forecast anual al continuo", GOLD),
           ("AI", "inteligencia", "Predicciones con machine learning", NAVY)]
mx = Inches(0.85); mw = Inches(2.83)
for tag, big, d, col in metrics:
    card = rect(s, mx, Inches(2.35), mw, Inches(1.7), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    soft_shadow(card)
    txt(s, mx+Inches(0.3), Inches(2.5), mw-Inches(0.6), Inches(0.7),
        [[(tag+" ", 30, WHITE, FONT_H, True),(big, 16, RGBColor(0xDD,0xE7,0xF0), FONT_L, False)]],
        space_after=0)
    txt(s, mx+Inches(0.3), Inches(3.35), mw-Inches(0.6), Inches(0.6),
        [one(d, 11.5, RGBColor(0xE6,0xED,0xF3), FONT_T, False)], line_spacing=1.15, space_after=0)
    mx += Inches(3.0)

benefits = [
    ("Una sola fuente de la verdad", "Todos planifican sobre los mismos datos gobernados."),
    ("Ciclos más rápidos", "Automatización de cálculos y forecasts recurrentes."),
    ("Mejores decisiones", "Escenarios, simulaciones y predicción integrados."),
    ("Escalable en la nube", "Sin infraestructura; se adapta al crecimiento."),
]
by = Inches(4.5)
for i,(h,d) in enumerate(benefits):
    r,c = divmod(i,2)
    x = Inches(0.85) + c*Inches(5.95)
    y = by + r*Inches(0.98)
    icon_circle(s, x, y, Inches(0.5), TEAL, "✔", 16)
    txt(s, x+Inches(0.7), y-Inches(0.03), Inches(5.0), Inches(0.35),
        [one(h, 14.5, DARKTXT, FONT_H, True)], space_after=0)
    txt(s, x+Inches(0.7), y+Inches(0.33), Inches(5.0), Inches(0.4),
        [one(d, 11.5, GREY, FONT_T, False)], line_spacing=1.12, space_after=0)
footer(s, pnum())


# ============================================================================
# SLIDE 10 — CIERRE
# ============================================================================
s = slide(); bg(s, NAVY)
rect(s, Inches(-1), Inches(4.9), Inches(16), Inches(4), RGBColor(0x0C,0x3A,0x5C), shape=MSO_SHAPE.PARALLELOGRAM)
rect(s, Inches(10.6), Inches(-1.2), Inches(3.4), Inches(3.4), RGBColor(0x0F,0x3E,0x60), shape=MSO_SHAPE.OVAL)
rect(s, Inches(0.9), Inches(2.2), Inches(0.16), Inches(1.0), GOLD)

txt(s, Inches(1.25), Inches(2.15), Inches(9), Inches(0.5),
    [one("EN RESUMEN", 13, TEAL, FONT_H, True)], space_after=0)
txt(s, Inches(1.22), Inches(2.75), Inches(11), Inches(1.6),
    [ one("Planificar el futuro,", 40, WHITE, FONT_H, True),
      one("con los datos del presente.", 40, GOLD, FONT_L, False) ],
    line_spacing=1.03, space_after=2)
txt(s, Inches(1.25), Inches(4.7), Inches(10.3), Inches(1.0),
    [one("SAP Analytics Cloud Planning unifica análisis y planificación en una "
         "sola plataforma en la nube: colaborativa, inteligente y conectada.",
         16, RGBColor(0xC7,0xD6,0xE3), FONT_T, False)], line_spacing=1.3, space_after=0)

txt(s, Inches(1.25), Inches(6.4), Inches(9), Inches(0.4),
    [[("¿Preguntas?  ", 14, WHITE, FONT_H, True),
      ("Gracias por su atención.", 14, RGBColor(0x8A,0x9C,0xAD), FONT_T, False)]], space_after=0)


# ----------------------------------------------------------------------------
out = "/home/user/Johnmeg/SAP_Analytics_Cloud_Planning.pptx"
prs.save(out)
print("Guardado:", out, "| Slides:", len(prs.slides._sldIdLst))
