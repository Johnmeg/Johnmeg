from openpyxl import load_workbook
from copy import copy
import uuid, html, os
SP=os.path.dirname(os.path.abspath(__file__))

# --- leer casos SAC ---
s=load_workbook(f"{SP}/in/sac_script.xlsx")["Casos de Prueba"]
cases=[]
for r in range(3, s.max_row+1):
    idv=s.cell(r,1).value
    if not idv: continue
    cases.append(dict(id=idv, caso=s.cell(r,2).value, cat=s.cell(r,3).value, mod=s.cell(r,4).value,
                      pre=s.cell(r,5).value, pasos=s.cell(r,6).value, esp=s.cell(r,7).value, pri=s.cell(r,8).value))
print("casos SAC:", len(cases))

# --- abrir plantilla (formato 4HH-01MT-0003) y trabajar sobre ella ---
wb=load_workbook(f"{SP}/in/emergency.xlsx")
ws=wb["Test Cases"]
NCOL=21
# capturar estilos de la fila de datos (fila 6) ANTES de tocar nada
tpl={c:{'font':copy(ws.cell(6,c).font),'fill':copy(ws.cell(6,c).fill),
        'border':copy(ws.cell(6,c).border),'alignment':copy(ws.cell(6,c).alignment),
        'nf':ws.cell(6,c).number_format} for c in range(1,NCOL+1)}

# --- metadatos ---
ws["B3"]="9/8/2026, 10:00:00 AM"   # created at
SCOPE_GUID="55f82077-4266-43cd-9c10-6204c0ca8340"; SCOPE_NAME="Alcance Greenfield Fanalca"
PRI={"Alta":"High","Media":"Medium","Baja":"Low"}
def esc(x): return html.escape(str(x), quote=False)
def instr(pre,pasos):
    p=""
    if pre and str(pre).strip() not in ("—","-",""): p+=f"<p><strong>Precondiciones:</strong> {esc(pre)}</p>"
    return p+f"<p>{esc(pasos)}</p>"

# --- limpiar filas de datos existentes (6..max) ---
for r in range(6, ws.max_row+1):
    for c in range(1,NCOL+1): ws.cell(r,c).value=None

# --- escribir casos SAC en el formato ---
row=6
for cs in cases:
    v={1:str(uuid.uuid4()), 2:f"{cs['id']} {cs['caso']}", 3:SCOPE_GUID, 4:SCOPE_NAME,
       5:"",6:"",7:"",8:"",9:"",10:"", 11:PRI.get((cs['pri'] or '').strip(),"Medium"), 12:"", 13:"In Preparation",
       14:str(uuid.uuid4()), 15:f"{cs['mod']} · {cs['cat']}", 16:"",17:"",
       18:str(uuid.uuid4()), 19:"1 Ejecutar y verificar",
       20:instr(cs['pre'],cs['pasos']), 21:f"<p>{esc(cs['esp'])}</p>"}
    for c in range(1,NCOL+1):
        cell=ws.cell(row,c); cell.value=v[c]; st=tpl[c]
        cell.font=copy(st['font']); cell.fill=copy(st['fill']); cell.border=copy(st['border'])
        cell.alignment=copy(st['alignment']); cell.number_format=st['nf']
    if row in ws.row_dimensions: ws.row_dimensions[row].height=None
    row+=1

out=f"{SP}/out/Script_de_pruebas_SAC_(formato_4HH-01MT-0003).xlsx"
wb.save(out)
print("GUARDADO:", out, "| filas de datos:", row-6)
