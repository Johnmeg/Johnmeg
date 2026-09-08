# -*- coding: utf-8 -*-
"""Guion de pruebas (test script) para los modelos de SAP Analytics Cloud (SAC).
Sencillo y alineado con buenas prácticas de SAP. Genera un libro Excel con:
Instrucciones, Casos de Prueba, Resumen (tablero), Defectos y Listas."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList

NAVY="1F3864"; NAVY2="2E4C7E"; BLUE="2E75B6"; GREY="595959"; LIGHT="F2F6FC"
HFILL=PatternFill("solid", fgColor=NAVY)
WHITE=Font(name="Calibri", color="FFFFFF", bold=True, size=10)
THIN=Side(style="thin", color="BFBFBF")
BORDER=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
TITLE=Font(name="Calibri", bold=True, size=15, color="FFFFFF")
CEN=Alignment(horizontal="center", vertical="center", wrap_text=True)
LEF=Alignment(horizontal="left", vertical="center", wrap_text=True)
TOPL=Alignment(horizontal="left", vertical="top", wrap_text=True)
F9=Font(name="Calibri", size=9); F9B=Font(name="Calibri", size=9, bold=True)

ESTADOS={"Pendiente":("EDEDED","595959"),"Aprobado":("C6EFCE","006100"),
         "Fallido":("FFC7CE","9C0006"),"Bloqueado":("FFE699","7F6000"),
         "No aplica":("DDEBF7","1F3864")}
PRIOR={"Alta":("FCE4D6","833C00"),"Media":("FFF2CC","7F6000"),"Baja":("EDEDED","595959")}
CATS=["Datos maestros y dimensiones","Integración de datos","Cálculos (Data Actions)",
      "Versiones y escenarios","Entrada de datos","Seguridad y autorizaciones",
      "Reportes y dashboards","Conciliación","Rendimiento","Regresión / general"]
MODELOS=["Ingresos","Costos y Gastos","Estados Financieros","Todos"]

# (id, caso, categoria, modelo, precondiciones, pasos, esperado, prioridad)
C=CATS
CASES=[
("TC-001","Verificar dimensiones y jerarquías del modelo",C[0],"Todos","Modelo creado; maestros importados","Abrir el modelo y revisar las dimensiones (Cuenta, Sociedad, CEBE, CECO, Versión, Fecha) y sus jerarquías","Todas las dimensiones y jerarquías existen, sin duplicados ni nodos huérfanos","Alta"),
("TC-002","Validar miembros de dimensión contra la fuente",C[0],"Todos","Maestros disponibles en DataSphere","Comparar el conteo y los valores de miembros (regionales, CECOs, cuentas) con la fuente","Coinciden 100% con la fuente","Alta"),
("TC-003","Validar propiedades de la dimensión Cuenta",C[0],"Todos","—","Revisar tipo de cuenta (INC/EXP/AST/LEQ), agregación y signo de cada cuenta","Cada cuenta tiene tipo, agregación y signo correctos","Alta"),
("TC-004","Validar descripciones y atributos de miembros",C[0],"Todos","—","Revisar descripciones y atributos de cuentas y centros","Descripciones correctas y completas; sin miembros sin descripción","Media"),
("TC-010","Ejecutar importación del real desde DataSphere",C[1],"Todos","Conexión a DataSphere configurada","Ejecutar el flujo de importación del dato real para un periodo de prueba","La importación finaliza sin errores","Alta"),
("TC-011","Conciliar saldos reales SAC vs. ERP",C[1],"Estados Financieros","Real importado","Comparar totales por cuenta y CEBE contra el balance del ERP","Diferencia = 0","Alta"),
("TC-012","Verificar el mapeo de campos (origen → dimensión/medida)",C[1],"Todos","—","Revisar que cada campo origen mapea a la dimensión o medida correcta","Mapeo correcto; sin datos en «no asignado»","Alta"),
("TC-013","Probar el refresco programado del real",C[1],"Todos","Calendario de refresco configurado","Verificar la ejecución del refresco y su bitácora","Refresco exitoso según el calendario","Media"),
("TC-014","Manejo de errores de importación (datos inválidos)",C[1],"Todos","—","Cargar un registro con un miembro inexistente","El sistema rechaza el registro e informa el error con claridad","Media"),
("TC-020","Data Action de ingresos P×Q",C[2],"Ingresos","Kilos y tarifa cargados","Ejecutar la Data Action de P×Q sobre un juego de datos controlado","Importe = Kilos × Tarifa (coincide con el cálculo manual)","Alta"),
("TC-021","Data Action de crecimiento (volumen y tarifa)",C[2],"Ingresos","Base N-1 y % definidos","Ejecutar la Data Action de crecimiento","Resultado = base × (1 + %), por regional","Alta"),
("TC-022","Data Action de estacionalidad",C[2],"Ingresos","Total anual calculado","Ejecutar la Data Action de distribución mensual","La suma mensual = total anual (100%)","Alta"),
("TC-023","Data Action de costos (fletes / mano de obra)",C[2],"Costos y Gastos","Kilos proyectados","Ejecutar las Data Actions de costos","Costo = índice × kg (coincide con lo esperado)","Alta"),
("TC-024","Consolidación al EEFF (Cross-Model Copy)",C[2],"Estados Financieros","Ingresos y costos calculados","Ejecutar la Data Action de consolidación","El P&G por CEBE/regional cuadra con la suma de origen","Alta"),
("TC-025","Idempotencia / re-ejecución de Data Actions",C[2],"Todos","—","Ejecutar dos veces la misma Data Action","El resultado no se duplica (mismos valores)","Media"),
("TC-026","Cálculo sobre versión privada",C[2],"Todos","—","Ejecutar el cálculo en una versión privada de pruebas","El cálculo afecta solo a la versión privada","Media"),
("TC-030","Crear y publicar una versión de presupuesto",C[3],"Todos","—","Copiar de Actual a Presupuesto, editar y publicar","La versión pública queda con los valores publicados","Alta"),
("TC-031","Aislamiento entre versiones",C[3],"Todos","—","Editar la versión privada y revisar la pública","Las versiones son independientes","Alta"),
("TC-032","Bloqueo de datos (Data Locking)",C[3],"Todos","Versión aprobada","Bloquear el presupuesto aprobado e intentar editar","La edición se impide para usuarios sin permiso","Alta"),
("TC-033","Simulación / escenario",C[3],"Estados Financieros","—","Crear un escenario (p. ej. nueva licitación) y recalcular","El escenario refleja el impacto esperado","Media"),
("TC-040","Cargar el plan mediante plantilla (Input Form)",C[4],"Todos","Plantilla diseñada","Diligenciar la plantilla y cargar los datos","Los datos se cargan en el modelo y versión correctos","Alta"),
("TC-041","Opción de carga: Reemplazar vs. Acumular",C[4],"Todos","—","Cargar el mismo archivo dos veces con la opción «Reemplazar»","Los valores se reemplazan, no se suman","Alta"),
("TC-042","Validaciones de entrada de datos",C[4],"Todos","—","Ingresar un valor negativo o no permitido","El sistema valida y muestra el mensaje","Media"),
("TC-043","Complemento de Excel (SAC Add-in)",C[4],"Todos","Add-in instalado","Leer y escribir datos desde Excel","Lectura y escritura correctas","Baja"),
("TC-050","Rol de planeación (escritura)",C[5],"Todos","Usuario de planeación","Iniciar sesión y editar los datos de su sociedad","Puede leer y escribir en su sociedad","Alta"),
("TC-051","Rol de consulta (solo lectura)",C[5],"Todos","Usuario de gerencia","Intentar editar datos","Solo lectura; la edición se impide","Alta"),
("TC-052","Data Access Control por sociedad (prueba negativa)",C[5],"Todos","Usuario de la sociedad A","Intentar ver datos de la sociedad B","No visualiza datos de otra sociedad","Alta"),
("TC-053","Segregación de funciones",C[5],"Todos","—","Verificar que un rol no acumula permisos indebidos","Permisos acordes al rol","Media"),
("TC-060","Cifras de la historia (Story) vs. modelo",C[6],"Todos","Story creada","Comparar los KPIs y cifras de la story con el modelo","Coinciden","Alta"),
("TC-061","Filtros e interacciones de la story",C[6],"Todos","—","Aplicar filtros (sociedad, regional, periodo)","Los filtros responden y los datos se actualizan","Media"),
("TC-062","Comparación Presupuesto vs. Real (variación)",C[6],"Todos","Real y presupuesto disponibles","Revisar la variación en $ y %","La variación se calcula correctamente","Alta"),
("TC-063","Exportación (PDF / PPT / Excel)",C[6],"Todos","—","Exportar la story","La exportación conserva datos y formato","Baja"),
("TC-070","Conciliación de ingresos SAC vs. fuente (Silotrans)",C[7],"Ingresos","Histórico importado","Comparar kilos y pesos por regional contra la fuente","Coinciden con la fuente","Alta"),
("TC-071","Conciliación del EEFF vs. contabilidad",C[7],"Estados Financieros","—","Comparar el P&G con la contabilidad","Coinciden","Alta"),
("TC-080","Tiempo de ejecución de las Data Actions",C[8],"Todos","Volumen representativo","Medir el tiempo de ejecución de la Data Action","≤ 60 s (objetivo)","Media"),
("TC-081","Tiempo de apertura y refresco de la story",C[8],"Todos","—","Medir el tiempo de carga de la story","Dentro del umbral definido","Baja"),
("TC-090","Pruebas de regresión tras cambios",C[9],"Todos","Cambio aplicado","Re-ejecutar el conjunto de casos críticos","Sin regresiones","Alta"),
("TC-091","Trazabilidad requerimiento ↔ caso de prueba",C[9],"Todos","—","Verificar que cada caso enlaza a un requerimiento","Cobertura completa de requerimientos","Media"),
]

wb=Workbook()

# =============================== LISTAS (validación) ===============================
ls=wb.active; ls.title="Listas"
listcols=[("A","Estados",list(ESTADOS)),("B","Prioridad",list(PRIOR)),
          ("C","Modelos",MODELOS),("D","Categorías",CATS),
          ("E","EstadoDefecto",["Abierto","En análisis","En corrección","Resuelto","Cerrado"]),
          ("F","Severidad",["Crítica","Alta","Media","Baja"])]
for col,title,items in listcols:
    ls[col+"1"]=title; ls[col+"1"].font=F9B
    for i,v in enumerate(items): ls[col+str(2+i)]=v
ls.sheet_state="hidden"

# =============================== CASOS DE PRUEBA ===============================
ws=wb.create_sheet("Casos de Prueba")
headers=["ID","Caso de prueba","Categoría","Modelo","Precondiciones","Pasos",
         "Resultado esperado","Prioridad","Estado","Resultado obtenido","Tester","Fecha","Evidencia"]
widths=[8,30,22,15,24,34,30,10,12,24,12,11,14]
for j,wd in enumerate(widths): ws.column_dimensions[get_column_letter(1+j)].width=wd
ws.merge_cells("A1:M1"); t=ws["A1"]
t.value="GUION DE PRUEBAS — MODELOS DE SAP ANALYTICS CLOUD (SAC)"
t.font=TITLE; t.fill=HFILL; t.alignment=Alignment("center","center"); ws.row_dimensions[1].height=24
for j,h in enumerate(headers):
    c=ws.cell(2,1+j,h); c.fill=HFILL; c.font=WHITE; c.alignment=CEN; c.border=BORDER
ws.row_dimensions[2].height=30
for i,cs in enumerate(CASES):
    r=3+i
    idv,caso,cat,mod,pre,pasos,esp,pri=cs
    vals=[idv,caso,cat,mod,pre,pasos,esp,pri,"Pendiente","","","",""]
    for j,v in enumerate(vals):
        c=ws.cell(r,1+j,v); c.border=BORDER; c.font=F9
        c.alignment={0:CEN,1:TOPL,2:LEF,3:CEN,4:TOPL,5:TOPL,6:TOPL,7:CEN,8:CEN,9:TOPL,10:LEF,11:CEN,12:LEF}[j]
    ws.cell(r,2).font=F9B
    ws.row_dimensions[r].height=44
last=2+len(CASES)
ws.freeze_panes="C3"; ws.auto_filter.ref=f"A2:M{last}"
# validaciones
for colL,ref,n in [("D","C",len(MODELOS)+1),("H","B",len(PRIOR)+1),("I","A",len(ESTADOS)+1),("C","D",len(CATS)+1)]:
    dv=DataValidation(type="list", formula1=f"Listas!${ref}$2:${ref}${n}", allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"{colL}3:{colL}{last}")
# formato condicional Estado (col I) y Prioridad (col H)
for est,(fill,fc) in ESTADOS.items():
    ws.conditional_formatting.add(f"I3:I{last}", CellIsRule(operator="equal", formula=[f'"{est}"'],
        fill=PatternFill("solid",fgColor=fill), font=Font(color=fc, bold=est in("Aprobado","Fallido"))))
for pr,(fill,fc) in PRIOR.items():
    ws.conditional_formatting.add(f"H3:H{last}", CellIsRule(operator="equal", formula=[f'"{pr}"'],
        fill=PatternFill("solid",fgColor=fill), font=Font(color=fc)))

# =============================== INSTRUCCIONES ===============================
ins=wb.create_sheet("Instrucciones",0)
ins.sheet_view.showGridLines=False
for col,wd in [("A",22),("B",20),("C",20),("D",20),("E",20),("F",16)]:
    ins.column_dimensions[col].width=wd
ins.merge_cells("A1:F1"); h=ins["A1"]
h.value="GUION DE PRUEBAS — MODELOS DE SAP ANALYTICS CLOUD (SAC)"
h.font=TITLE; h.fill=HFILL; h.alignment=Alignment("center","center"); ins.row_dimensions[1].height=26
def block(row, title, text):
    ins.merge_cells(f"A{row}:F{row}"); a=ins[f"A{row}"]; a.value=title; a.font=F9B
    ins.merge_cells(f"A{row+1}:F{row+1}"); b=ins[f"A{row+1}"]; b.value=text
    b.font=Font(size=9.5); b.alignment=Alignment("left","top",wrap_text=True); ins.row_dimensions[row+1].height=30
block(3,"Objetivo","Validar que los modelos de planeación de SAC (Ingresos, Costos y Gastos y Estados Financieros) "
      "funcionan correctamente en datos, cálculos, versiones, seguridad y reportes, antes de su puesta en productivo.")
block(6,"Alcance","Aplica a los tres modelos de las sociedades del programa. Cubre pruebas de datos maestros, integración, "
      "Data Actions, versiones/escenarios, entrada de datos, seguridad (DAC), reportes, conciliación y rendimiento.")
block(9,"Cómo usar","Ejecute cada caso de la hoja «Casos de Prueba», registre el Estado (Pendiente/Aprobado/Fallido/Bloqueado/"
      "No aplica), el resultado obtenido, el tester, la fecha y la evidencia. Los fallos se registran en la hoja «Defectos». "
      "La hoja «Resumen» muestra el avance automáticamente.")
# leyenda de estados
ins["A13"]="Leyenda de estados"; ins["A13"].font=F9B
for i,(e,(fill,fc)) in enumerate(ESTADOS.items()):
    cc=ins.cell(14,1+i, e); cc.fill=PatternFill("solid",fgColor=fill); cc.font=Font(size=9,color=fc,bold=True)
    cc.alignment=CEN; cc.border=BORDER
# buenas prácticas SAP
ins["A16"]="Buenas prácticas de SAP consideradas"; ins["A16"].font=F9B
bps=[
 "Ejecutar las pruebas en un tenant/entorno de prueba, nunca en productivo.",
 "Validar la integridad de datos maestros y jerarquías antes de probar cálculos.",
 "Usar versiones privadas para probar Data Actions y publicar solo lo aprobado.",
 "Comparar cada Data Action contra un cálculo manual o un valor esperado conocido.",
 "Conciliar los datos importados contra la fuente (DataSphere / ERP / Silotrans).",
 "Probar la seguridad con usuarios de distintos roles e incluir pruebas negativas (DAC).",
 "Verificar la gestión de versiones: copia, publicación y bloqueo de datos (Data Locking).",
 "Medir el rendimiento de Data Actions y stories con volúmenes representativos.",
 "Documentar evidencia (capturas) y registrar cada fallo como defecto con su severidad.",
 "Mantener la trazabilidad requerimiento ↔ caso de prueba y ejecutar regresión tras cada cambio.",
]
for i,bp in enumerate(bps):
    r=17+i; ins.merge_cells(f"A{r}:F{r}"); c=ins[f"A{r}"]; c.value="•  "+bp
    c.font=Font(size=9.3); c.alignment=Alignment("left","center",wrap_text=True)

# =============================== RESUMEN (tablero) ===============================
rs=wb.create_sheet("Resumen",1)
rs.sheet_view.showGridLines=False
for col,wd in [("A",26),("B",12),("C",12),("D",20),("E",12),("F",10),("G",3)]:
    rs.column_dimensions[col].width=wd
rs.merge_cells("A1:F1"); hh=rs["A1"]; hh.value="RESUMEN DE EJECUCIÓN DE PRUEBAS"
hh.font=TITLE; hh.fill=HFILL; hh.alignment=Alignment("center","center"); rs.row_dimensions[1].height=26
P="'Casos de Prueba'!"
def kpi(cl,cv,label,formula,fmt=None):
    a=rs[cl]; a.value=label; a.font=F9B; a.alignment=Alignment("left","center")
    b=rs[cv]; b.value=formula; b.font=Font(size=15,bold=True,color=NAVY); b.alignment=CEN
    if fmt: b.number_format=fmt
    a.fill=PatternFill("solid",fgColor="F2F5FA"); b.fill=PatternFill("solid",fgColor="F2F5FA")
    a.border=BORDER; b.border=BORDER
kpi("A3","B3","Total de casos",f"=COUNTA({P}A3:A{last})")
kpi("A4","B4","Ejecutados",f'=COUNTIF({P}I3:I{last},"<>Pendiente")-COUNTBLANK({P}I3:I{last})')
kpi("A5","B5","Aprobados",f'=COUNTIF({P}I3:I{last},"Aprobado")')
kpi("A6","B6","Fallidos",f'=COUNTIF({P}I3:I{last},"Fallido")')
kpi("A7","B7","% ejecución",f'=IFERROR((COUNTA({P}A3:A{last})-COUNTIF({P}I3:I{last},"Pendiente"))/COUNTA({P}A3:A{last}),0)',"0%")
kpi("A8","B8","% aprobación",f'=IFERROR(COUNTIF({P}I3:I{last},"Aprobado")/COUNTA({P}A3:A{last}),0)',"0%")
# tabla por estado (pie)
rs["D3"]="Por estado"; rs["D3"].font=F9B
for j,tt in enumerate(["Estado","Casos"]):
    c=rs.cell(4,4+j,tt); c.fill=HFILL; c.font=WHITE; c.alignment=CEN; c.border=BORDER
for i,e in enumerate(ESTADOS):
    r=5+i; rs.cell(r,4,e).border=BORDER; rs.cell(r,4).font=F9
    rs.cell(r,4).fill=PatternFill("solid",fgColor=ESTADOS[e][0])
    rs.cell(r,5,f'=COUNTIF({P}I3:I{last},"{e}")').border=BORDER; rs.cell(r,5).alignment=CEN; rs.cell(r,5).font=F9
# tabla por categoría (bar)
rs["A10"]="Por categoría"; rs["A10"].font=F9B
for j,tt in enumerate(["Categoría","Casos","Aprobados"]):
    c=rs.cell(11,1+j,tt); c.fill=HFILL; c.font=WHITE; c.alignment=CEN; c.border=BORDER
for i,cat in enumerate(CATS):
    r=12+i
    rs.cell(r,1,cat).border=BORDER; rs.cell(r,1).font=Font(size=8.5); rs.cell(r,1).alignment=LEF
    rs.cell(r,2,f'=COUNTIF({P}C3:C{last},"{cat}")').border=BORDER; rs.cell(r,2).alignment=CEN; rs.cell(r,2).font=F9
    rs.cell(r,3,f'=COUNTIFS({P}C3:C{last},"{cat}",{P}I3:I{last},"Aprobado")').border=BORDER; rs.cell(r,3).alignment=CEN; rs.cell(r,3).font=F9
# gráfico pie estados
pie=PieChart(); pie.title="Casos por estado"; pie.height=6.2; pie.width=9.5
pdata=Reference(rs,min_col=5,min_row=4,max_row=9); pcats=Reference(rs,min_col=4,min_row=5,max_row=9)
pie.add_data(pdata,titles_from_data=True); pie.set_categories(pcats)
pie.dataLabels=DataLabelList(); pie.dataLabels.showVal=True
rs.add_chart(pie,"D11")
# gráfico barras por categoría
bar=BarChart(); bar.type="bar"; bar.title="Casos por categoría"; bar.height=8.0; bar.width=11
bdata=Reference(rs,min_col=2,min_row=11,max_row=21); bcats=Reference(rs,min_col=1,min_row=12,max_row=21)
bar.add_data(bdata,titles_from_data=True); bar.set_categories(bcats); bar.legend=None
rs.add_chart(bar,"D23")

# =============================== DEFECTOS ===============================
df=wb.create_sheet("Defectos")
dh=["ID","Caso (TC)","Descripción","Severidad","Estado","Responsable","Fecha","Resolución"]
dw=[9,10,34,12,14,14,11,30]
for j,wd in enumerate(dw): df.column_dimensions[get_column_letter(1+j)].width=wd
df.merge_cells("A1:H1"); t=df["A1"]; t.value="REGISTRO DE DEFECTOS"
t.font=TITLE; t.fill=HFILL; t.alignment=Alignment("center","center"); df.row_dimensions[1].height=22
for j,h in enumerate(dh):
    c=df.cell(2,1+j,h); c.fill=HFILL; c.font=WHITE; c.alignment=CEN; c.border=BORDER
example=["DEF-001","TC-024","(ejemplo) El P&G consolidado no cuadra por CEBE","Alta","Abierto","(responsable)","","(pendiente)"]
rows=[example]+[[ "" ]*8 for _ in range(14)]
for i,row in enumerate(rows):
    r=3+i
    for j,v in enumerate(row):
        c=df.cell(r,1+j,v); c.border=BORDER; c.font=F9
        c.alignment=TOPL if j in (2,7) else (CEN if j in (0,1,3,4,6) else LEF)
    df.row_dimensions[r].height=26
dlast=2+len(rows)
df.freeze_panes="A3"; df.auto_filter.ref=f"A2:H{dlast}"
for colL,ref,n in [("D","F",5),("E","E",6)]:
    dv=DataValidation(type="list", formula1=f"Listas!${ref}$2:${ref}${n}", allow_blank=True)
    df.add_data_validation(dv); dv.add(f"{colL}3:{colL}{dlast}")

wb.active=0
fn="Guion_Pruebas_Modelos_SAC.xlsx"
wb.save(fn)
print("GUARDADO:", fn, "| casos:", len(CASES), "| hojas:", wb.sheetnames)
