import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# ─── COLORES CORPORATIVOS ─────────────────────────────────────────────────────
SAP_BLUE   = RGBColor(0x00, 0x61, 0x9A)
SAP_DARK   = RGBColor(0x00, 0x33, 0x69)
SAP_GOLD   = RGBColor(0xF0, 0xAB, 0x00)
SAP_LIGHT  = RGBColor(0xE8, 0xF4, 0xFD)
SAP_GREEN  = RGBColor(0x10, 0x7E, 0x3E)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_MED   = RGBColor(0x75, 0x75, 0x75)
GRAY_LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
ORANGE     = RGBColor(0xE8, 0x6A, 0x19)
PURPLE     = RGBColor(0x7B, 0x3F, 0xA0)

# ─── HELPERS DOCX ─────────────────────────────────────────────────────────────
def set_cell_bg(cell, rgb):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  '%02X%02X%02X' % (rgb[0], rgb[1], rgb[2]))
    tcPr.append(shd)

def set_cell_borders(cell, sides=('top','bottom','left','right'), color='B0C4D8', sz=4):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for s in sides:
        b = OxmlElement(f'w:{s}')
        b.set(qn('w:val'),   'single')
        b.set(qn('w:sz'),    str(sz))
        b.set(qn('w:color'), color)
        tcBorders.append(b)
    tcPr.append(tcBorders)

def add_heading(doc, text, level=1, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    if level == 1:
        run.font.size  = Pt(16)
        run.font.bold  = True
        run.font.color.rgb = SAP_DARK
        pPr  = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bot  = OxmlElement('w:bottom')
        bot.set(qn('w:val'),   'single')
        bot.set(qn('w:sz'),    '6')
        bot.set(qn('w:color'), '003369')
        pBdr.append(bot)
        pPr.append(pBdr)
    elif level == 2:
        run.font.size  = Pt(13)
        run.font.bold  = True
        run.font.color.rgb = SAP_BLUE
    elif level == 3:
        run.font.size  = Pt(11)
        run.font.bold  = True
        run.font.color.rgb = SAP_BLUE

def add_body(doc, text, space_before=2, space_after=4, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.size   = Pt(10)
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before    = Pt(1)
    p.paragraph_format.space_after     = Pt(1)
    p.paragraph_format.left_indent     = Cm(0.5 + level * 0.5)
    run = p.add_run(text)
    run.font.size = Pt(10)

def header_row(table, cols, bg=None):
    bg = bg or SAP_DARK
    row = table.rows[0]
    for i, txt in enumerate(cols):
        cell = row.cells[i]
        set_cell_bg(cell, bg)
        set_cell_borders(cell, color='003369', sz=6)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(txt)
        run.font.size  = Pt(9)
        run.font.bold  = True
        run.font.color.rgb = WHITE

def data_row(table, row_idx, values, alt=False, center_cols=None, bold_col=None):
    center_cols = center_cols or []
    row = table.rows[row_idx]
    bg  = SAP_LIGHT if alt else WHITE
    for i, val in enumerate(values):
        cell = row.cells[i]
        set_cell_bg(cell, bg)
        set_cell_borders(cell, color='C5D8EA', sz=2)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p   = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(val))
        run.font.size = Pt(9)
        if bold_col is not None and i == bold_col:
            run.font.bold  = True
            run.font.color.rgb = SAP_DARK

def add_info_box(doc, title, lines, bg_title=None, bg_body=None):
    bg_title = bg_title or SAP_BLUE
    bg_body  = bg_body  or SAP_LIGHT
    tbl = doc.add_table(rows=1 + len(lines), cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tc = tbl.rows[0].cells[0]
    set_cell_bg(tc, bg_title)
    set_cell_borders(tc, color='003369', sz=6)
    p = tc.paragraphs[0]
    run = p.add_run(title)
    run.font.size  = Pt(10)
    run.font.bold  = True
    run.font.color.rgb = WHITE
    for i, line in enumerate(lines, 1):
        tc = tbl.rows[i].cells[0]
        set_cell_bg(tc, bg_body)
        set_cell_borders(tc, color='B0C4D8', sz=2)
        p = tc.paragraphs[0]
        run = p.add_run(line)
        run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_footer(doc):
    section = doc.sections[0]
    footer  = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('SAP BUILD | Fanalca S.A. — SAP Analytics Cloud | Documento Confidencial | 2026')
    run.font.size   = Pt(8)
    run.font.color.rgb = GRAY_MED
    run.font.italic = True

# ─── FLOWCHART ────────────────────────────────────────────────────────────────
def generar_flowchart(path):
    fig, ax = plt.subplots(figsize=(22, 15))
    ax.set_xlim(0, 22)
    ax.set_ylim(0, 15)
    ax.axis('off')
    fig.patch.set_facecolor('#F5F8FC')

    def band(y, h, color, label):
        ax.add_patch(plt.Rectangle((0.3, y), 21.4, h, color=color, zorder=0, alpha=0.22))
        ax.text(0.52, y + h/2, label, va='center', ha='left',
                fontsize=8, color='#2C3E50', fontweight='bold', rotation=90, zorder=2)

    def box(x, y, w, h, label, sub='', fc='#00619A', tc='white', fs=9, subfs=7.5, radius=0.2):
        ax.add_patch(FancyBboxPatch((x, y), w, h,
                     boxstyle=f'round,pad=0.04,rounding_size={radius}',
                     fc=fc, ec='white', lw=1.8, zorder=3))
        cy = y + h/2 + (0.13 if sub else 0)
        ax.text(x + w/2, cy, label, ha='center', va='center',
                fontsize=fs, color=tc, fontweight='bold', zorder=4, multialignment='center')
        if sub:
            ax.text(x + w/2, y + h/2 - 0.22, sub, ha='center', va='center',
                    fontsize=subfs, color=tc, style='italic', zorder=4)

    def arrow(x1, y1, x2, y2, color='#003369', lw=1.4):
        ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle='->', color=color, lw=lw))

    # Bands
    band(0.4,  1.8,  '#D6E8F7', 'FUENTES DE DATOS SAP')
    band(2.4,  1.4,  '#FFF5D6', 'CAPA DE INTEGRACIÓN')
    band(4.0,  4.2,  '#E8F5E9', 'MODELOS SAP ANALYTICS CLOUD')
    band(8.4,  3.8,  '#F3E5F5', 'UNIDADES DE NEGOCIO / PROCESOS')
    band(12.4, 2.2,  '#FFEBEE', 'SALIDAS — DASHBOARDS & REPORTES')

    # Fuentes de datos
    box(0.8,  0.6,  2.8, 1.1, 'SAP FI', 'Contabilidad Financiera\nCuentas · Sociedades', fc='#1565C0', fs=9)
    box(4.0,  0.6,  2.8, 1.1, 'SAP CO', 'Controlling\nCEBES · CECOS · CUENTAS', fc='#1565C0', fs=9)
    box(7.2,  0.6,  2.8, 1.1, 'SAP SD', 'Ventas y Distribución\nClientes · Referencias', fc='#1565C0', fs=9)
    box(10.4, 0.6,  2.8, 1.1, 'SAP MM', 'Gestión de Materiales\nReferencias / SKU', fc='#1565C0', fs=9)
    box(13.6, 0.6,  2.8, 1.1, 'BW/4HANA', 'Data Warehouse\nConsolidación DWH', fc='#0D47A1', fs=9)
    box(16.8, 0.6,  2.8, 1.1, 'SLT / DS', 'Replicación y\nTransformación', fc='#0D47A1', fs=9)
    box(19.8, 0.6,  1.8, 1.1, 'TCURC\nMONEDA', fc='#1976D2', fs=8)

    # Integración
    box(3.5,  2.6,  14.5, 0.9, 'SAP Analytics Cloud — Import Connection  |  Live Connection  |  OData / CSV / BW Connector',
        fc='#F57F17', tc='#1A1A1A', fs=9.5)
    for x in [2.2, 5.4, 8.6, 11.8, 15.0, 18.2]:
        arrow(x, 1.7, x, 2.6)
    arrow(20.7, 1.7, 18.0, 2.6)

    # Modelos SAC
    # Ingresos
    box(0.8,  4.3,  6.0, 3.6,
        'Modelo Ingresos\nFanalca',
        'Medida: Importe (Decimal)\n9 Dimensiones:\nVersion  ·  RATIO  ·  Date\nAUDITORIA  ·  CEBES  ·  CLIENTES\nMONEDA  ·  REFERENCIA  ·  SOCIEDAD',
        fc='#2E7D32', fs=10.5, subfs=8)

    # EEFF
    box(7.8,  4.3,  6.0, 3.6,
        'Modelo EEFF\nFanalca',
        'Medida: Importe (Decimal)\n7 Dimensiones:\nVersion  ·  CUENTA  ·  Date\nAUDITORIA  ·  CEBES\nMONEDA  ·  SOCIEDAD',
        fc='#6A1B9A', fs=10.5, subfs=8)

    # Costos y Gastos
    box(14.8, 4.3,  6.2, 3.6,
        'Modelo Costos\ny Gastos Fanalca',
        'Medida: Importe (Decimal)\n8 Dimensiones:\nVersion  ·  CUENTAS  ·  Date\nAUDITORIA  ·  CEBES  ·  CECOS\nMONEDA  ·  SOCIEDAD',
        fc='#E65100', fs=10.5, subfs=8)

    for x in [3.8, 10.8, 17.9]:
        arrow(x, 3.5, x, 4.3)

    # Shared dims label
    ax.text(11, 3.75, '★ Dimensiones compartidas: SOCIEDAD · CEBES · MONEDA · AUDITORIA · Version · Date',
            ha='center', va='center', fontsize=7.5, color='#003369',
            style='italic', fontweight='bold',
            bbox=dict(boxstyle='round,pad=0.3', fc='#E3F2FD', ec='#1565C0', lw=0.8))

    # Unidades de negocio
    box(0.8,  8.7,  5.8, 3.0,
        'Fanalca Metalmecánica',
        'Manufactura de autopartes\ny componentes metálicos\n→ Control Costos Producción\n→ Análisis Gastos Operativos',
        fc='#BF360C', fs=10, subfs=8.5)

    box(7.8,  8.7,  6.0, 3.0,
        'Honda Supermotos',
        'Importación y distribución\nde motocicletas Honda\n→ Ingresos por línea/cliente\n→ Análisis de márgenes',
        fc='#558B2F', fs=10, subfs=8.5)

    box(14.8, 8.7,  6.2, 3.0,
        'Honda Autos',
        'Distribución vehículos\nlivianos Honda Colombia\n→ Estados Financieros\n→ Rentabilidad por sociedad',
        fc='#4527A0', fs=10, subfs=8.5)

    arrow(3.8, 7.9, 3.8, 8.7)
    arrow(10.8, 7.9, 10.8, 8.7)
    arrow(17.9, 7.9, 17.9, 8.7)

    # Dashboards
    dash = [
        (1.0,  'Dashboard\nEjecutivo\n24/7'),
        (4.5,  'Análisis de\nIngresos\nGranular'),
        (7.8,  'EEFF\nConsolidado\nMulti-sociedad'),
        (11.2, 'Control\nCostos y\nGastos'),
        (14.6, 'KPIs\nVariación\nReal vs Plan'),
        (18.0, 'Reportes\nxP&A /\nPlaneación'),
    ]
    for x, lbl in dash:
        box(x, 12.6, 3.0, 1.6, lbl, fc='#B71C1C', fs=8.5)
    arrow(11, 11.7, 11, 12.6)

    # Archivos validación
    box(19.0, 4.3,  2.5, 1.2, 'Archivos\nValidación', 'SAC vs SAP FI', fc='#37474F', fs=8.5, subfs=7.5)
    arrow(19.0, 3.5, 19.0+1.25, 4.3)

    # Title
    ax.text(11, 14.6, 'Arquitectura de Modelos SAC — Fanalca S.A. | Grupo Fanalca',
            ha='center', va='center', fontsize=14, fontweight='bold', color='#003369')
    ax.text(11, 14.15, 'Proyecto SAP BUILD  ·  SAP Analytics Cloud Planning  ·  Metalmecánica | Honda Supermotos | Honda Autos  ·  2026',
            ha='center', va='center', fontsize=9, color='#555', style='italic')

    # Legend
    items = [
        (mpatches.Patch(color='#1565C0'), 'Módulos SAP ERP'),
        (mpatches.Patch(color='#F57F17'), 'Integración SAC'),
        (mpatches.Patch(color='#2E7D32'), 'Modelo Ingresos'),
        (mpatches.Patch(color='#6A1B9A'), 'Modelo EEFF'),
        (mpatches.Patch(color='#E65100'), 'Modelo Costos y Gastos'),
        (mpatches.Patch(color='#BF360C'), 'Metalmecánica'),
        (mpatches.Patch(color='#558B2F'), 'Honda Supermotos'),
        (mpatches.Patch(color='#4527A0'), 'Honda Autos'),
        (mpatches.Patch(color='#B71C1C'), 'Dashboards / Salidas'),
    ]
    fig.legend([h for h, _ in items], [l for _, l in items],
               loc='lower center', ncol=5, fontsize=8.5,
               framealpha=0.9, edgecolor='#BBBBBB',
               bbox_to_anchor=(0.5, -0.01))

    plt.tight_layout(rect=[0, 0.05, 1, 1])
    plt.savefig(path, dpi=150, bbox_inches='tight', facecolor='#F5F8FC')
    plt.close()
    print(f'Flowchart: {path}')

# ─── DOCUMENTO WORD ───────────────────────────────────────────────────────────
def generar_doc(flowchart_path, doc_path):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10)
    add_footer(doc)

    # ── PORTADA ──────────────────────────────────────────────────────────────
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, SAP_DARK)

    def cline(text, sz, bold=True, color=WHITE, sb=4, sa=4, italic=False):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(sb)
        p.paragraph_format.space_after  = Pt(sa)
        r = p.add_run(text)
        r.font.size = Pt(sz); r.font.bold = bold
        r.font.color.rgb = color; r.font.italic = italic

    cline('', 6, sb=20, sa=0)
    cline('SAP BUILD', 13, color=SAP_GOLD, italic=True, sa=2)
    cline('GRUPO FANALCA', 11, color=RGBColor(0xCC, 0xDD, 0xEE), bold=False)
    cline('', 6, sb=8, sa=8)
    cline('Documento de Diseño de Solución', 14, color=WHITE)
    cline('SAP Analytics Cloud', 24, color=WHITE, sa=2)
    cline('Modelos de Planeación Financiera', 16, color=SAP_GOLD, sb=2)
    cline('', 6, sb=12, sa=12)

    meta_tbl = doc.add_table(rows=6, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate([
        ('Cliente',   'Fanalca S.A.'),
        ('Proyecto',  'SAP BUILD — Implementación SAP S/4HANA + SAC'),
        ('Unidades',  'Metalmecánica | Honda Supermotos | Honda Autos'),
        ('Versión',   'v1.0'),
        ('Fecha',     '05 de junio de 2026'),
        ('Estado',    'Borrador para revisión y aprobación'),
    ]):
        r = meta_tbl.rows[i]
        set_cell_bg(r.cells[0], SAP_BLUE)
        set_cell_bg(r.cells[1], SAP_LIGHT)
        set_cell_borders(r.cells[0], color='003369', sz=4)
        set_cell_borders(r.cells[1], color='B0C4D8', sz=2)
        pk = r.cells[0].paragraphs[0]
        pk.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rk = pk.add_run(k)
        rk.font.size = Pt(9); rk.font.bold = True; rk.font.color.rgb = WHITE
        pv = r.cells[1].paragraphs[0]
        rv = pv.add_run(v)
        rv.font.size = Pt(9); rv.font.color.rgb = SAP_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    doc.add_page_break()

    # ── TABLA DE CONTENIDO ───────────────────────────────────────────────────
    add_heading(doc, 'Tabla de Contenido', level=1, space_before=4)
    toc = [
        ('1.', 'Propósito del documento', '4'),
        ('2.', 'Cómo usar el documento', '4'),
        ('3.', 'Escenario / Proceso de negocio', '4'),
        ('3.1.', 'Objetivos empresariales y beneficios esperados', '5'),
        ('3.2.', 'Descripción a alto nivel de Requisitos empresariales', '5'),
        ('4.', 'Estructura organizativa relevante', '6'),
        ('5.', 'Diseño y configuración de soluciones', '7'),
        ('5.1.', 'Alcance del proceso de solución', '7'),
        ('5.2.', 'Valor de Actividad de Configuración', '9'),
        ('5.3.', 'Roles', '10'),
        ('5.4.', 'Inventario de datos maestros susceptibles de migración', '11'),
        ('5.5.', 'Componente de Solución Técnica Relacionada', '12'),
    ]
    tt = doc.add_table(rows=len(toc), cols=3)
    for i, (n, t, p) in enumerate(toc):
        r = tt.rows[i]
        bg = SAP_LIGHT if i % 2 == 0 else WHITE
        for c in r.cells: set_cell_bg(c, bg)
        r.cells[0].paragraphs[0].add_run(n).font.size = Pt(9)
        p1 = r.cells[1].paragraphs[0]
        p1.paragraph_format.left_indent = Cm(0 if '.' not in n[1:] else 0.5)
        run = p1.add_run(t)
        run.font.size = Pt(9)
        run.font.bold = ('.' not in n[1:])
        r.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r.cells[2].paragraphs[0].add_run(p).font.size = Pt(9)
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # 1. PROPÓSITO DEL DOCUMENTO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '1. Propósito del documento', level=1)
    add_body(doc,
        'El presente documento describe el diseño y la configuración de la solución '
        'SAP Analytics Cloud (SAC) implementada para Fanalca S.A., empresa colombiana '
        'líder en el sector automotriz, metalmecánico y de distribución de vehículos Honda, '
        'con más de 70 años de trayectoria en el mercado. Este artefacto constituye el '
        'Blueprint de Diseño de Solución del proyecto SAP BUILD, detallando los tres modelos '
        'analíticos de planeación financiera, sus dimensiones, las fuentes de datos, los '
        'componentes técnicos y los datos maestros involucrados.')
    add_body(doc,
        'El documento cubre los tres modelos centrales del proyecto, que soportan la '
        'analítica financiera y comercial de las unidades de negocio Metalmecánica, '
        'Honda Supermotos y Honda Autos de Fanalca S.A.:')

    add_info_box(doc, 'Modelos SAC en Alcance',
        ['Modelo de Ingresos Fanalca — 9 dimensiones | Planning Model (Read/Write)',
         'Modelo de Estados Financieros (EEFF) Fanalca — 7 dimensiones | Planning Model (Read/Write)',
         'Modelo de Costos y Gastos Fanalca — 8 dimensiones | Planning Model (Read/Write)',
         'Dimensiones compartidas: SOCIEDAD · CEBES · MONEDA · AUDITORIA · Version · Date'],
        bg_title=SAP_DARK)

    add_body(doc, 'Alcance del documento:')
    for b in [
        'Unidades de negocio: Metalmecánica, Honda Supermotos, Honda Autos',
        'Plataforma: SAP Analytics Cloud (SAC) — Tenencia en la nube SAP BTP',
        'Período de referencia: Ejercicios fiscales 2024–2026',
        'Versión SAC: 2024.Q4 en adelante',
        'Fuera de alcance: Modelo de Flujo de Caja y módulos operativos distintos de FI/CO/SD/MM',
    ]:
        add_bullet(doc, b)

    # ══════════════════════════════════════════════════════════════════════════
    # 2. CÓMO USAR EL DOCUMENTO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '2. Cómo usar el documento', level=1)
    add_body(doc,
        'Este documento está dirigido a múltiples perfiles del proyecto. Utilice la '
        'Tabla de Contenido para navegar directamente a la sección de interés. Las tablas '
        'de la Sección 5 contienen los valores de configuración que deben implementarse '
        'en SAC en la secuencia descrita.')

    uso_t = doc.add_table(rows=6, cols=3)
    header_row(uso_t, ['Perfil', 'Uso del Documento', 'Secciones Clave'])
    for i, (p, u, s) in enumerate([
        ('Arquitectos SAP/SAC',   'Revisión del diseño técnico de modelos y dimensiones',          'Secciones 4, 5.1'),
        ('Analistas Funcionales', 'Validación del mapeo de requisitos de negocio',                  'Secciones 3, 5.2'),
        ('Equipo TI / BASIS',     'Referencia de configuración e integración SAP ERP ↔ SAC',        'Secciones 5.1, 5.5'),
        ('Usuarios Clave',        'Validación de la lógica de negocio y pruebas UAT',               'Secciones 3, 5.3'),
        ('Gestión del Proyecto',  'Seguimiento del alcance, actividades, roles y entregables',      'Secciones 5.2, 5.3'),
    ], 1):
        data_row(uso_t, i, [p, u, s], alt=(i % 2 == 0), bold_col=0)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # 3. ESCENARIO / PROCESO DE NEGOCIO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '3. Escenario / Proceso de negocio', level=1)
    add_body(doc,
        'Fanalca S.A. es un conglomerado industrial colombiano con más de 70 años de '
        'trayectoria, conformado por tres grandes unidades de negocio estratégicas. '
        'La compañía requiere una plataforma de analítica centralizada que integre los '
        'datos financieros y operativos de sus divisiones para soportar la toma de '
        'decisiones gerenciales y ejecutivas en tiempo real. Actualmente el proceso de '
        'reporte financiero depende de hojas de cálculo manuales y consolidaciones en '
        'Excel, generando demoras en el cierre mensual y limitando la capacidad de '
        'análisis por unidad de negocio.')

    # Tabla unidades de negocio
    un_t = doc.add_table(rows=4, cols=4)
    header_row(un_t, ['Unidad de Negocio', 'Descripción', 'Proceso Clave SAC', 'Video Sesión'])
    for i, (u, d, p, v) in enumerate([
        ('Fanalca Metalmecánica',
         'Manufactura de autopartes, componentes metálicos y ensambles para la industria automotriz colombiana.',
         'Control de Costos de Producción y análisis de Gastos Operativos por CECOS/CEBES',
         'Fanalca Metalmecanica.mp4'),
        ('Honda Supermotos',
         'Importación, distribución y comercialización de motocicletas Honda en Colombia. Red nacional de concesionarios.',
         'Análisis de Ingresos por línea de producto (referencia/SKU), cliente y región',
         'Fanalca Supermotos.mp4'),
        ('Honda Autos',
         'Distribución de vehículos livianos Honda. Gestión de inventario, financiamiento y rentabilidad.',
         'Estados Financieros (EEFF) consolidados y análisis de márgenes por sociedad',
         'Fanalca Autos Honda.mp4'),
    ], 1):
        data_row(un_t, i, [u, d, p, v], alt=(i % 2 == 0), bold_col=0)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Flowchart
    add_heading(doc, 'Arquitectura de Modelos SAC — Fanalca S.A.', level=2, space_before=8)
    doc.add_picture(flowchart_path, width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph('Figura 1. Arquitectura de modelos SAP Analytics Cloud — Fanalca S.A. | Proyecto SAP BUILD 2026')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(8); cap.runs[0].font.italic = True
    cap.runs[0].font.color.rgb = GRAY_MED
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3.1
    add_heading(doc, '3.1. Objetivos empresariales y beneficios esperados', level=2)
    obj_t = doc.add_table(rows=6, cols=3)
    header_row(obj_t, ['#', 'Objetivo Empresarial', 'Beneficio Esperado'])
    for i, (n, o, b) in enumerate([
        ('1', 'Visibilidad financiera consolidada',
              'Vista unificada de resultados de las 3 unidades en tiempo real; eliminación de silos en Excel'),
        ('2', 'Análisis de ingresos granular',
              'Desagregación por cliente, referencia SKU, canal, CEBES y moneda para decisiones comerciales ágiles'),
        ('3', 'Control de costos y gastos',
              'Seguimiento detallado por centro de costo (CECOS) y centro de beneficio (CEBES) para gestión de OPEX'),
        ('4', 'Cierre financiero ágil',
              'Reducción del tiempo de cierre mensual en 60% mediante automatización de EEFF integrado con SAP FI'),
        ('5', 'Planeación integrada xP&A',
              'Base para planificación extendida que integre planes de ventas, costos y gastos de todas las unidades'),
    ], 1):
        data_row(obj_t, i, [n, o, b], alt=(i % 2 == 0), center_cols=[0], bold_col=1)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # KPIs
    kpi_t = doc.add_table(rows=5, cols=3)
    header_row(kpi_t, ['Área', 'Beneficio Esperado', 'KPI de Medición'], bg=SAP_BLUE)
    for i, (a, b, k) in enumerate([
        ('Finanzas',      'Reducción tiempo de cierre mensual 60%',     '# Días de cierre mensual'),
        ('Comercial',     'Incremento en análisis de rentabilidad',      '# Reportes en autoservicio SAC'),
        ('Operaciones',   'Control desviación presupuestal < 5%',        '% Variación Real vs. Plan'),
        ('Alta Dirección','Dashboard ejecutivo 24/7 sin dependencia TI', 'NPS Usuarios / % Adopción SAC'),
    ], 1):
        data_row(kpi_t, i, [a, b, k], alt=(i % 2 == 0), bold_col=0)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3.2
    add_heading(doc, '3.2. Descripción a alto nivel de Requisitos empresariales', level=2)
    req_t = doc.add_table(rows=8, cols=4)
    header_row(req_t, ['ID', 'Requisito', 'Descripción', 'Modelo SAC'])
    for i, (rid, req, desc, mod) in enumerate([
        ('REQ-001', 'Modelo de Ingresos',
         'Reportar ingresos por sociedad, cliente, referencia SKU, CEBES, moneda y ratio con comparativos Real vs. Plan',
         'Ingresos'),
        ('REQ-002', 'Modelo EEFF',
         'Generar Balance General y Estado de Resultados por cuenta contable, sociedad y CEBES con visión multi-sociedad',
         'EEFF'),
        ('REQ-003', 'Modelo Costos y Gastos',
         'Analizar costos/gastos por cuenta, CECOS, CEBES y sociedad con granularidad mensual y acumulado anual',
         'Costos y Gastos'),
        ('REQ-004', 'Archivos de Validación',
         'Proceso de validación y conciliación entre datos SAC y SAP ERP/FI para garantizar integridad de información',
         'Todos'),
        ('REQ-005', 'Multi-moneda',
         'Soporte para análisis en COP, USD y EUR a través de la dimensión MONEDA compartida entre modelos',
         'Todos'),
        ('REQ-006', 'Versionamiento',
         'Manejo de versiones Real, Plan, Forecast y Budget mediante la dimensión estándar Version de SAC',
         'Todos'),
        ('REQ-007', 'Seguridad por Sociedad',
         'Control de acceso a datos por sociedad (SOCIEDAD) para garantizar confidencialidad entre unidades de negocio',
         'Todos'),
    ], 1):
        data_row(req_t, i, [rid, req, desc, mod], alt=(i % 2 == 0), center_cols=[0, 3], bold_col=1)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # 4. ESTRUCTURA ORGANIZATIVA RELEVANTE
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '4. Estructura organizativa relevante', level=1)
    add_body(doc,
        'La estructura organizativa de Fanalca S.A. relevante para SAC refleja la jerarquía '
        'de entidades contables y analíticas definidas en SAP ERP (S/4HANA), las cuales se '
        'mapean directamente a las dimensiones de los modelos SAC. Las dimensiones marcadas '
        'con ★ son compartidas entre los tres modelos, garantizando consistencia y análisis '
        'integrado.')

    org_t = doc.add_table(rows=13, cols=4)
    header_row(org_t, ['Dimensión SAC', 'Tipo', 'Descripción / Equivalente SAP', 'Modelos'])
    for i, (d, tp, desc, mod) in enumerate([
        ('★ SOCIEDAD',  'Genérica', 'Entidad jurídica / Company Code SAP FI (tabla T001)',                    'TODOS'),
        ('★ CEBES',     'Genérica', 'Centro de Beneficio / Profit Center SAP CO-PCA (tabla CEPC)',            'TODOS'),
        ('★ MONEDA',    'Genérica', 'Dimensión multimoneda COP/USD/EUR (tabla TCURC)',                        'TODOS'),
        ('★ AUDITORIA', 'Genérica', 'Control y trazabilidad de carga de datos por origen/fecha',              'TODOS'),
        ('★ Version',   'Estándar SAC', 'Versión de datos: Real / Plan / Forecast / Budget',                 'TODOS'),
        ('★ Date',      'Estándar SAC', 'Dimensión temporal Año/Mes (granularidad mensual)',                  'TODOS'),
        ('CUENTA',      'Genérica', 'Cuenta contable Plan de Cuentas SAP FI (tablas SKA1/SKB1)',              'EEFF'),
        ('CUENTAS',     'Genérica', 'Cuenta de costos/gastos SAP CO (tablas CSKA/CSKB)',                     'Costos/Gastos'),
        ('CECOS',       'Genérica', 'Centro de Costo SAP CO-CCA (tabla CSKS)',                               'Costos/Gastos'),
        ('CLIENTES',    'Genérica', 'Maestro de clientes SAP SD (tabla KNA1)',                               'Ingresos'),
        ('REFERENCIA',  'Genérica', 'Referencia/SKU de producto SAP MM (tablas MARA/MVKE)',                  'Ingresos'),
        ('RATIO',       'Genérica', 'Ratio precio/volumen para análisis de variaciones de ingresos',          'Ingresos'),
    ], 1):
        data_row(org_t, i, [d, tp, desc, mod], alt=(i % 2 == 0),
                 bold_col=0, center_cols=[1, 3])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_info_box(doc, 'Nota — Dimensiones Compartidas (★)',
        ['Las dimensiones SOCIEDAD, CEBES, MONEDA, AUDITORIA, Version y Date son compartidas '
         'entre los 3 modelos SAC.',
         'Esto garantiza consistencia en el análisis integrado y permite reportes cross-modelo '
         'sin transformaciones adicionales.',
         'Las dimensiones específicas por modelo (CUENTA, CECOS, CLIENTES, REFERENCIA, RATIO) '
         'reflejan la granularidad requerida por cada proceso de negocio.'],
        bg_title=SAP_BLUE)

    # Jerarquías organizativas
    add_heading(doc, 'Jerarquías y Sociedades de Fanalca S.A.', level=2, space_before=8)
    soc_t = doc.add_table(rows=5, cols=4)
    header_row(soc_t, ['Sociedad', 'Unidad de Negocio', 'Módulos SAP', 'Dimensiones Clave'])
    for i, (s, u, m, d) in enumerate([
        ('Fanalca S.A. Metalmecánica', 'Manufactura autopartes',       'FI · CO · MM · PP',
         'CECOS · CEBES · CUENTAS · SOCIEDAD'),
        ('Honda Supermotos Colombia',  'Distribución motos Honda',      'FI · CO · SD · MM',
         'CLIENTES · REFERENCIA · RATIO · CEBES'),
        ('Honda Autos (Autos Honda)',  'Distribución vehículos livianos','FI · CO · SD · MM',
         'CUENTA · CEBES · MONEDA · SOCIEDAD'),
    ], 1):
        data_row(soc_t, i, [s, u, m, d], alt=(i % 2 == 0), bold_col=0)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ══════════════════════════════════════════════════════════════════════════
    # 5. DISEÑO Y CONFIGURACIÓN
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, '5. Diseño y configuración de soluciones', level=1)

    # 5.1
    add_heading(doc, '5.1. Alcance del proceso de solución', level=2)
    add_body(doc,
        'El alcance de SAC para Fanalca S.A. contempla 3 modelos Planning interconectados. '
        'Los modelos de Ingresos y Costos/Gastos alimentan el modelo EEFF para la generación '
        'del Estado de Resultados consolidado y el Balance General por sociedad y unidad de negocio.')

    # 3 Models detail tables
    for title, mtype, dims, extra in [
        ('Modelo 1: Ingresos Fanalca',
         'Planning Model (Read/Write) | Medida: Importe (Decimal) | 9 Dimensiones',
         [('Version★',   'Estándar SAC', 'Versión de datos (Real / Plan / Forecast / Budget)'),
          ('RATIO',      'Genérica',     'Ratio precio/volumen para análisis de variaciones de ingresos'),
          ('Date★',      'Estándar SAC', 'Dimensión temporal Año/Mes'),
          ('AUDITORIA★', 'Genérica',     'Control y trazabilidad de carga de datos'),
          ('CEBES★',     'Genérica',     'Centro de Beneficio / Profit Center SAP CO (CEPC)'),
          ('CLIENTES',   'Genérica',     'Maestro de clientes SAP SD (KNA1) — Supermotos / Autos Honda'),
          ('MONEDA★',    'Genérica',     'Dimensión de moneda COP / USD / EUR'),
          ('REFERENCIA', 'Genérica',     'Referencia / SKU de producto SAP MM (MARA/MVKE)'),
          ('SOCIEDAD★',  'Genérica',     'Sociedad / Company Code SAP FI (T001)')],
         'Unidades: Honda Supermotos · Honda Autos | Proceso: Análisis ingresos por cliente, referencia y CEBES'),

        ('Modelo 2: EEFF Fanalca (Estados Financieros)',
         'Planning Model (Read/Write) | Medida: Importe (Decimal) | 7 Dimensiones',
         [('Version★',   'Estándar SAC', 'Versión de datos (Real / Plan / Forecast / Budget)'),
          ('CUENTA',     'Genérica',     'Cuenta contable Plan de Cuentas SAP FI (SKA1/SKB1)'),
          ('Date★',      'Estándar SAC', 'Dimensión temporal Año/Mes'),
          ('AUDITORIA★', 'Genérica',     'Control y trazabilidad de carga de datos'),
          ('CEBES★',     'Genérica',     'Centro de Beneficio / Profit Center SAP CO (CEPC)'),
          ('MONEDA★',    'Genérica',     'Dimensión de moneda COP / USD / EUR'),
          ('SOCIEDAD★',  'Genérica',     'Sociedad / Company Code SAP FI (T001)')],
         'Unidades: Todas | Proceso: Balance General · Estado de Resultados · Análisis de márgenes'),

        ('Modelo 3: Costos y Gastos Fanalca',
         'Planning Model (Read/Write) | Medida: Importe (Decimal) | 8 Dimensiones',
         [('Version★',   'Estándar SAC', 'Versión de datos (Real / Plan / Forecast / Budget)'),
          ('CUENTAS',    'Genérica',     'Cuenta de costos/gastos SAP CO (CSKA/CSKB)'),
          ('Date★',      'Estándar SAC', 'Dimensión temporal Año/Mes'),
          ('AUDITORIA★', 'Genérica',     'Control y trazabilidad de carga de datos'),
          ('CEBES★',     'Genérica',     'Centro de Beneficio / Profit Center SAP CO'),
          ('CECOS',      'Genérica',     'Centro de Costo SAP CO-CCA (CSKS) — Metalmecánica'),
          ('MONEDA★',    'Genérica',     'Dimensión de moneda COP / USD / EUR'),
          ('SOCIEDAD★',  'Genérica',     'Sociedad / Company Code SAP FI (T001)')],
         'Unidades: Metalmecánica (principal) | Proceso: Control OPEX · CAPEX · Costos de producción'),
    ]:
        add_heading(doc, title, level=3, space_before=8)
        add_body(doc, mtype, italic=True, color=SAP_BLUE)
        mt = doc.add_table(rows=len(dims)+1, cols=3)
        header_row(mt, ['Dimensión', 'Tipo', 'Descripción / Equivalente SAP'], bg=SAP_BLUE)
        for j, (d, tp, de) in enumerate(dims, 1):
            is_shared = '★' in d
            data_row(mt, j, [d, tp, de], alt=(j % 2 == 0),
                     bold_col=0 if is_shared else None, center_cols=[1])
        add_body(doc, f'Aplica para: {extra}', italic=True, color=GRAY_MED, space_before=3, space_after=6)

    # Versiones
    add_heading(doc, 'Configuración de Versiones (Dimensión Version — Estándar SAC)', level=3, space_before=8)
    ver_t = doc.add_table(rows=5, cols=3)
    header_row(ver_t, ['Versión', 'Descripción', 'Uso en Fanalca'])
    for i, (v, d, u) in enumerate([
        ('Real/Actual', 'Ejecución contable real extraída de SAP ERP/FI',    'Cierre mensual · Seguimiento vs Plan'),
        ('Plan',        'Presupuesto anual aprobado por la alta dirección',   'Línea base para control presupuestal'),
        ('Forecast',    'Proyección actualizada intra-año (rolling forecast)', 'Re-estimación trimestral'),
        ('Budget',      'Anteproyecto / versión de trabajo presupuestal',     'Proceso de elaboración del presupuesto'),
    ], 1):
        data_row(ver_t, i, [v, d, u], alt=(i % 2 == 0), bold_col=0)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.2
    add_heading(doc, '5.2. Valor de Actividad de Configuración', level=2)
    cfg_t = doc.add_table(rows=11, cols=5)
    header_row(cfg_t, ['Actividad', 'Área', 'Estado', 'Prioridad', 'Notas'])
    for i, (a, ar, st, pr, n) in enumerate([
        ('Creación de los 3 modelos SAC', 'Modelado SAC', 'COMPLETO', 'Alta',
         '3 modelos Planning en ambiente de Desarrollo'),
        ('Configuración de dimensiones', 'Modelado SAC', 'COMPLETO', 'Alta',
         'Version, Date (estándar) + 10 dimensiones genéricas'),
        ('Carga de datos maestros', 'Datos Maestros', 'EN PROGRESO', 'Alta',
         'CSV/OData desde SAP ERP por dimensión (CLIENTES, CECOS, CEBES, CUENTA)'),
        ('Integración SAP ERP → SAC', 'Integración', 'EN PROGRESO', 'Alta',
         'Conexión Import/Live desde SAP BW/4HANA y S/4HANA'),
        ('Archivos de Validación SAC vs FI', 'Calidad de Datos', 'EN PROGRESO', 'Media',
         'Reconciliación automática de datos cargados vs. SAP FI'),
        ('Configuración de versiones', 'Modelado SAC', 'COMPLETO', 'Alta',
         'Real, Plan, Forecast, Budget configurados en dimensión Version'),
        ('Creación de Stories/Dashboards', 'Reporting', 'PENDIENTE', 'Media',
         'Visualizaciones por unidad de negocio y KPIs ejecutivos'),
        ('Perfilamiento de acceso y roles', 'Seguridad', 'PENDIENTE', 'Alta',
         'Roles SAC + restricciones por SOCIEDAD (Data Access Control)'),
        ('Pruebas de Usuario UAT', 'Calidad', 'PENDIENTE', 'Alta',
         'Validación con key users de las 3 unidades de negocio'),
        ('Formación y capacitación', 'Change Mgmt', 'PENDIENTE', 'Media',
         'Talleres SAC para usuarios finales y administradores'),
    ], 1):
        data_row(cfg_t, i, [a, ar, st, pr, n], alt=(i % 2 == 0),
                 bold_col=0, center_cols=[2, 3])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.3
    add_heading(doc, '5.3. Roles', level=2)
    rol_t = doc.add_table(rows=8, cols=4)
    header_row(rol_t, ['Rol SAC', 'Responsabilidades', 'Área', 'Asignación Fanalca'])
    for i, (r, resp, ar, asig) in enumerate([
        ('SAC Administrator',
         'Gestión del tenant SAC, usuarios, roles, conexiones y ciclo de vida DEV→QA→PRD',
         'TI / BASIS', 'Equipo TI Fanalca'),
        ('SAC Modeler',
         'Creación y mantenimiento de modelos, dimensiones, medidas y data actions',
         'TI / Analytics', 'Consultor SAP + TI Fanalca'),
        ('SAC Story Developer',
         'Desarrollo de Stories, dashboards, KPIs ejecutivos y reportes operativos',
         'Analytics CoE', 'Consultor SAP / Analítica'),
        ('Finance Key User',
         'Valida modelos EEFF y Costos/Gastos; aprueba datos de cierre mensual; UAT',
         'Finanzas', 'Fanalca Finanzas'),
        ('Commercial Key User',
         'Valida modelo Ingresos; gestiona clientes, referencias y RATIOs; UAT',
         'Comercial / Ventas', 'Fanalca Comercial'),
        ('Data Owner',
         'Gobierno, calidad y aprobación de datos maestros por área (CECOS, CEBES, CLIENTES)',
         'Varias áreas', 'Dirección de cada área'),
        ('End User / Viewer',
         'Consumo de reportes y dashboards en SAC. Sin acceso de escritura a modelos',
         'Toda la organización', 'Usuarios finales Fanalca'),
    ], 1):
        data_row(rol_t, i, [r, resp, ar, asig], alt=(i % 2 == 0), bold_col=0)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5.4
    add_heading(doc, '5.4. Inventario de datos maestros susceptibles de migración', level=2)
    add_body(doc,
        'Los siguientes datos maestros deben ser cargados en SAC desde los sistemas SAP ERP '
        '(S/4HANA, FI, CO, SD, MM). La carga se realiza mediante archivos CSV o conectores '
        'OData según la dimensión. Los maestros de alta complejidad requieren coordinación '
        'con los equipos FI, CO y SD:')

    dm_t = doc.add_table(rows=11, cols=5)
    header_row(dm_t, ['Dato Maestro', 'Fuente SAP', 'Método de Carga', 'Alcance', 'Complejidad'])
    for i, (d, f, m, a, c) in enumerate([
        ('Clientes (CLIENTES)',            'SD – KNA1',       'OData / Import CSV',   'Supermotos / Autos Honda',     'ALTA'),
        ('Cuentas Contables (CUENTA)',     'FI – SKA1/SKB1',  'Import CSV / BW',      'Todas las sociedades',          'ALTA'),
        ('Cuentas de Costos (CUENTAS)',    'CO – CSKA/CSKB',  'Import CSV / BW',      'Fanalca Metalmecánica',         'ALTA'),
        ('Centros de Costo (CECOS)',       'CO – CSKS',       'Import CSV / OData',   'Fanalca Metalmecánica',         'ALTA'),
        ('Centros de Beneficio (CEBES)',   'CO-PCA – CEPC',   'Import CSV / OData',   'Todas las unidades',            'ALTA'),
        ('Sociedades (SOCIEDAD)',          'FI – T001',       'Import CSV manual',    'Todas las unidades',            'MEDIA'),
        ('Referencias Producto (REF)',     'MM – MARA/MVKE',  'Import CSV / BW',      'Supermotos / Autos Honda',     'ALTA'),
        ('Monedas (MONEDA)',               'FI – TCURC',      'Import CSV manual',    'Todas las unidades',            'BAJA'),
        ('Dim. Auditoría (AUDITORIA)',     'Definición manual','Carga manual SAC',    'Equipo TI',                     'BAJA'),
        ('Versiones (VERSION)',            'Estándar SAC',    'Configuración SAC',    'Equipo implementación',         'BAJA'),
    ], 1):
        data_row(dm_t, i, [d, f, m, a, c], alt=(i % 2 == 0), bold_col=0, center_cols=[4])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_info_box(doc, 'Prioridad de carga de datos maestros',
        ['FASE 1 — Crítico (pre Go-Live): CEBES, SOCIEDAD, MONEDA, Versiones',
         'FASE 2 — Alto (configuración modelos): CECOS, CUENTA, CUENTAS, CLIENTES',
         'FASE 3 — Medio (historiales): REFERENCIA, datos transaccionales 2024-2025',
         'Todos los datos maestros deben alinearse con definiciones de módulos SD/FI/CO antes de carga en SAC'],
        bg_title=SAP_GREEN)

    # 5.5
    add_heading(doc, '5.5. Componente de Solución Técnica Relacionada', level=2)
    tech_t = doc.add_table(rows=9, cols=4)
    header_row(tech_t, ['Componente', 'Descripción', 'Entorno', 'Versión / Estado'])
    for i, (c, d, e, v) in enumerate([
        ('SAP Analytics Cloud (SAC)',
         'Plataforma SaaS en SAP BTP. Región us10 (América). 3 tenants: DEV / QA / PRD',
         'SAP BTP – SaaS', 'PRD · QA · DEV — Activo'),
        ('SAP S/4HANA / ERP',
         'Sistema fuente principal. Módulos FI, CO, SD, MM como fuentes de datos primarias',
         'On-Premise Colombia', 'Release SAP compatible'),
        ('SAP BW/4HANA',
         'Capa Data Warehouse para consolidación, transformación y carga de datos hacia SAC',
         'On-Premise / Cloud', 'BW/4HANA 2.0+'),
        ('SAP Data Services / SLT',
         'Replicación y transformación de datos desde SAP ERP hacia BW y SAC',
         'On-Premise', 'SAP DS 4.3+'),
        ('SAC – Import Connection',
         'Conector batch para carga programada: archivos CSV, OData y conectores SAP BW',
         'SAC (cloud)', 'Estándar SAC'),
        ('SAC – Live Connection',
         'Conexión en tiempo real hacia SAP BW/4HANA o SAP HANA para reportes live',
         'SAC ↔ BW/HANA', 'Requiere SAP Connector'),
        ('SAP Identity Auth. (IAS)',
         'Single Sign-On (SSO) para acceso de usuarios Fanalca a SAC',
         'SAP BTP', 'SAP IAS Cloud — Activo'),
        ('SAC Lifecycle Management',
         'Gestión del ciclo de vida y transporte de contenido: DEV → QA → PRD',
         'SAC', 'SAC Lifecycle Mgmt'),
    ], 1):
        data_row(tech_t, i, [c, d, e, v], alt=(i % 2 == 0), bold_col=0, center_cols=[2])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Nota final
    add_info_box(doc, 'Próximos pasos — Proyecto SAP BUILD Fanalca',
        ['1.  Revisión y aprobación de este documento de diseño por parte del comité del proyecto',
         '2.  Carga de datos maestros Fase 1: CEBES, SOCIEDAD, MONEDA, VERSION en SAC',
         '3.  Configuración de conexiones SAC ↔ SAP BW/4HANA y S/4HANA',
         '4.  Carga de historiales 2024-2025 para análisis Real vs. Plan',
         '5.  Desarrollo de Stories/Dashboards por unidad de negocio',
         '6.  Pruebas UAT con Key Users de las 3 unidades de negocio',
         '7.  Formación a usuarios finales y pase a producción (Go-Live)'],
        bg_title=SAP_DARK)

    doc.save(doc_path)
    print(f'Documento: {doc_path}')

# ─── MAIN ─────────────────────────────────────────────────────────────────────
if __name__ == '__main__':
    BASE = '/home/user/Johnmeg'
    fp = os.path.join(BASE, 'arquitectura_fanalca.png')
    dp = os.path.join(BASE, 'Diseno_Modelos_SAC_Fanalca_v1.docx')
    print('Generando flowchart...')
    generar_flowchart(fp)
    print('Generando documento Word...')
    generar_doc(fp, dp)
    print('Listo!')
