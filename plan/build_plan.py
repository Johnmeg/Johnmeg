# -*- coding: utf-8 -*-
"""Genera el Plan de seguimiento de la configuración de SAC (Excel).
3 sociedades (Fanalca, Ciudad Limpia, Transprensa) x 3 modelos (Ingresos,
Costos/Gastos, Estados Financieros). Modelos + dimensiones + vinculación con
SAP DataSphere = base ya completada. El plan cubre las actividades restantes."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from datetime import date, timedelta

# ---------- paleta ----------
NAVY="1F3864"; NAVY2="2E4C7E"; BLUE="2E75B6"; STEEL="8EAADB"
HFILL=PatternFill("solid", fgColor=NAVY)
WHITE=Font(name="Calibri", color="FFFFFF", bold=True, size=10)
THIN=Side(style="thin", color="BFBFBF")
BORDER=Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLEFONT=Font(name="Calibri", bold=True, size=15, color="FFFFFF")

TRACK_COLOR={"Base":"8EAADB","Transversal":"BFBFBF","Fanalca":NAVY,
             "Ciudad Limpia":"2E7D32","Transprensa":"C55A11"}
SOC_FILL={"Fanalca":"D6DCE9","Ciudad Limpia":"E2EFDA","Transprensa":"FCE4D6",
          "Transversal":"EDEDED"}
ESTADOS={  # texto -> (relleno, color fuente)
 "Completado":("C6EFCE","006100"),
 "En progreso":("BDD7EE","1F3864"),
 "No iniciado":("EDEDED","595959"),
 "En riesgo":("FFE699","7F6000"),
 "Bloqueado":("FFC7CE","9C0006")}
FASES=["0. Base configurada","1. Integración de datos","2. Plantillas de carga",
       "3. Cálculos (Data Actions)","4. Versiones y escenarios",
       "5. Reportes y dashboards","6. Seguridad y roles","7. Pruebas y UAT",
       "8. Despliegue y estabilización"]

# ---------- datos específicos por sociedad ----------
SOC={
"Fanalca":{
 "resp":"Planeación Fanalca",
 "maestros":"Validar cuentas, líneas de negocio (motos CKD/CBU, autos, repuestos, servicios), CEBE y CECO importados",
 "ing_pl":"Plantilla de ingresos por línea (motos CKD/CBU, autos, repuestos, servicios) por CEBE y moneda",
 "cg_pl":"Plantilla de gastos (paquete UNOE→Gastos) y CECOS transversales (PDB, PPC)",
 "ing_da":"Data Action de ingresos: P×Q por línea, precios y TRM, con crecimiento y estacionalidad mensual",
 "cg_da":"Data Action de distribución de CECOS transversales (PDB, PPC) y costeo por línea",
 "ver":"Versiones Actual/Presupuesto/Forecast y escenarios Pesimista/Optimista/Medio y Plan Bancos 1-2; bloqueo de datos",
 "dash":"Dashboard Fanalca: ingresos por línea, margen, EBITDA, escenarios y sensibilidad a la TRM"},
"Ciudad Limpia":{
 "resp":"Planeación C. Limpia",
 "maestros":"Validar cuentas, componentes de servicio (CEBE), CECO por vehículo y clientes de CGS/RH",
 "ing_pl":"Plantilla de ingresos ordinarios por componente (recolección, barrido, corte, poda) y P×Q para CGS/RH",
 "cg_pl":"Plantilla de mano de obra (base de personal de TH), flota (CECO por vehículo) y gastos por ~17 áreas",
 "ing_da":"Data Action de ingresos por componente y P×Q (tarifa × toneladas/kilos) por cliente y tipo de residuo",
 "cg_da":"Data Action de mano de obra (~80% del costo), costos de flota y consolidación de gastos por área",
 "ver":"Versiones Actual/Presupuesto/Forecast con regla de reemplazo del dato autorizado; bloqueo de datos",
 "dash":"Dashboard Ciudad Limpia: KPIs operativos (toneladas, km, usuarios, árboles), ejecución por componente y balance por municipio (Bogotá/Cali)"},
"Transprensa":{
 "resp":"Planeación Transprensa",
 "maestros":"Validar cuentas, regionales (21), tipos de servicio, CRM (46) y clientes especiales (Ingreso, T1, Soy Más)",
 "ing_pl":"Plantilla P×Q por regional (21) y servicio (paqueteo, masivo, dedicados, almacenamiento) y clientes especiales (Ingreso, T1, Soy Más)",
 "cg_pl":"Plantilla de fletes (expedición/reexpedición), mano de obra variable por kg, comisiones y arrendamientos",
 "ing_da":"Data Action de ingresos: P×Q desde histórico N-1 (Silotrans), crecimiento en volumen y tarifa (IPC+ICTC), estacionalidad y crédito/contado por CRM",
 "cg_da":"Data Action de fletes por ruta y margen, mano de obra por índice persona/kg y gastos fijos (N-1 × IPC)",
 "ver":"Versiones Presupuesto Comercial, Presupuesto Financiero y Forecast; simulaciones (nueva licitación / cliente grande); bloqueo de datos",
 "dash":"Dashboard Transprensa: kilos e ingresos por regional y servicio, margen bruto y operativo (PCGA)"},
}

# ---------- construcción de la lista de actividades ----------
# cada fila: (fase, soc, modelo, actividad, detalle, resp, dur, estado, avance, track, late)
rows=[]
def add(fase,soc,modelo,act,det,resp,dur,estado="No iniciado",av=0,track=None,late=False,dep=""):
    rows.append(dict(fase=fase,soc=soc,modelo=modelo,act=act,det=det,resp=resp,
                     dur=dur,estado=estado,av=av,track=track or soc,late=late,dep=dep))

# Base (completada)
add(FASES[0],"Transversal","Los 3 modelos","Modelos de planeación creados en SAC",
    "9 modelos (Ingresos, Costos/Gastos y EEFF por sociedad) creados y habilitados","Consultor SAC",10,"Completado",100,"Base")
add(FASES[0],"Transversal","Los 3 modelos","Dimensiones y jerarquías configuradas",
    "Cuenta, Sociedad, CEBE, CECO, Cliente, Versión, Fecha y Auditoría con sus jerarquías","Consultor SAC",8,"Completado",100,"Base")
add(FASES[0],"Transversal","N/A","Vinculación con SAP DataSphere establecida",
    "Conexión SAC–DataSphere operativa y modelos de datos (CDS) expuestos","Consultor DataSphere",5,"Completado",100,"Base")
# Transversal temprano
add(FASES[1],"Transversal","N/A","Gobierno del dato y calendario de refresco",
    "Catálogo de flujos de DataSphere, frecuencia de carga del real y responsables","Consultor DataSphere",4,track="Transversal")
add(FASES[6],"Transversal","N/A","Matriz de roles y licencias",
    "Definir roles (planificación vs. lectura BI) y asignación de licencias por usuario","PMO",4,track="Transversal")
add(FASES[7],"Transversal","N/A","Estrategia de pruebas y criterios de aceptación",
    "Plan de pruebas unitarias, de integración y UAT; criterios de salida","PMO",3,track="Transversal")

for s in ["Fanalca","Ciudad Limpia","Transprensa"]:
    D=SOC[s]; R=D["resp"]
    # F1 integración
    add(FASES[1],s,"N/A","Validar maestros importados desde DataSphere",D["maestros"],"Consultor DataSphere",4,dep="Base")
    add(FASES[1],s,"Los 3 modelos","Configurar flujos de datos reales (Actual)",
        "Import/flows del real desde DataSphere a los 3 modelos y programación de refresco","Consultor DataSphere",5,dep="Base")
    add(FASES[1],s,"Estados Financieros","Cuadre de saldos reales vs. ERP",
        "Validación de integridad del real cargado frente al ERP",R+" / Consultor DataSphere",4)
    # F2 plantillas (por modelo)
    add(FASES[2],s,"Ingresos","Plantilla de carga de Ingresos",D["ing_pl"],"Consultor SAC",4,dep="F1")
    add(FASES[2],s,"Costos y Gastos","Plantilla de carga de Costos y Gastos",D["cg_pl"],"Consultor SAC",4,dep="F1")
    add(FASES[2],s,"Estados Financieros","Plantilla de cuentas de Balance y Flujo",
        "Plantilla de cuentas propias de balance y flujo (activos, cartera, deuda, CAPEX) y ajustes","Consultor SAC",3,dep="F1")
    # F3 data actions (por modelo)
    add(FASES[3],s,"Ingresos","Data Action de cálculo de Ingresos",D["ing_da"],"Consultor SAC",5,dep="F2")
    add(FASES[3],s,"Costos y Gastos","Data Action de cálculo de Costos y Gastos",D["cg_da"],"Consultor SAC",5,dep="F2")
    add(FASES[3],s,"Estados Financieros","Data Action de consolidación del EEFF",
        "Consolidación Cross-Model (Ingresos + Costos/Gastos → P&G), depreciación, balance y flujo","Consultor SAC",5,dep="F3 Ing/CG")
    # F4 versiones
    add(FASES[4],s,"Los 3 modelos","Configurar versiones y escenarios",D["ver"],"Consultor SAC",4,dep="F3")
    # F5 reportes
    add(FASES[5],s,"Estados Financieros","Construir reportes financieros base",
        "P&G, Balance, Flujo de Caja y comparación Presupuesto vs. Real ($ y %)","Consultor SAC",5,dep="F3")
    add(FASES[5],s,"Los 3 modelos","Construir dashboard de la sociedad",D["dash"],"Consultor SAC",4,dep="F5")
    # F6 seguridad
    add(FASES[6],s,"Los 3 modelos","Configurar DAC, roles y pruebas de acceso",
        "Data Access Control por Sociedad/dimensiones; roles de la sociedad y validación de accesos","Consultor SAC",3,dep="Matriz roles")
    # F7 pruebas
    add(FASES[7],s,"Los 3 modelos","Pruebas unitarias de cálculos",
        "Validar Data Actions de Ingresos, Costos/Gastos y EEFF",R+" / Consultor SAC",4,dep="F3")
    add(FASES[7],s,"Los 3 modelos","Pruebas de integración (real + plan → EEFF)",
        "Flujo completo: carga, cálculo, consolidación y reporte",R,3,dep="F5,F6")
    add(FASES[7],s,"Los 3 modelos","UAT y ajustes",
        "Pruebas de aceptación con usuarios de la sociedad y ajustes finales",R,5,dep="Pruebas int.")
    # F8 despliegue
    add(FASES[8],s,"Los 3 modelos","Go-live y hypercare de la sociedad",
        "Puesta en productivo del presupuesto de la sociedad y soporte post-productivo","PMO",4,dep="UAT")

# Transversal tardío
add(FASES[8],"Transversal","N/A","Plan de capacitación y gestión del cambio",
    "Material y sesiones de formación a usuarios finales de las tres sociedades","PMO",6,track="Transversal",late=True)
add(FASES[8],"Transversal","N/A","Cierre del proyecto y transición a soporte",
    "Documentación final, lecciones aprendidas y entrega a soporte/hypercare","PMO",5,track="Transversal",late=True)

# ---------- programación de fechas ----------
def first_wd(d):
    while d.weekday()>=5: d+=timedelta(days=1)
    return d
def next_wd(d):
    d+=timedelta(days=1)
    return first_wd(d)
def add_wd(start,dur):
    d=start
    for _ in range(dur-1): d=next_wd(d)
    return d

cur={"Base":date(2026,7,6),"Transversal":date(2026,8,4),"Fanalca":date(2026,8,4),
     "Ciudad Limpia":date(2026,8,4),"Transprensa":date(2026,8,4)}
for r in rows:
    if r["late"]: continue
    t=r["track"]; s=first_wd(cur[t]); e=add_wd(s,r["dur"])
    r["ini"],r["fin"]=s,e; cur[t]=next_wd(e)
gmax=max(r["fin"] for r in rows if not r["late"])
lc=first_wd(gmax)
for r in rows:
    if r["late"]:
        s=first_wd(lc); e=add_wd(s,r["dur"]); r["ini"],r["fin"]=s,e; lc=next_wd(e)

print("Actividades:",len(rows),"| rango:",min(r['ini'] for r in rows),"->",max(r['fin'] for r in rows))

# =====================================================================
#  ESCRITURA DEL LIBRO EXCEL
# =====================================================================
CEN=Alignment(horizontal="center", vertical="center", wrap_text=True)
LEF=Alignment(horizontal="left", vertical="center", wrap_text=True)
TOPL=Alignment(horizontal="left", vertical="top", wrap_text=True)
F9=Font(name="Calibri", size=9)
F9B=Font(name="Calibri", size=9, bold=True)
lastrow=2+len(rows)
PLAN="'Plan de Actividades'!"

wb=Workbook()
# ---------------- Listas (validación) ----------------
ls=wb.active; ls.title="Listas"
cols=[("A","Estados",list(ESTADOS)),
      ("B","Sociedades",["Fanalca","Ciudad Limpia","Transprensa","Transversal"]),
      ("C","Modelos",["Ingresos","Costos y Gastos","Estados Financieros","Los 3 modelos","N/A"]),
      ("D","Fases",FASES)]
for col,title,items in cols:
    ls[col+"1"]=title; ls[col+"1"].font=F9B
    for i,v in enumerate(items): ls[col+str(2+i)]=v
ls.sheet_state="hidden"

# ---------------- Plan de Actividades ----------------
ws=wb.create_sheet("Plan de Actividades")
headers=["ID","Actividad","Fase","Sociedad","Modelo","Detalle / Entregable",
         "Responsable","Dur (d)","Inicio","Fin","% Avance","Estado","Dependencias","Observaciones"]
widths=[5,32,19,14,15,44,19,7,10,10,9,13,15,20]
for j,w in enumerate(widths): ws.column_dimensions[get_column_letter(1+j)].width=w
ws.column_dimensions[get_column_letter(15)].width=2   # espaciador (O)

# semanas del Gantt
start0=min(r["ini"] for r in rows); endN=max(r["fin"] for r in rows)
wk0=start0-timedelta(days=start0.weekday())
weeks=[]; w=wk0
while w<=endN: weeks.append(w); w+=timedelta(days=7)
GC0=16
for k in range(len(weeks)): ws.column_dimensions[get_column_letter(GC0+k)].width=3.1

# fila 1 títulos
ws.merge_cells("A1:N1"); t=ws["A1"]
t.value="PLAN DE SEGUIMIENTO — CONFIGURACIÓN DE SAP ANALYTICS CLOUD (SAC)"
t.font=TITLEFONT; t.alignment=Alignment("center","center"); t.fill=HFILL
g1=get_column_letter(GC0)+"1"; g2=get_column_letter(GC0+len(weeks)-1)+"1"
ws.merge_cells(f"{g1}:{g2}"); gt=ws[g1]; gt.value="CRONOGRAMA (semana de)"
gt.font=TITLEFONT; gt.alignment=Alignment("center","center"); gt.fill=PatternFill("solid",fgColor=NAVY2)
ws.row_dimensions[1].height=26

# fila 2 encabezados
for j,h in enumerate(headers):
    c=ws.cell(2,1+j,h); c.fill=HFILL; c.font=WHITE; c.alignment=CEN; c.border=BORDER
for k,wk in enumerate(weeks):
    c=ws.cell(2,GC0+k,wk.strftime("%d/%m")); c.fill=PatternFill("solid",fgColor=NAVY2)
    c.font=Font(color="FFFFFF", size=7); c.border=BORDER
    c.alignment=Alignment("center","center",text_rotation=90)
ws.row_dimensions[2].height=42

# filas de datos
for i,r in enumerate(rows):
    rr=3+i
    vals=[i+1, r["act"], r["fase"], r["soc"], r["modelo"], r["det"], r["resp"],
          r["dur"], r["ini"], r["fin"], r["av"]/100.0, r["estado"], r["dep"], ""]
    for j,v in enumerate(vals):
        c=ws.cell(rr,1+j,v); c.border=BORDER; c.font=F9
        c.alignment={0:CEN,1:TOPL,2:LEF,3:CEN,4:CEN,5:TOPL,6:LEF,7:CEN,8:CEN,9:CEN,10:CEN,11:CEN,12:LEF,13:TOPL}[j]
    ws.cell(rr,2).font=F9B
    ws.cell(rr,6).font=Font(name="Calibri", size=8.5)
    ws.cell(rr,4).fill=PatternFill("solid",fgColor=SOC_FILL[r["soc"]])
    ws.cell(rr,9).number_format="dd/mm/yy"; ws.cell(rr,10).number_format="dd/mm/yy"
    ws.cell(rr,11).number_format="0%"
    # barras Gantt
    col=TRACK_COLOR[r["track"]]
    for k,wk in enumerate(weeks):
        we=wk+timedelta(days=6); g=ws.cell(rr,GC0+k); g.border=BORDER
        if not (we<r["ini"] or wk>r["fin"]):
            g.fill=PatternFill("solid",fgColor=col)
    ws.row_dimensions[rr].height=42

ws.freeze_panes="C3"
ws.auto_filter.ref=f"A2:N{lastrow}"

# validación de datos
for rng,ref in [("D","B"),("E","C"),("L","A"),("C","D")]:
    n={"A":len(ESTADOS)+1,"B":5,"C":6,"D":len(FASES)+1}[ref]
    dv=DataValidation(type="list", formula1=f"Listas!${ref}$2:${ref}${n}", allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"{rng}3:{rng}{lastrow}")
# formato condicional de estado
for est,(fill,fc) in ESTADOS.items():
    ws.conditional_formatting.add(f"L3:L{lastrow}", CellIsRule(operator="equal",
        formula=[f'"{est}"'], fill=PatternFill("solid",fgColor=fill),
        font=Font(color=fc, bold=est in ("Completado","Bloqueado","En riesgo"))))
# barra de datos en % avance
ws.conditional_formatting.add(f"K3:K{lastrow}", DataBarRule(start_type="num",
    start_value=0, end_type="num", end_value=1, color=BLUE, showValue=True))

# ---------------- Resumen (tablero) ----------------
rs=wb.create_sheet("Resumen",0)
for col,wd in [("A",27),("B",14),("C",12),("D",19),("E",12),("F",10),("G",3)]:
    rs.column_dimensions[col].width=wd
rs.merge_cells("A1:F1"); h=rs["A1"]
h.value="TABLERO DE SEGUIMIENTO — CONFIGURACIÓN SAC"
h.font=TITLEFONT; h.fill=HFILL; h.alignment=Alignment("center","center"); rs.row_dimensions[1].height=28
rs.merge_cells("A2:F2"); s=rs["A2"]
s.value=("3 sociedades (Fanalca, Ciudad Limpia, Transprensa) × 3 modelos (Ingresos, Costos/Gastos, "
         "Estados Financieros). Base ya configurada: modelos, dimensiones y vinculación con SAP DataSphere. "
         "Este tablero resume el avance de las actividades de configuración restantes.")
s.font=Font(size=9, italic=True, color="595959"); s.alignment=Alignment("left","center",wrap_text=True)
rs.row_dimensions[2].height=42

def kpi(cellL,cellV,label,formula,fmt=None,big=True):
    a=rs[cellL]; a.value=label; a.font=F9B; a.alignment=Alignment("left","center")
    b=rs[cellV]; b.value=formula; b.font=Font(size=16,bold=True,color=NAVY); b.alignment=Alignment("center","center")
    if fmt: b.number_format=fmt
kpi("A4","B4","Total de actividades",f"=COUNTA({PLAN}B3:B{lastrow})")
kpi("A5","B5","Avance global",f"=AVERAGE({PLAN}K3:K{lastrow})","0%")
kpi("A6","B6","Completadas",f'=COUNTIF({PLAN}L3:L{lastrow},"Completado")')
kpi("A7","B7","En progreso",f'=COUNTIF({PLAN}L3:L{lastrow},"En progreso")')
kpi("A8","B8","No iniciadas",f'=COUNTIF({PLAN}L3:L{lastrow},"No iniciado")')
kpi("A9","B9","En riesgo / bloqueadas",f'=COUNTIF({PLAN}L3:L{lastrow},"En riesgo")+COUNTIF({PLAN}L3:L{lastrow},"Bloqueado")')
for rr in range(4,10):
    rs[f"A{rr}"].fill=PatternFill("solid",fgColor="F2F5FA"); rs[f"B{rr}"].fill=PatternFill("solid",fgColor="F2F5FA")
    rs[f"A{rr}"].border=BORDER; rs[f"B{rr}"].border=BORDER

# tabla por sociedad
rs["D4"]="Por sociedad"; rs["D4"].font=F9B
for j,t in enumerate(["Sociedad","Actividades","Avance"]):
    c=rs.cell(5,4+j,t); c.fill=HFILL; c.font=WHITE; c.alignment=CEN; c.border=BORDER
for i,soc in enumerate(["Fanalca","Ciudad Limpia","Transprensa","Transversal"]):
    rr=6+i
    rs.cell(rr,4,soc).border=BORDER; rs.cell(rr,4).font=F9
    rs.cell(rr,4).fill=PatternFill("solid",fgColor=SOC_FILL[soc])
    rs.cell(rr,5,f'=COUNTIF({PLAN}D3:D{lastrow},"{soc}")').border=BORDER; rs.cell(rr,5).alignment=CEN; rs.cell(rr,5).font=F9
    a=rs.cell(rr,6,f'=IFERROR(AVERAGEIF({PLAN}D3:D{lastrow},"{soc}",{PLAN}K3:K{lastrow}),0)')
    a.number_format="0%"; a.border=BORDER; a.alignment=CEN; a.font=F9

# tabla por estado (para gráfico de torta)
rs["D12"]="Por estado"; rs["D12"].font=F9B
for j,t in enumerate(["Estado","Cantidad"]):
    c=rs.cell(13,4+j,t); c.fill=HFILL; c.font=WHITE; c.alignment=CEN; c.border=BORDER
for i,est in enumerate(ESTADOS):
    rr=14+i
    rs.cell(rr,4,est).border=BORDER; rs.cell(rr,4).font=F9
    rs.cell(rr,4).fill=PatternFill("solid",fgColor=ESTADOS[est][0])
    rs.cell(rr,5,f'=COUNTIF({PLAN}L3:L{lastrow},"{est}")').border=BORDER; rs.cell(rr,5).alignment=CEN; rs.cell(rr,5).font=F9

# tabla por fase
rs["A12"]="Por fase"; rs["A12"].font=F9B
for j,t in enumerate(["Fase","Act.","Avance"]):
    c=rs.cell(13,1+j,t); c.fill=HFILL; c.font=WHITE; c.alignment=CEN; c.border=BORDER
for i,fa in enumerate(FASES):
    rr=14+i
    rs.cell(rr,1,fa).border=BORDER; rs.cell(rr,1).font=Font(size=8.5); rs.cell(rr,1).alignment=LEF
    rs.cell(rr,2,f'=COUNTIF({PLAN}C3:C{lastrow},"{fa}")').border=BORDER; rs.cell(rr,2).alignment=CEN; rs.cell(rr,2).font=F9
    a=rs.cell(rr,3,f'=IFERROR(AVERAGEIF({PLAN}C3:C{lastrow},"{fa}",{PLAN}K3:K{lastrow}),0)')
    a.number_format="0%"; a.border=BORDER; a.alignment=CEN; a.font=F9

# gráfico de barras: avance por sociedad
bar=BarChart(); bar.type="col"; bar.title="Avance por sociedad"; bar.height=6.5; bar.width=11
data=Reference(rs,min_col=6,min_row=6,max_row=9); cats=Reference(rs,min_col=4,min_row=6,max_row=9)
bar.add_data(data,titles_from_data=False); bar.set_categories(cats)
bar.y_axis.numFmt="0%"; bar.y_axis.majorGridlines=None; bar.legend=None
bar.dataLabels=DataLabelList(); bar.dataLabels.showVal=True; bar.dataLabels.numFmt="0%"
rs.add_chart(bar,"H4")
# gráfico de torta: estados
pie=PieChart(); pie.title="Actividades por estado"; pie.height=6.5; pie.width=11
pdata=Reference(rs,min_col=5,min_row=13,max_row=18); pcats=Reference(rs,min_col=4,min_row=14,max_row=18)
pie.add_data(pdata,titles_from_data=True); pie.set_categories(pcats)
pie.dataLabels=DataLabelList(); pie.dataLabels.showVal=True
rs.add_chart(pie,"H20")

# leyenda de colores del Gantt
rs["A24"]="Leyenda del cronograma (color de barra):"; rs["A24"].font=F9B
for i,(k,v) in enumerate([("Fanalca",TRACK_COLOR["Fanalca"]),("Ciudad Limpia",TRACK_COLOR["Ciudad Limpia"]),
                          ("Transprensa",TRACK_COLOR["Transprensa"]),("Transversal",TRACK_COLOR["Transversal"]),
                          ("Base (completada)",TRACK_COLOR["Base"])]):
    rr=25+i
    cc=rs.cell(rr,1,k); cc.font=F9; cc.alignment=Alignment("left","center")
    rs.cell(rr,2).fill=PatternFill("solid",fgColor=v); rs.cell(rr,2).border=BORDER
rs.sheet_view.showGridLines=False

wb.active=0
fn="Plan_Seguimiento_Configuracion_SAC.xlsx"
wb.save(fn)
print("GUARDADO:",fn,"| hojas:",wb.sheetnames,"| semanas Gantt:",len(weeks))
